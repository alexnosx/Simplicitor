"""Verify the bytes users review, including literal proposals and provenance."""
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook
import pytest

from extraction.models import ColumnSpec, ExtractionResult, FieldProposal, FieldResult, Issue
from extraction.xlsx_writer import read_candidate, write_candidate


def result_for(tmp_path, values, *, issues=(), coverage=None):
    """Two same-named attachments remain distinct records."""
    return ExtractionResult(
        ("one", "two"),
        {(source, column): field for source in ("one", "two")
         for column, field in values.items()},
        {"one": tmp_path / "a" / "same.docx", "two": tmp_path / "b" / "same.docx"},
        issues, coverage or {"one#p:0": "processed", "two#p:0": "processed"},
    )


def field(value, typed=None, flagged=False, issues=(), alternatives=()):
    return FieldResult(FieldProposal(value, f"Recorded {value}", "one#p:0"),
                       value if typed is None else typed, flagged, issues, alternatives)


def test_saved_values_types_evidence_and_selection_mapping(tmp_path):
    columns = (ColumnSpec("id", "ID", "Identifier"),
               ColumnSpec("total", "Total", "Gross amount", "decimal"),
               ColumnSpec("count", "Count", "Count", "integer"),
               ColumnSpec("due", "Due", "Due date", "date"))
    result = result_for(tmp_path, {"id": field("00123"),
                                  "total": field("GBP 2,400.00", Decimal("2400.00")),
                                  "count": field("12", 12),
                                  "due": field("March 15th, 2026", date(2026, 3, 15))})
    candidate = write_candidate(result, columns, tmp_path / "candidate.xlsx")
    wb = load_workbook(candidate.path)
    assert wb.sheetnames == ["Data", "Evidence"]
    assert list(wb["Data"].values)[0] == ("Record", "File", "ID", "Total", "Count", "Due")
    assert list(wb["Data"].values)[1] == ("one", "same.docx", "00123", 2400, 12,
                                        datetime(2026, 3, 15))
    assert wb["Data"]["C2"].data_type == "s"
    assert wb["Data"]["D2"].data_type == "n"
    assert wb["Data"]["D2"].number_format == "0.00"
    assert wb["Data"]["F2"].number_format == "yyyy-mm-dd"
    wb.close()
    cells = read_candidate(candidate)
    assert [(c.row, c.column) for c in cells] == [(r, c) for r in (2, 3) for c in range(3, 7)]
    total = cells[1]
    assert total.value == 2400 and total.data_type == "n"
    assert total.evidence[0]["Value"] == "GBP 2,400.00"
    assert total.evidence[0]["Quote"] == "Recorded GBP 2,400.00"
    assert total.evidence[0]["Data cell"] == "D2"
    assert cells[4].evidence[0]["Record"] == "two"
    assert str(tmp_path) not in repr([c.evidence for c in cells])
    # Grid is reconstructed from disk, rather than an in-memory result.
    wb = load_workbook(candidate.path)
    wb["Data"]["C2"] = "reviewed on disk"
    wb.save(candidate.path)
    wb.close()
    assert read_candidate(candidate)[0].value == "reviewed on disk"


@pytest.mark.parametrize("text", ["00123", "NaN", "inf", "1_000", "=1+1",
                                      '=HYPERLINK("https://example.com","click")'])
@pytest.mark.parametrize("flagged", [False, True])
def test_strings_stay_literal_on_both_sheets(tmp_path, text, flagged):
    candidate = write_candidate(result_for(tmp_path, {"x": field(text, flagged=flagged)}),
                                (ColumnSpec("x", "Value", "Value"),),
                                tmp_path / "candidate.xlsx")
    wb = load_workbook(candidate.path, data_only=False)
    for cell in (wb["Data"]["C2"], wb["Evidence"]["E2"]):
        assert cell.value == text and cell.data_type == "s"
        assert cell.hyperlink is None
    assert bool(wb["Data"]["C2"].fill.patternType) == flagged
    assert bool(candidate.issues) == flagged
    wb.close()


def test_missing_highlight_alternatives_and_file_coverage(tmp_path):
    alternative = FieldProposal("500", "Total 500", "one#p:1")
    result = result_for(tmp_path, {
        "missing": FieldResult(FieldProposal(None, "", ""), None, True, ("missing_value",)),
        "total": field("100", 100, alternatives=(alternative,)),
    }, issues=(Issue("request_failed", "one", "one#p:2", "A section failed."),),
       coverage={"one#p:2": "failed", "two#p:0": "processed"})
    candidate = write_candidate(result, (ColumnSpec("missing", "Missing", "Missing"),
                                         ColumnSpec("total", "Total", "Total", "integer")),
                                tmp_path / "candidate.xlsx")
    wb = load_workbook(candidate.path)
    assert wb["Data"]["C2"].value is None and wb["Data"]["C2"].fill.patternType == "solid"
    assert wb["Data"]["D2"].value == 100 and wb["Data"]["D2"].fill.patternType is None
    wb.close()
    evidence = read_candidate(candidate)[1].evidence
    assert [e["Value"] for e in evidence if e["Status"] != "coverage"] == ["100", "500"]
    assert any(e["Status"] == "coverage" and e["Anchor"] == "one#p:2" for e in evidence)
    assert any(i.code == "request_failed" for i in candidate.issues)


@pytest.mark.parametrize("bad", ["bad\x00text", "x" * 32768, "bad\ud800"],
                         ids=["control", "too-long", "surrogate"])
@pytest.mark.parametrize("location", ["value", "quote", "label"])
def test_unrepresentable_text_fails_without_a_candidate(tmp_path, bad, location):
    proposal = FieldProposal(bad if location == "value" else "ok",
                             bad if location == "quote" else "ok", "one#p:0")
    result = result_for(tmp_path, {"x": FieldResult(proposal, proposal.value, False)})
    path = tmp_path / "candidate.xlsx"
    with pytest.raises(ValueError, match="Excel"):
        write_candidate(result, (ColumnSpec("x", bad if location == "label" else "X", "X"),), path)
    assert not path.exists()
    assert not list(tmp_path.glob("*.tmp"))


@pytest.mark.parametrize("bad", [Decimal("123456789012345.67"), 12345678901234567,
                                   Decimal("NaN"), Decimal("Infinity"), Decimal("1e-400"),
                                   Decimal("1e1000000")])
def test_unsafe_numeric_precision_saves_literal_flag_and_other_fields(tmp_path, bad):
    result = result_for(tmp_path, {"x": field(str(bad), bad), "safe": field("12", 12)})
    candidate = write_candidate(result, (ColumnSpec("x", "X", "X", "decimal"),
                                         ColumnSpec("safe", "Safe", "Safe", "integer")),
                                tmp_path / "candidate.xlsx")
    wb = load_workbook(candidate.path)
    assert wb["Data"]["C2"].value == str(bad) and wb["Data"]["C2"].data_type == "s"
    assert wb["Data"]["C2"].fill.patternType == "solid"
    assert wb["Data"]["D2"].value == 12 and wb["Data"]["D2"].data_type == "n"
    assert wb["Data"]["D2"].fill.patternType is None
    assert wb["Evidence"]["E2"].value == str(bad)
    assert wb["Evidence"]["H2"].value == "flagged"
    assert wb["Evidence"]["I2"].value == "excel_precision"
    wb.close()
    assert any(issue.code == "excel_precision" for issue in candidate.issues)
    assert read_candidate(candidate)[0].evidence[0]["Issue"] == "excel_precision"
    assert result.fields[("one", "x")].flagged is False


@pytest.mark.parametrize("kind,value", [("integer", "000452"), ("decimal", "007")])
def test_leading_zero_proposal_is_saved_as_highlighted_literal(tmp_path, kind, value):
    from extraction.grounding import validate_field
    from extraction.models import SourceUnit
    column = ColumnSpec("x", "ID", "Identifier", kind)
    unit = SourceUnit("one#p:0", f"ID {value}", "paragraph", 0)
    validated = validate_field(FieldProposal(value, unit.text, unit.anchor),
                               column, {unit.anchor: unit})
    candidate = write_candidate(result_for(tmp_path, {"x": validated}), (column,),
                                tmp_path / "candidate.xlsx")
    wb = load_workbook(candidate.path)
    assert wb["Data"]["C2"].value == value and wb["Data"]["C2"].data_type == "s"
    assert wb["Data"]["C2"].fill.patternType == "solid"
    wb.close()
    assert read_candidate(candidate)[0].evidence[0]["Issue"] == "leading_zero"


def test_excel_precision_preserves_verbatim_currency_and_requires_acknowledgement(tmp_path):
    from extraction.grounding import validate_field
    from extraction.jobs import save_candidate
    from extraction.models import SourceUnit
    value = "GBP 123456789012345.67"
    column = ColumnSpec("x", "Amount", "Amount", "decimal")
    unit = SourceUnit("one#p:0", value, "paragraph", 0)
    validated = validate_field(FieldProposal(value, value, unit.anchor),
                               column, {unit.anchor: unit})
    assert not validated.flagged
    candidate = write_candidate(result_for(tmp_path, {"x": validated}), (column,),
                                tmp_path / "candidate.xlsx")
    assert read_candidate(candidate)[0].value == value
    with pytest.raises(ValueError, match="acknowledg"):
        save_candidate(candidate, tmp_path / "output.xlsx", False)
    save_candidate(candidate, tmp_path / "output.xlsx", True)


def test_failed_rerun_invalidates_old_candidate(tmp_path, monkeypatch):
    from openpyxl.workbook.workbook import Workbook
    path = tmp_path / "candidate.xlsx"
    result = result_for(tmp_path, {"x": field("old")})
    columns = (ColumnSpec("x", "X", "X"),)
    old = write_candidate(result, columns, path)
    def fail_save(self, destination):
        raise OSError("disk full")
    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(OSError):
        write_candidate(result, columns, path)
    assert not path.exists()
    with pytest.raises(OSError):
        read_candidate(old)
    assert not list(tmp_path.glob("*.tmp"))


def test_successful_rerun_replaces_saved_values(tmp_path):
    path = tmp_path / "candidate.xlsx"
    columns = (ColumnSpec("x", "X", "X"),)
    write_candidate(result_for(tmp_path, {"x": field("old")}), columns, path)
    new = write_candidate(result_for(tmp_path, {"x": field("new")}), columns, path)
    assert read_candidate(new)[0].value == "new"


def test_real_job_candidate_saves_and_reopens_with_acknowledgement(tmp_path):
    from extraction.jobs import create_job, discard_job, save_candidate
    job = create_job(tmp_path / "app")
    result = result_for(tmp_path, {"x": field("=1+1", flagged=True)})
    candidate = write_candidate(result, (ColumnSpec("x", "X", "X"),), job / "candidate.xlsx")
    destination = tmp_path / "export.xlsx"
    with pytest.raises(ValueError, match="acknowledg"):
        save_candidate(candidate, destination, False)
    save_candidate(candidate, destination, True)
    wb = load_workbook(destination)
    assert wb["Data"]["C2"].value == "=1+1" and wb["Data"]["C2"].data_type == "s"
    assert wb["Evidence"]["E2"].value == "=1+1"
    wb.close()
    discard_job(job)
    assert destination.exists() and not job.exists()


def test_failed_validation_rerun_removes_old_review(tmp_path):
    columns = (ColumnSpec("x", "X", "X"),)
    path = tmp_path / "candidate.xlsx"
    write_candidate(result_for(tmp_path, {"x": field("old")}), columns, path)
    with pytest.raises(ValueError, match="Excel"):
        write_candidate(result_for(tmp_path, {"x": field("bad\x00text")}), columns, path)
    assert not path.exists()


def test_candidate_writer_refuses_source_before_invalidating_it(tmp_path):
    path = tmp_path / "candidate.xlsx"
    path.write_bytes(b"read-only source")
    result = ExtractionResult(("one",), {("one", "x"): field("value")},
                              {"one": path}, (), {})
    with pytest.raises(ValueError, match="source"):
        write_candidate(result, (ColumnSpec("x", "X", "X"),), path)
    assert path.read_bytes() == b"read-only source"


def test_flagged_alternative_evidence_is_highlighted(tmp_path):
    alternative = FieldProposal("500", "Total 500", "one#p:1")
    result = result_for(tmp_path, {"x": field("100", flagged=True, issues=("conflict",),
                                             alternatives=(alternative,))})
    candidate = write_candidate(result, (ColumnSpec("x", "X", "X"),), tmp_path / "candidate.xlsx")
    wb = load_workbook(candidate.path)
    assert wb["Evidence"]["E3"].value == "500"
    assert wb["Evidence"]["E3"].fill.patternType == "solid"
    wb.close()


@pytest.mark.parametrize("value", [Decimal("0e-100000"), Decimal("1e-30"),
                                  Decimal("0.123456789012345"), Decimal("123456789012345")])
def test_safe_numeric_extremes_round_trip(tmp_path, value):
    candidate = write_candidate(result_for(tmp_path, {"x": field(str(value), value)}),
                                (ColumnSpec("x", "X", "X", "decimal"),),
                                tmp_path / "candidate.xlsx")
    assert Decimal(str(read_candidate(candidate)[0].value)) == value


def test_missing_field_fails_instead_of_dropping_it(tmp_path):
    with pytest.raises(ValueError, match="missing a requested field"):
        write_candidate(result_for(tmp_path, {}), (ColumnSpec("x", "X", "X"),),
                        tmp_path / "candidate.xlsx")
    assert not (tmp_path / "candidate.xlsx").exists()


def test_failed_candidate_replace_leaves_no_previous_review(tmp_path, monkeypatch):
    from extraction import xlsx_writer
    path = tmp_path / "candidate.xlsx"
    columns = (ColumnSpec("x", "X", "X"),)
    write_candidate(result_for(tmp_path, {"x": field("old")}), columns, path)
    def fail(source, destination):
        raise PermissionError("replace failed")
    monkeypatch.setattr(xlsx_writer.os, "replace", fail)
    with pytest.raises(OSError):
        write_candidate(result_for(tmp_path, {"x": field("new")}), columns, path)
    assert not path.exists() and not list(tmp_path.glob("*.tmp"))
