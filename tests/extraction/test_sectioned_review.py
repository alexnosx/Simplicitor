"""Sectioned sources require review without changing saved values or column types."""
from datetime import date
import json
from threading import Event

import pytest

from extraction.models import ColumnSpec
from extraction.pipeline import extract
from extraction.xlsx_writer import read_candidate, write_candidate
from tests.extraction.test_pipeline import document, profile


class SourceClient:
    def generate(self, model, prompt, system, **kwargs):
        payload = json.loads(prompt)
        fields = {}
        values = {"id": "00123", "total": "12,500.00", "due": "15 March 2026"}
        for column in payload["columns"]:
            value = values.get(column["id"])
            unit = next((u for u in payload["source_units"]
                         if value is not None and value in u["text"]), None)
            fields[column["id"]] = {"value": value if unit else None,
                                     "quote": unit["text"] if unit else "",
                                     "anchor": unit["anchor"] if unit else ""}
        return json.dumps({"records": [{"record_id": payload["record_id"], "fields": fields}]})


def test_sectioned_cells_are_flagged_with_saved_types_and_evidence_preserved(tmp_path):
    facts = "ID 00123. Total 12,500.00. Due 15 March 2026. "
    docs = (document([facts + "x" * 1200, "Tail " + "y" * 1200], "large"),
            document([facts], "small"))
    columns = (ColumnSpec("id", "ID", "Identifier."),
               ColumnSpec("total", "Total", "Amount.", "decimal"),
               ColumnSpec("due", "Due", "Payment date.", "date"),
               ColumnSpec("missing", "Missing", "Absent field."))
    calls = []
    result = extract(docs, columns, "", profile(1900), SourceClient(), Event(),
                     lambda source, *_: calls.append(source))
    assert calls.count("large") == 2 and calls.count("small") == 1
    assert any(i.code == "sectioned_source" and i.source_id == "large" for i in result.issues)
    assert not any(i.code == "sectioned_source" and i.source_id == "small" for i in result.issues)
    candidate = write_candidate(result, columns, tmp_path / "candidate.xlsx")
    from extraction.jobs import save_candidate
    with pytest.raises(ValueError, match="review|acknowledge"):
        save_candidate(candidate, tmp_path / "export.xlsx", acknowledge_issues=False)
    cells = read_candidate(candidate)
    large, small = cells[:4], cells[4:]
    assert large[0].value == small[0].value == "00123"
    assert large[1].value == small[1].value == 12500
    assert large[1].data_type == "n" and large[1].number_format == "0.00"
    assert large[2].value.date() == small[2].value.date() == date(2026, 3, 15)
    assert large[2].data_type == "d" and large[2].number_format == "yyyy-mm-dd"
    assert large[3].value is None
    for cell in large:
        entries = [e for e in cell.evidence if e["Status"] != "alternative"]
        assert len(entries) == 1 and entries[0]["Status"] == "flagged"
        assert "sectioned_source" in entries[0]["Issue"]
        assert not any(e["Status"] == "coverage" for e in cell.evidence)
    assert all(c.evidence[0]["Status"] == "verified" for c in small[:3])
    assert "sectioned_source" not in small[3].evidence[0]["Issue"]
    assert "missing_value" in large[3].evidence[0]["Issue"]
    from openpyxl import load_workbook
    book = load_workbook(candidate.path)
    assert all(book["Data"].cell(2, c.column).fill.patternType == "solid" for c in large)
    assert all(book["Data"].cell(3, c.column).fill.patternType is None for c in small[:3])
    book.close()


def test_sectioned_review_keeps_invalid_literal_proposals_and_original_issues(tmp_path):
    from extraction.models import Issue
    from tests.extraction.test_xlsx_writer import field, result_for
    result = result_for(tmp_path, {"x": field("=1+1", flagged=True, issues=("conversion_failed",))},
                        issues=(Issue("sectioned_source", "one", "", "Review every value."),))
    candidate = write_candidate(result, (ColumnSpec("x", "Amount", "Amount.", "decimal"),),
                                tmp_path / "candidate.xlsx")
    cell = read_candidate(candidate)[0]
    assert cell.value == "=1+1" and cell.data_type == "s"
    assert cell.evidence[0]["Status"] == "flagged"
    assert "sectioned_source" in cell.evidence[0]["Issue"]
    assert "conversion_failed" in cell.evidence[0]["Issue"]
