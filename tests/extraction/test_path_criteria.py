"""CLI path gates depend on reopened cells, independently of aggregate fractions."""
import json

from docx import Document
from openpyxl import load_workbook
import pytest

from tests.extraction.test_evaluation import _pipeline_fixture


class SourceClient:
    """Stand in for Ollama using only source units and requested field IDs."""

    def __init__(self, *_args):
        pass

    def generate(self, model, prompt, system, **kwargs):
        payload = json.loads(prompt)
        fields = {}
        for column in payload["columns"]:
            value = "12,500.00" if column["id"] == "total" else "00123"
            unit = next((u for u in payload["source_units"] if value in u["text"]), None)
            fields[column["id"]] = {"value": value if unit else None,
                                     "quote": unit["text"] if unit else "",
                                     "anchor": unit["anchor"] if unit else ""}
        return json.dumps({"records": [{"record_id": payload["record_id"], "fields": fields}]})


@pytest.mark.parametrize("whole_fields,target,flagged,exit_code,aggregate_passed", [
    (100, "section", False, 1, True),  # Section requires every cell to be flagged.
    (100, "section", True, 0, True),
    (2, "section", True, 0, False),  # Section accuracy is reported, not gated.
    (2, "whole", True, 1, False),  # Whole-file accuracy still gates.
])
def test_cli_applies_each_saved_output_paths_criterion(tmp_path, monkeypatch,
        whole_fields, target, flagged, exit_code, aggregate_passed):
    from scripts import evaluate_extraction as cli
    cases = []
    for name in ("whole", "section"):
        folder = tmp_path / name
        folder.mkdir()
        _pipeline_fixture(folder, large=name == "section")
        case = json.loads((folder / "manifest.json").read_text())["cases"][0]
        case.update(id=name, file=f"{name}/one.docx", labels=f"{name}/labels.json")
        if name == "whole":
            case["columns"] = [{"id": f"id{i}", "label": f"ID {i}",
                                 "description": "Account identifier.", "kind": "text"}
                                for i in range(whole_fields)]
            (folder / "labels.json").write_text(json.dumps(
                {f"id{i}": "00123" for i in range(whole_fields)}))
        cases.append(case)
    (tmp_path / "manifest.json").write_text(json.dumps({"cases": cases}))
    monkeypatch.setattr(cli, "OllamaClient", SourceClient)
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})
    writer = cli.write_candidate

    def change_saved_cell(result, columns, path):
        candidate = writer(result, columns, path)
        if result.ordered_source_ids == (target,):
            book = load_workbook(candidate.path)
            book["Data"]["C2"] = "wrong saved identifier"
            book["Evidence"]["H2"] = "flagged" if flagged else "verified"
            book.save(candidate.path)
            book.close()
        return candidate

    monkeypatch.setattr(cli, "write_candidate", change_saved_cell)
    report = tmp_path / "report.json"
    assert cli.main(["--full-pipeline", "--manifest", str(tmp_path / "manifest.json"),
                     "--profiles", str(tmp_path / "whole/profiles.json"),
                     "--report", str(report)]) == exit_code
    candidate = json.loads(report.read_text())["candidates"][0]
    assert candidate["aggregate_passed"] is aggregate_passed
    assert candidate["passed"] is (exit_code == 0)
    assert candidate["paths"]["sectioned"]["passed"] is (target != "section" or flagged)
    assert candidate["paths"]["whole_file"]["passed"] is (target != "whole")


def test_sectioned_review_flags_cannot_hide_incomplete_coverage(tmp_path, monkeypatch):
    from scripts import evaluate_extraction as cli
    _pipeline_fixture(tmp_path)
    source = Document(tmp_path / "one.docx")
    source.add_paragraph("Unreadably large indivisible paragraph " + "x" * 70000)
    source.save(tmp_path / "one.docx")
    monkeypatch.setattr(cli, "OllamaClient", SourceClient)
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})
    report = cli.evaluate(tmp_path / "manifest.json", tmp_path / "profiles.json", "http://localhost",
                          full_pipeline=True, output_dir=tmp_path / "saved")
    candidate = report["candidates"][0]
    assert candidate["correct"] == 2 and candidate["unflagged_wrong"] == 0
    assert candidate["fixtures"][0]["request_error"] == "incomplete_coverage"
    assert not candidate["paths"]["sectioned"]["passed"] and not report["passed"]


def test_correct_but_unflagged_sectioned_value_fails_the_cli(tmp_path, monkeypatch):
    from scripts import evaluate_extraction as cli
    _pipeline_fixture(tmp_path, large=True)
    monkeypatch.setattr(cli, "OllamaClient", SourceClient)
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})
    writer = cli.write_candidate

    def remove_review_flag(*args):
        candidate = writer(*args)
        book = load_workbook(candidate.path)
        book["Evidence"]["H2"] = "verified"
        book.save(candidate.path)
        book.close()
        return candidate

    monkeypatch.setattr(cli, "write_candidate", remove_review_flag)
    report = tmp_path / "report.json"
    assert cli.main(["--full-pipeline", "--manifest", str(tmp_path / "manifest.json"),
                     "--profiles", str(tmp_path / "profiles.json"), "--report", str(report)]) == 1
    summary = json.loads(report.read_text())["candidates"][0]["paths"]["sectioned"]
    assert summary["correct"] == 2 and summary["unflagged_wrong"] == 0
    assert summary["unflagged_values"] == 1 and not summary["passed"]
