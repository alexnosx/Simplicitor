"""Reader contracts use real synthetic Office/PDF files."""
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter


def test_docx_keeps_paragraphs_table_cells_and_long_tail(tmp_path):
    from extraction.source_readers import read_source

    source = tmp_path / "invoice.docx"
    doc = Document()
    doc.add_paragraph("Invoice 00123")
    row = doc.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Total"
    row.cells[1].text = "USD 12,500.00"
    doc.add_paragraph("x" * 51000 + " END-OF-FILE")
    doc.save(source)
    before = source.read_bytes()
    result = read_source(source, "file-1")

    assert [u.anchor for u in result.units[:3]] == [
        "file-1#p:0", "file-1#t:0:r:0:c:0", "file-1#t:0:r:0:c:1"
    ]
    assert result.units[2].text == "USD 12,500.00"
    assert result.units[-1].text.endswith("END-OF-FILE")
    assert sum(len(u.text) for u in result.units) > 51000
    assert source.read_bytes() == before


def test_docx_merged_cells_are_not_repeated(tmp_path):
    from extraction.source_readers import read_source

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Single supplier"
    source = tmp_path / "merged.docx"
    doc.save(source)
    assert [u.text for u in read_source(source, "one").units] == ["Single supplier"]


def test_empty_pdf_pages_are_reported_not_dropped(tmp_path):
    from extraction.source_readers import SourceReadError, read_source

    source = tmp_path / "blank.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(source)
    with pytest.raises(SourceReadError) as exc:
        read_source(source, "one")
    assert [i.anchor for i in exc.value.issues] == ["one#page:1"]


def test_unsupported_and_malformed_files_fail_safely(tmp_path):
    from extraction.source_readers import SourceReadError, read_source

    for suffix in [".xlsx", ".docx", ".pdf"]:
        source = tmp_path / ("bad" + suffix)
        source.write_bytes(b"not a document")
        with pytest.raises(SourceReadError):
            read_source(source, "one")


def test_job_page_limit_and_distinct_source_ids(tmp_path):
    from extraction.source_readers import SourceReadError, read_sources

    sources = []
    for folder in ["first", "second"]:
        (tmp_path / folder).mkdir()
        source = tmp_path / folder / "same.docx"
        doc = Document()
        doc.add_paragraph("z" * 450000)
        doc.save(source)
        sources.append(source)
    result = read_sources(sources)
    assert [d.source_id for d in result] == ["source-1", "source-2"]
    assert [d.page_cost for d in result] == [150, 150]
    with pytest.raises(SourceReadError):
        read_sources([*sources, sources[0]])


def test_per_file_limit_is_checked_before_parsing(tmp_path):
    from extraction.source_readers import SourceReadError, read_source

    source = tmp_path / "oversize.docx"
    with source.open("wb") as stream:
        stream.truncate(50 * 1024 * 1024 + 1)
    with pytest.raises(SourceReadError, match="50 MiB"):
        read_source(source, "one")


def test_pdf_keeps_page_text_and_flags_only_near_zero_page(tmp_path):
    from extraction.source_readers import read_source
    from scripts.build_extraction_fixtures import write_pdf

    source = tmp_path / "pages.pdf"
    write_pdf(source, [["Invoice total: £12,500.00. The payment date is 15 March 2026."],
                       ["x" * 39], ["x" * 40]])
    result = read_source(source, "one")
    assert [u.anchor for u in result.units] == ["one#page:1", "one#page:2", "one#page:3"]
    assert "£12,500.00" in result.units[0].text
    assert [i.anchor for i in result.issues] == ["one#page:2"]


def test_simple_word_fields_are_disclosed_as_unsupported(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from extraction.source_readers import read_source
    doc = Document()
    paragraph = doc.add_paragraph("Invoice due date: ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "15 March 2026"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    source = tmp_path / "field.docx"
    doc.save(source)
    result = read_source(source, "one")
    assert any(i.code == "unsupported_structure" for i in result.issues)


def test_empty_docx_fails_before_model_work(tmp_path):
    from extraction.source_readers import SourceReadError, read_source
    source = tmp_path / "empty.docx"
    Document().save(source)
    with pytest.raises(SourceReadError, match="readable"):
        read_source(source, "one")


def test_docx_reads_header_footer_paragraphs_and_cells_once(tmp_path):
    from docx.enum.section import WD_SECTION_START
    from docx.shared import Inches
    from extraction.source_readers import read_source

    doc = Document()
    doc.add_paragraph("Services were provided to Blue Finch Ltd.")
    header = doc.sections[0].header
    header.paragraphs[0].text = "Alder Services Ltd issues invoice 00123."
    header.add_table(rows=1, cols=1, width=Inches(4)).cell(0, 0).text = "USD 12,500.00"
    doc.sections[0].footer.paragraphs[0].text = "Contact accounts@alder.example."
    doc.add_section(WD_SECTION_START.NEW_PAGE)  # Inherits the same header/footer parts.
    doc.sections[0].different_first_page_header_footer = True
    doc.sections[0].first_page_header.paragraphs[0].text = "First page invoice"
    doc.settings.odd_and_even_pages_header_footer = True
    doc.sections[0].even_page_footer.paragraphs[0].text = "Even page contact"
    source = tmp_path / "header-invoice.docx"
    doc.save(source)
    before = source.read_bytes()
    result = read_source(source, "one")
    by_text = {u.text: u for u in result.units}
    assert by_text["Alder Services Ltd issues invoice 00123."].anchor == "one#header:0:p:0"
    assert by_text["USD 12,500.00"].anchor == "one#header:0:t:0:r:0:c:0"
    assert by_text["Contact accounts@alder.example."].anchor == "one#footer:0:p:0"
    assert "#header:" in by_text["First page invoice"].anchor
    assert "#footer:" in by_text["Even page contact"].anchor
    assert sum(u.text == "Alder Services Ltd issues invoice 00123." for u in result.units) == 1
    assert len({u.anchor for u in result.units}) == len(result.units)
    assert not result.issues
    assert source.read_bytes() == before


def test_docx_header_only_text_is_readable_and_counts_toward_page_cost(tmp_path):
    from extraction.source_readers import read_source

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "h" * 3001
    source = tmp_path / "header-only.docx"
    doc.save(source)
    result = read_source(source, "one")
    assert result.page_cost == 2
    assert any(u.text == "h" * 3001 for u in result.units)
