"""Actual-file whole/sectioned corpus boundaries and the 300-page job contract."""
import json
from pathlib import Path
from threading import Event

from docx import Document
import pytest

from extraction.models import ColumnSpec, ExtractionProfile
from extraction.pipeline import ExtractionError, extract
from extraction.sectioning import make_sections
from extraction.source_readers import SourceReadError, read_source, read_sources
from extraction.xlsx_writer import read_candidate, write_candidate

ROOT = Path(__file__).parent / "fixtures"


def test_full_corpus_retains_original_cases_and_large_facts_cross_sections():
    original = json.loads((ROOT / "manifest.json").read_text())["cases"]
    full = json.loads((ROOT / "full-pipeline-manifest.json").read_text())["cases"]
    assert full[:len(original)] == original
    authored = json.loads((ROOT / "sectioned_cases.json").read_text())["cases"]
    assert len(full[len(original):]) == len(authored) == 4
    for case, text in zip(full[len(original):], authored):
        source = read_source(ROOT / case["file"], case["id"])
        columns = tuple(ColumnSpec(**c) for c in case["columns"])
        sections = make_sections((source,), ExtractionProfile("qwen3:8b"), columns=columns)
        assert len(sections) >= 3 and not any(s.excluded for s in sections)
        assert set(json.loads((ROOT / case["labels"]).read_text())) == {c.id for c in columns}
        locations = []
        for block in text["blocks"]:
            first = " ".join(block[0].split())
            unit = next(u for u in source.units if first in " ".join(u.text.split()))
            locations.append(next(index for index, s in enumerate(sections) if unit.anchor in s.unit_ids))
        assert len(set(locations)) == 3


def test_actual_300_page_job_extracts_and_301_pages_reject_before_model(tmp_path):
    path = tmp_path / "limit.docx"
    doc = Document()
    for index in range(300):
        text = "The account identifier is 00123. " if index == 0 else ""
        doc.add_paragraph(text + "a" * (3000 - len(text)))
    doc.save(path)
    extra = tmp_path / "extra.docx"
    doc = Document()
    doc.add_paragraph("Additional supporting source.")
    doc.save(extra)
    sources = read_sources((path,))
    assert sum(s.page_cost for s in sources) == 300
    columns = (ColumnSpec("id", "ID", "Account identifier."),)

    class Client:
        calls = 0

        def generate(self, model, prompt, system, **kwargs):
            self.calls += 1
            payload = json.loads(prompt)
            unit = next((u for u in payload["source_units"] if "00123" in u["text"]), None)
            return json.dumps({"records": [{"record_id": payload["record_id"], "fields": {
                "id": {"value": "00123" if unit else None,
                       "quote": "The account identifier is 00123." if unit else "",
                       "anchor": unit["anchor"] if unit else ""}}}]})

    client = Client()
    result = extract(sources, columns, "", ExtractionProfile("model"), client, Event())
    assert set(result.coverage.values()) == {"processed"} and client.calls > 1
    candidate = write_candidate(result, columns, tmp_path / "candidate.xlsx")
    assert read_candidate(candidate)[0].value == "00123"
    with pytest.raises(SourceReadError, match="300"):
        read_sources((path, extra))
    before = client.calls
    with pytest.raises(ExtractionError, match="300"):
        extract(sources + (read_source(extra, "extra"),), columns, "",
                ExtractionProfile("model"), client, Event())
    assert client.calls == before
