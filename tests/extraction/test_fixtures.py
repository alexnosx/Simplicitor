"""Fixture rendering preserves authored prose and real Word header/footer content."""
import json

from docx import Document

from scripts.build_extraction_fixtures import build


def test_narrative_invoice_renders_header_footer_without_label_table(tmp_path):
    case = {
        "id": "narrative", "format": "docx", "kind": "invoice", "narrative": True,
        "date_order": "DMY", "header": ["Alder Services Ltd submits invoice 00081."],
        "footer": ["Contact accounts@alder.example for payment questions."],
        "pages": [["We supplied advisory services to Blue Finch Ltd.",
                   "Payment of USD 12,500.00 is due on 15th March 2026."]],
        "labels": {"document_id": "00081"},
    }
    (tmp_path / "authored_cases.json").write_text(json.dumps({"cases": [case]}))
    build(tmp_path)
    doc = Document(tmp_path / "narrative.docx")
    assert doc.sections[0].header.paragraphs[0].text == "Alder Services Ltd submits invoice 00081."
    assert doc.sections[0].footer.paragraphs[0].text.startswith("Contact accounts@alder.example")
    assert not doc.tables
    assert doc.paragraphs[0].text == "We supplied advisory services to Blue Finch Ltd."
    assert json.loads((tmp_path / "labels/narrative.json").read_text()) == {"document_id": "00081"}
