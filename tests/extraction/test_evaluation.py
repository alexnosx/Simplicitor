"""Scoring is independent of grounding and expected labels never affect flags."""
from datetime import date

import pytest


@pytest.mark.parametrize("correct,wrong,passed", [
    (199, 1, True), (198, 2, True), (197, 3, False), (190, 2, True), (189, 0, False),
])
def test_quality_boundaries(correct, wrong, passed):
    from scripts.evaluate_extraction import passes_quality_gate
    assert passes_quality_gate(correct, wrong, 200) is passed


def test_real_quote_in_wrong_column_is_counted_unflagged():
    from extraction.grounding import validate_field
    from extraction.models import ColumnSpec, FieldProposal, SourceUnit
    from scripts.evaluate_extraction import ExpectedField, ScoredField, score_results

    unit = SourceUnit("invoice#p:0", "Invoice date: 15 March 2026", "paragraph", 0)
    field = validate_field(FieldProposal("15 March 2026", unit.text, unit.anchor),
                           ColumnSpec("due_date", "Due date", "", "date"),
                           {unit.anchor: unit})
    assert not field.flagged
    report = score_results({("invoice", "due_date"): ScoredField(field.data_value, False)},
                           {("invoice", "due_date"): ExpectedField(date(2026, 4, 15), "date")})
    assert report.correct == 0
    assert report.unflagged_wrong == 1


def test_missing_slots_and_correct_flag_review_burden():
    from scripts.evaluate_extraction import ExpectedField, ScoredField, score_results
    labels = {("one", "id"): ExpectedField("00123", "text"),
              ("two", "id"): ExpectedField("00456", "text")}
    report = score_results({("one", "id"): ScoredField("00123", True)}, labels)
    assert report.total == 2 and report.correct == 1
    assert report.unflagged_wrong == 0
    assert report.flagged_correct == 1
    assert report.flag_rate_correct == 1.0


def test_response_schema_keeps_all_values_as_literal_strings():
    from extraction.models import ColumnSpec, build_response_schema
    schema = build_response_schema((ColumnSpec("amount", "Amount", "", "decimal"),),
                                   ("one",))
    field = schema["properties"]["records"]["items"]["properties"]["fields"]
    assert field["required"] == ["amount"]
    assert field["properties"]["amount"]["properties"]["value"]["type"] == ["string", "null"]


def test_direct_call_sends_whole_file_schema_and_thinking_off(monkeypatch, tmp_path):
    from extraction.models import ColumnSpec, SourceDocument, SourceUnit
    from scripts.evaluate_extraction import call_ollama

    class FakeSession:
        trust_env = True
        body = None

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def post(self, url, json, timeout):
            self.body = json
            class Response:
                def raise_for_status(self):
                    pass

                def json(self):
                    return {"response": '{"records": []}'}
            return Response()

    session = FakeSession()
    monkeypatch.setattr("scripts.evaluate_extraction.requests.Session", lambda: session)
    source = SourceDocument("one", tmp_path / "source.docx", "source.docx", (
        SourceUnit("one#p:0", "BEGIN Total: USD 12,500.00", "paragraph", 0),
        SourceUnit("one#p:1", "END-OF-SOURCE", "paragraph", 1),
    ), 1)
    call_ollama("http://localhost:11434", "model", source,
                (ColumnSpec("total", "Total", "", "decimal"),), {"num_ctx": 16384})
    assert not session.trust_env
    assert session.body["think"] is False and session.body["stream"] is False
    assert "BEGIN" in session.body["prompt"] and "END-OF-SOURCE" in session.body["prompt"]
    assert isinstance(session.body["format"], dict)
    assert "labels" not in session.body["prompt"]


def test_remote_transport_is_rejected_before_a_request(tmp_path):
    from extraction.models import SourceDocument
    from scripts.evaluate_extraction import EvaluationError, call_ollama
    with pytest.raises(EvaluationError, match="loopback"):
        call_ollama("https://remote.example", "model",
                    SourceDocument("one", tmp_path / "source.docx", "source", (), 1), (), {})


def test_invalid_output_retains_every_field(tmp_path):
    from extraction.models import ColumnSpec, SourceDocument
    from scripts.evaluate_extraction import parse_fields
    source = SourceDocument("one", tmp_path / "source.docx", "source", (), 1)
    columns = (ColumnSpec("id", "ID", ""), ColumnSpec("total", "Total", "", "decimal"))
    for output in ["not JSON", '{"records": []}', '{"records":[{"record_id":"other","fields":{}}]}']:
        fields = parse_fields(output, source, columns)
        assert set(fields) == {"id", "total"}
        assert all(f.flagged for f in fields.values())


def test_invalid_evidence_keeps_identifiable_proposed_value(tmp_path):
    import json
    from extraction.models import ColumnSpec, SourceDocument
    from scripts.evaluate_extraction import parse_fields

    output = json.dumps({"records": [{"record_id": "one", "fields": {
        "id": {"value": "00123", "quote": None, "anchor": "one#p:0"}
    }}]})
    fields = parse_fields(output, SourceDocument("one", tmp_path / "a.docx", "a", (), 1),
                           (ColumnSpec("id", "ID", ""),))
    assert fields["id"].flagged
    assert fields["id"].data_value == "00123"


@pytest.mark.parametrize("extra_location", ["root", "record", "fields"])
def test_extra_schema_properties_flag_but_retain_known_proposal(tmp_path, extra_location):
    import json
    from extraction.models import ColumnSpec, SourceDocument, SourceUnit
    from scripts.evaluate_extraction import parse_fields
    payload = {"records": [{"record_id": "one", "fields": {
        "id": {"value": "00123", "quote": "ID 00123", "anchor": "one#p:0"}
    }}]}
    target = payload if extra_location == "root" else payload["records"][0]
    if extra_location == "fields":
        target = target["fields"]
    target["unexpected"] = "extra"
    source = SourceDocument("one", tmp_path / "a.docx", "a", (
        SourceUnit("one#p:0", "ID 00123", "paragraph", 0),
    ), 1)
    result = parse_fields(json.dumps(payload), source, (ColumnSpec("id", "ID", ""),))["id"]
    assert result.flagged
    assert result.data_value == "00123"


@pytest.mark.parametrize("value", ["", " ", "\t\n", "null", "NULL", " NuLl\t"])
def test_blank_model_values_are_absent_and_score_correctly(tmp_path, value):
    import json
    from extraction.models import ColumnSpec, SourceDocument
    from scripts.evaluate_extraction import ExpectedField, ScoredField, parse_fields, score_results

    response = json.dumps({"records": [{"record_id": "one", "fields": {
        "email": {"value": value, "quote": "", "anchor": ""}
    }}]})
    source = SourceDocument("one", tmp_path / "a.docx", "a", (), 1)
    field = parse_fields(response, source, (ColumnSpec("email", "Email", ""),))["email"]
    assert field.proposal.value is None and field.data_value is None
    assert field.flagged
    score = score_results({("one", "email"): ScoredField(field.data_value, field.flagged)},
                          {("one", "email"): ExpectedField(None, "text")})
    assert score.correct == 1 and score.flagged_expected_missing == 1


@pytest.mark.parametrize("value", ["NULL-009", "null reference", "Nullable Ltd"])
def test_text_containing_null_is_preserved_as_a_grounded_value(value):
    from extraction.grounding import validate_field
    from extraction.models import ColumnSpec, FieldProposal, SourceUnit
    unit = SourceUnit("one#p:0", value, "paragraph", 0)
    field = validate_field(FieldProposal(value, value, unit.anchor),
                           ColumnSpec("id", "ID", "Identifier."), {unit.anchor: unit})
    assert not field.flagged and field.data_value == value


def test_unknown_extra_record_flags_but_retains_the_identifiable_requested_proposal(tmp_path):
    import json
    from extraction.models import ColumnSpec, SourceDocument, SourceUnit
    from scripts.evaluate_extraction import parse_fields

    source = SourceDocument("one", tmp_path / "a.docx", "a", (
        SourceUnit("one#p:0", "ID 00123", "paragraph", 0),
    ), 1)
    response = json.dumps({"records": [
        {"record_id": "unknown", "fields": {}},
        {"record_id": "one", "fields": {
            "id": {"value": "00123", "quote": "ID 00123", "anchor": "one#p:0"}}},
    ]})
    field = parse_fields(response, source, (ColumnSpec("id", "ID", ""),))["id"]
    assert field.flagged and "invalid_response_schema" in field.issues
    assert field.data_value == "00123"


def test_actual_file_evaluation_counts_a_failed_request_instead_of_skipping_it(tmp_path, monkeypatch):
    import json
    from docx import Document
    from scripts import evaluate_extraction as cli

    doc = Document()
    doc.add_paragraph("ID 00123")
    doc.save(tmp_path / "one.docx")
    (tmp_path / "labels.json").write_text(json.dumps({"id": "00123"}))
    (tmp_path / "manifest.json").write_text(json.dumps({"cases": [{
        "id": "one", "file": "one.docx", "labels": "labels.json", "columns": [
            {"id": "id", "label": "ID", "description": "Identifier", "kind": "text"}]}]}))
    (tmp_path / "profiles.json").write_text(json.dumps({"candidates": [
        {"name": "qwen3:8b", "model": "qwen"}]}))
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})
    def fail_request(*args):
        raise cli.EvaluationError("ollama_timeout")
    monkeypatch.setattr(cli, "call_ollama", fail_request)
    report = cli.evaluate(tmp_path / "manifest.json", tmp_path / "profiles.json", "http://localhost")
    candidate = report["candidates"][0]
    assert candidate["total"] == 1 and candidate["correct"] == 0
    assert candidate["unflagged_wrong"] == 0
    assert candidate["fixtures"][0]["request_error"] == "ollama_timeout"


def _pipeline_fixture(tmp_path, *, large=False):
    import json
    from docx import Document
    doc = Document()
    doc.add_paragraph("The account identifier is 00123.")
    if large:
        for _ in range(100):
            doc.add_paragraph("Delivery documentation is reviewed by the operations team. " * 12)
    doc.add_paragraph("The gross total is USD 12,500.00.")
    doc.save(tmp_path / "one.docx")
    (tmp_path / "labels.json").write_text(json.dumps({"id": "00123", "total": "12500.00"}))
    (tmp_path / "manifest.json").write_text(json.dumps({"cases": [{
        "id": "one", "file": "one.docx", "labels": "labels.json", "columns": [
            {"id": "id", "label": "ID", "description": "Account identifier.", "kind": "text"},
            {"id": "total", "label": "Total", "description": "Gross total.", "kind": "decimal"}]}]}))
    (tmp_path / "profiles.json").write_text(json.dumps({"candidates": [
        {"name": "qwen3:8b", "model": "qwen"}]}))


class _GroundedClient:
    def __init__(self, *_args, **_kwargs):
        self.calls = []

    def generate(self, model, prompt, system, **kwargs):
        import json
        self.calls.append((prompt, kwargs))
        payload = json.loads(prompt)
        fields = {}
        for name, value in (("id", "00123"), ("total", "12,500.00")):
            unit = next((u for u in payload["source_units"] if value in u["text"]), None)
            fields[name] = {"value": value if unit else None,
                            "quote": unit["text"] if unit else "",
                            "anchor": unit["anchor"] if unit else ""}
        return json.dumps({"records": [{"record_id": payload["record_id"], "fields": fields}]})


@pytest.mark.parametrize("large,path", [(False, "whole_file"), (True, "sectioned")])
def test_full_pipeline_reads_actual_files_sections_and_scores_saved_cells(tmp_path, monkeypatch,
                                                                       large, path):
    from scripts import evaluate_extraction as cli
    _pipeline_fixture(tmp_path, large=large)
    client = _GroundedClient()
    monkeypatch.setattr(cli, "OllamaClient", lambda *_args, **_kwargs: client)
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})
    report = cli.evaluate(tmp_path / "manifest.json", tmp_path / "profiles.json", "http://localhost",
                          full_pipeline=True, output_dir=tmp_path / "saved")
    candidate = report["candidates"][0]
    assert candidate["correct"] == 2 and candidate["unflagged_wrong"] == 0
    assert candidate["paths"][path]["total"] == 2
    assert candidate["paths"][path]["passed"]
    fixture = candidate["fixtures"][0]
    assert fixture["path"] == path and fixture["saved"]
    assert fixture["coverage"]["processed"] > 0
    assert not fixture["coverage"].get("failed", 0)
    assert len(client.calls) > 1 if large else len(client.calls) == 1
    assert all(c[1]["think"] is False and c[1]["local_only"] for c in client.calls)
    assert all("12500.00" not in p for p, _ in client.calls)  # Scoring label never enters prompts.
    assert len(tuple((tmp_path / "saved").rglob("*.xlsx"))) == 1


def test_full_pipeline_scores_reopened_data_and_saved_evidence_flags(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from scripts import evaluate_extraction as cli
    _pipeline_fixture(tmp_path)
    monkeypatch.setattr(cli, "OllamaClient", _GroundedClient)
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})
    writer = cli.write_candidate

    def change_saved_data(*args, **kwargs):
        candidate = writer(*args, **kwargs)
        book = load_workbook(candidate.path)
        book["Data"]["C2"] = "wrong saved ID"
        book["Evidence"]["H3"] = "flagged"
        book.save(candidate.path)
        book.close()
        return candidate

    monkeypatch.setattr(cli, "write_candidate", change_saved_data)
    report = cli.evaluate(tmp_path / "manifest.json", tmp_path / "profiles.json", "http://localhost",
                          full_pipeline=True, output_dir=tmp_path / "saved")
    candidate = report["candidates"][0]
    assert candidate["correct"] == 1 and candidate["unflagged_wrong"] == 1
    assert candidate["flagged_correct"] == 1


@pytest.mark.parametrize("large,path", [(False, "whole_file"), (True, "sectioned")])
def test_full_pipeline_failed_save_retains_denominator_and_route(tmp_path, monkeypatch, large, path):
    from scripts import evaluate_extraction as cli
    _pipeline_fixture(tmp_path, large=large)
    monkeypatch.setattr(cli, "OllamaClient", _GroundedClient)
    monkeypatch.setattr(cli, "_model_details", lambda *_: {"parameter_count": 8_000_000_000})

    def fail_save(*args, **kwargs):
        raise OSError("synthetic save failure")

    monkeypatch.setattr(cli, "write_candidate", fail_save)
    report = cli.evaluate(tmp_path / "manifest.json", tmp_path / "profiles.json", "http://localhost",
                          full_pipeline=True, output_dir=tmp_path / "saved")
    candidate = report["candidates"][0]
    assert candidate["total"] == candidate["paths"][path]["total"] == 2
    assert candidate["correct"] == candidate["unflagged_wrong"] == 0
    assert not candidate["fixtures"][0]["saved"]
    assert candidate["fixtures"][0]["request_error"] == "full_pipeline_failed"
    assert not candidate["passed"] and not report["passed"]
