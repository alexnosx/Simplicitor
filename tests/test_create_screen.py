from dataclasses import replace
from docx import Document
from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QTableWidget
from extraction.models import ColumnSpec, ExtractionResult, FieldProposal, FieldResult, Issue
from extraction.source_readers import read_sources
from extraction.xlsx_writer import read_candidate, write_candidate

def test_step_indicator_keeps_state_and_layout(qtbot, isolated_profile):
    from app.widgets.step_indicator import StepIndicator
    steps = StepIndicator()
    qtbot.addWidget(steps)
    steps.set_steps(["Describe", "Choose columns", "Review and save"])
    steps.set_current(1)
    steps.set_steps(["Describe", "Choose columns", "Review and save"])  # same route: no reset
    assert steps.current() == 1
    assert [label.text() for label in steps._labels] == ["Describe", "Choose columns", "Review and save"]
    count = steps._layout.count()
    steps.set_steps(["Describe", "Save"])
    steps.set_steps(["Describe", "Choose columns", "Review and save"])
    assert steps._layout.count() == count  # no accumulating stretches
    assert steps.current() == 0
    assert steps._labels[0].property("current") is True


def test_anchor_issue_and_coverage_text_are_plain():
    from app.widgets.plain_labels import describe_anchor, describe_coverage, describe_issues
    assert describe_anchor("source-1#page:3") == "page 3"
    assert describe_anchor("source-1#p:0") == "paragraph 1"
    assert describe_anchor("source-2#t:0:r:1:c:2") == "table 1, row 2, column 3"
    assert describe_anchor("source-1#header:0:p:1") == "paragraph 2 of the header"
    assert describe_anchor("") == ""
    assert describe_issues("sectioned_source") == "Long document: check every value."
    assert describe_issues("leading_zero; excel_precision") == (
        "Starts with 0, so it was kept as text to keep every digit. "
        "Too many digits for an Excel number, so it was kept as text.")
    assert describe_issues("") == ""
    assert describe_issues("something_new") == "Check this value."
    assert describe_coverage("invoice-02.pdf", "source-2#page:3", "unreadable_page") == (
        "invoice-02.pdf, page 3: this page has too little text to read.")
    assert describe_coverage("a.docx", "source-1#", "skipped_text_box") == (
        "a.docx: text in text boxes was not read.")
    assert describe_coverage("a.docx", "source-1", "skipped_footnotes") == (
        "a.docx: footnotes and endnotes were not read.")


from pathlib import Path


def _screen(qtbot, isolated_profile):
    from app.widgets.create_screen import CreateScreen
    screen = CreateScreen(isolated_profile)
    qtbot.addWidget(screen)
    screen.set_model("qwen3:8b")
    screen.set_ollama_connected(True)
    return screen


def test_without_sources_all_types_and_create_file_uses_chosen_folder(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    assert [screen._type_combo.itemText(i) for i in range(screen._type_combo.count())] == [
        "Word", "PDF", "Excel", "PowerPoint"]
    assert screen._primary_btn.text() == "Create file"
    screen._save_dir = str(tmp_path / "chosen")
    screen._prompt_edit.setPlainText("A one-page project summary")
    screen._type_combo.setCurrentText("PDF")
    with qtbot.waitSignal(screen.generate_requested) as blocker:
        screen._primary_btn.click()
    assert blocker.args == ["PDF (.pdf)", str(tmp_path / "chosen"), "A one-page project summary"]


def test_sources_force_excel_and_wait_for_inspection(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "invoice.docx"
    source.write_bytes(b"x")
    ignored = tmp_path / "notes.txt"
    ignored.write_bytes(b"x")
    screen._type_combo.setCurrentText("Word")
    screen._prompt_edit.setPlainText("Invoice number and total")
    with qtbot.waitSignal(screen.inspect_requested):
        screen.set_paths((source, ignored))
    assert screen._source_list.count() == 1
    assert screen.file_type() == "Excel (.xlsx)"
    word = screen._type_combo.findText("Word")
    assert screen._type_combo.model().item(word).isEnabled()  # retained Write choice
    assert not screen._type_combo.isVisible()
    assert screen._sources_note.isVisible()
    assert screen._primary_btn.text() == "Reading files..."
    assert not screen._primary_btn.isEnabled()  # still reading the files
    screen.show_inspection(())
    assert screen._primary_btn.isEnabled()
    with qtbot.waitSignal(screen.propose_requested):
        screen._primary_btn.click()


def test_same_named_files_stay_separate(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    a, b = tmp_path / "a" / "invoice.pdf", tmp_path / "b" / "invoice.pdf"
    for path in (a, b):
        path.parent.mkdir()
        path.write_bytes(b"x")
    screen.set_paths((a, b))
    assert screen._source_list.count() == 2


def test_removing_all_sources_restores_every_type(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "invoice.pdf"
    source.write_bytes(b"x")
    screen.set_paths((source,))
    screen.set_paths(())
    screen.set_route("write")
    word = screen._type_combo.findText("Word")
    assert screen._type_combo.model().item(word).isEnabled()
    assert screen._primary_btn.text() == "Create file"
    assert screen._steps.current() == 0


def test_busy_freezes_setup_and_reports(qtbot, isolated_profile):
    screen = _screen(qtbot, isolated_profile)
    with qtbot.waitSignal(screen.busy_changed) as blocker:
        screen.set_busy(True)
    assert blocker.args == [True]
    assert not screen._prompt_edit.isEnabled() and not screen._primary_btn.isEnabled()


def test_show_status_accepts_secondary_and_tooltip(qtbot, isolated_profile):
    screen = _screen(qtbot, isolated_profile)
    screen.show_status("Saved", False, "summary.pdf", "C:/out/summary.pdf")


def test_source_changes_are_ignored_while_busy(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "invoice.pdf"
    source.write_bytes(b"x")
    screen.set_busy(True)
    requests = []
    screen.inspect_requested.connect(requests.append)
    screen.set_paths((source,))
    assert screen._source_list.count() == 0 and requests == []


def test_sources_do_not_overwrite_the_retained_write_type(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    screen._type_combo.setCurrentText("PowerPoint")
    seen = []
    screen.file_type_changed.connect(seen.append)
    source = tmp_path / "invoice.docx"
    source.write_bytes(b"x")
    screen.set_paths((source,))
    assert seen == []
    screen.show_inspection(())
    screen.set_route("write")
    assert screen.file_type() == "PowerPoint (.pptx)"


def test_failed_inspection_ends_reading_and_keeps_setup(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    screen._prompt_edit.setPlainText("Invoice number")
    source = tmp_path / "invoice.docx"
    source.write_bytes(b"x")
    screen.set_paths((source,))
    assert not screen._primary_btn.isEnabled()
    screen.show_failure("Could not read the files.", ())
    screen.inspection_finished()
    assert screen._primary_btn.isEnabled() and screen._source_list.count() == 1


def test_cancel_is_outside_the_frozen_setup(qtbot, isolated_profile):
    screen = _screen(qtbot, isolated_profile)
    screen.set_busy(True)
    assert screen._cancel_btn.isEnabled()
    screen.set_busy(True, saving=True)
    assert not screen._cancel_btn.isEnabled()
    screen.set_busy(False)
    assert not screen._cancel_btn.isEnabled()

def test_saved_prompt_page_and_start_another_preserve_folder(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen._save_dir = str(tmp_path / "chosen")
    screen._prompt_edit.setPlainText("A synthetic summary")
    screen._type_combo.setCurrentText("PDF")
    screen.show_open_file_btn(str(tmp_path / "summary.pdf"))
    assert screen._pages.currentWidget() is screen._saved_page
    assert screen._steps.current() == 1
    assert screen._saved_label.text() == "Saved: summary.pdf"
    screen._another_btn.click()
    assert screen._pages.currentWidget() is screen._describe_page
    assert screen._prompt_edit.toPlainText() == ""
    assert screen.file_type() == "Word (.docx)"
    assert screen._save_dir == str(tmp_path / "chosen")
    assert screen._steps.current() == 0


def test_source_route_hides_save_folder_and_saved_reset_clears_sources(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    screen.set_paths((tmp_path / "invoice.docx",))
    assert not screen._save_folder.isVisible()
    screen.show_open_file_btn(str(tmp_path / "data.xlsx"))
    assert screen._steps.current() == 2
    screen._another_btn.click()
    assert screen.snapshot(False).paths == ()
    assert screen._save_folder.isVisible()

def _document(path):
    from extraction.models import SourceDocument, SourceUnit
    unit = SourceUnit("source-1#p:0", "Invoice number: 000731", "paragraph", 0, "source-1#p:0")
    return SourceDocument("source-1", path.resolve(), path.name, (unit,), 1, ())


def test_extract_before_reading_finishes_waits_and_resumes(qtbot, isolated_profile, tmp_path):
    from extraction.models import ColumnSpec
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "invoice.docx"
    source.write_bytes(b"x")
    screen._prompt_edit.setPlainText("Invoice number")
    screen.set_paths((source,))
    screen.set_columns((ColumnSpec("c1", "Invoice number", "The invoice number", "text"),))
    extracts = []
    screen.extract_requested.connect(extracts.append)
    screen._extract_btn.click()
    assert extracts == []  # waits for reading
    screen.show_inspection((_document(source),))
    assert extracts == []  # not during inspection_ready
    screen.resume_pending_confirmation()  # MainWindow calls this after the thread is cleared
    assert len(extracts) == 1


def test_columns_page_navigation_and_confirmation(qtbot, isolated_profile, tmp_path):
    from extraction.models import ColumnSpec
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "invoice.docx"
    source.write_bytes(b"x")
    screen._prompt_edit.setPlainText("Invoice number")
    screen.set_paths((source,))
    screen.show_inspection((_document(source),))
    screen.set_columns((ColumnSpec("c1", "Invoice number", "The invoice number", "text"),))
    assert screen._pages.currentWidget() is screen._columns_page
    assert screen._steps.current() == 1
    assert screen._extract_btn.text() == "Extract to Excel"
    with qtbot.waitSignal(screen.extract_requested):
        screen._extract_btn.click()
    assert screen.columns_confirmed()
    screen._back_btn.click()
    assert screen._pages.currentWidget() is screen._describe_page
    assert screen._steps.current() == 0


def test_review_uses_plain_words_and_resets_acknowledgement(qtbot, isolated_profile, tmp_path):
    from extraction.models import Candidate, Issue, ReviewCell
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    cell = ReviewCell(2, 3, "00123", "s", ({
        "Data cell": "C2", "Record": "source-1", "File": "invoice.docx", "Field": "Invoice number",
        "Value": "00123", "Quote": "Invoice number: 00123", "Anchor": "source-1#t:0:r:0:c:1",
        "Status": "flagged", "Issue": "sectioned_source"},))
    issues = (Issue("flagged_field", "source-1", "C2", "A field needs review before saving."),)
    candidate = Candidate("job", tmp_path / "candidate.xlsx", (tmp_path / "invoice.docx",), issues)
    screen.show_candidate(candidate, (cell,))
    assert screen._pages.currentWidget() is screen._review_page and screen._steps.current() == 2
    screen._grid.setCurrentCell(0, 1)
    text = screen._evidence.toPlainText()
    assert "Source text: Invoice number: 00123" in text
    assert "Found in: invoice.docx, table 1, row 1, column 2" in text
    assert "Check: Long document: check every value." in text
    assert "Anchor" not in text and "Evidence" not in text
    assert screen._ack_checkbox.isVisible() and not screen._save_btn.isEnabled()
    screen._ack_checkbox.setChecked(True)
    assert screen._save_btn.isEnabled()
    screen._prompt_edit.setPlainText("Changed request")
    assert screen.candidate() is None and not screen._ack_checkbox.isChecked()


import pytest


@pytest.mark.parametrize("highlighted,reading,label", [
    (True, False, "I have checked the highlighted values"),
    (False, True, "I have checked the reading problems"),
    (True, True, "I have checked the highlighted values and reading problems"),
    (False, False, ""),
])
def test_review_acknowledgement_labels_and_export_gate(qtbot, isolated_profile, tmp_path,
                                                     highlighted, reading, label):
    from extraction.models import Candidate, Issue, ReviewCell
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    issues = []
    if highlighted:
        issues.append(Issue("flagged_field", "source-1", "C2", "review"))
    if reading:
        issues.append(Issue("unreadable_page", "source-1", "source-1#page:3", "unreadable"))
    entry = {"Data cell": "C2", "Record": "source-1", "File": "invoice.pdf", "Field": "ID",
             "Value": "123", "Quote": "ID: 123", "Anchor": "source-1#page:1",
             "Status": "flagged" if highlighted else "verified",
             "Issue": "conflict" if highlighted else ""}
    cell = ReviewCell(2, 3, "123", "s", (entry,))
    candidate = Candidate("job", tmp_path / "candidate.xlsx", (tmp_path / "invoice.pdf",),
                          tuple(issues))
    saved = []
    screen.save_requested.connect(lambda value, ack: saved.append((value, ack)))
    screen.show_candidate(candidate, (cell,))
    assert screen._ack_checkbox.isVisible() == bool(issues)
    if issues:
        assert screen._ack_checkbox.text() == label
        screen._save_btn.click()
        assert saved == []
        screen._ack_checkbox.setChecked(True)
    screen._save_btn.click()
    assert saved == [(candidate, bool(issues))]
    screen.show_candidate(candidate, (cell,))
    assert not screen._ack_checkbox.isChecked()
    assert screen._save_btn.isEnabled() == (not issues)


def test_failure_lists_plain_coverage(qtbot, isolated_profile):
    from extraction.models import Issue
    screen = _screen(qtbot, isolated_profile)
    screen.show_failure("Could not read every page.",
                        (Issue("unreadable_page", "source-2", "source-2#page:3", "x"),))
    assert "page 3: this page has too little text to read." in screen._coverage_view.toPlainText()

@pytest.fixture
def panel(qtbot, isolated_profile, tmp_path):
    widget = _screen(qtbot, isolated_profile)
    widget.set_route("pull")
    path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("ID 00123; date 01/02/2026")
    document.save(path)
    widget.set_paths((path,))
    widget._prompt_edit.setPlainText("Extract the ID")
    widget.show_inspection(read_sources((path,)))
    return widget



def saved_candidate(tmp_path, *, flagged=True, coverage=False):
    issue = Issue("request_failed", "source-1", "source-1#p:1", "A section failed.")
    field = FieldResult(FieldProposal("000452", "ID 000452", "source-1#p:0"),
                        "000452", flagged, ("leading_zero",) if flagged else ())
    result = ExtractionResult(("source-1",), {("source-1", "id"): field},
                              {"source-1": tmp_path / "source.docx"},
                              (issue,) if coverage else (), {})
    return write_candidate(result, (ColumnSpec("id", "ID", "Identifier"),),
                           tmp_path / "candidate.xlsx")


@pytest.mark.parametrize("value,kind,expected", [
    ("12500.00", "decimal", "12500.00"),
    ("12.5000", "decimal", "12.5000"),
    ("-2.50", "decimal", "-2.50"),
    ("123.100000000000000", "decimal", "123.100000000000000"),
    ("9999999999999.9000", "decimal", "9999999999999.9000"),
    ("0.00000000000000001", "decimal", "1E-17"),
    ("12500", "integer", "12500"),
    ("000452", "text", "000452"),
])
def test_review_uses_saved_number_format_and_file_without_record(panel, tmp_path,
                                                                value, kind, expected):
    from decimal import Decimal
    typed = Decimal(value) if kind == "decimal" else int(value) if kind == "integer" else value
    field = FieldResult(FieldProposal(value, value, "source-1#p:0"), typed, False, ())
    result = ExtractionResult(("source-1",), {("source-1", "value"): field},
                              {"source-1": tmp_path / "source.docx"}, (), {})
    candidate = write_candidate(result, (ColumnSpec("value", "Value", "Value", kind),),
                                tmp_path / "candidate.xlsx")
    panel.show_candidate(candidate, read_candidate(candidate))
    assert panel._grid.columnCount() == 2
    assert panel._grid.horizontalHeaderItem(0).text() == "File"
    assert panel._grid.item(0, 0).text() == "source.docx"
    assert panel._grid.item(0, 1).text() == expected
    panel._grid.setCurrentCell(0, 1)
    assert "Found in: source.docx" in panel._evidence.toPlainText()
    book = load_workbook(candidate.path)
    assert book["Evidence"]["B1"].value == "Record"
    book.close()


def test_source_picker_preserves_duplicate_attachments_without_copying(panel, monkeypatch):
    path = panel._paths[0]
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *a, **k:
                        ([str(path), str(path)], ""))
    panel._browse_btn.click()
    assert panel._paths == (path, path, path)
    assert panel._source_list.count() == 3
    assert all(p == path for p in panel._paths)
    assert not panel._confirmed


def test_source_filename_is_drawn_once_and_remains_accessible(panel):
    from PySide6.QtWidgets import QLabel

    item = panel._source_list.item(0)
    row = panel._source_list.itemWidget(item)
    texts = [item.text()] + [label.text() for label in row.findChildren(QLabel)]
    assert texts.count(panel._paths[0].name) == 1
    assert item.data(Qt.ItemDataRole.AccessibleTextRole) == panel._paths[0].name


def test_columns_add_edit_remove_and_freeze_before_extract(panel):
    panel.set_columns((ColumnSpec("id", "ID", "Identifier", "unknown"),))
    assert panel.current_columns()[0].kind == "text"
    assert not panel.columns_confirmed()
    panel._add_column_btn.click()
    assert len(panel.current_columns()) == 2
    panel._columns_table.selectRow(1)
    panel._remove_column_btn.click()
    panel._columns_table.item(0, 0).setText("Customer ID")
    panel._extract_btn.click()
    assert panel._extract_btn.isEnabled()
    with pytest.raises(Exception):  # dataclass is frozen
        panel.snapshot().request = "changed"
    with pytest.raises(ValueError):
        panel.set_columns((ColumnSpec("id", "", ""),))
        panel.current_columns()


def test_numeric_defaults_overrides_and_conditional_date_order(panel):
    panel.set_columns((ColumnSpec("amount", "Amount", "Amount", "decimal"),
                       ColumnSpec("date", "Date", "Date", "date")))
    panel._columns_table.selectRow(0)
    assert panel._thousands_edit.text() == "," and panel._decimal_edit.text() == "."
    panel._thousands_edit.setText(".")
    panel._decimal_edit.setText(",")
    assert panel.current_columns()[0].decimal_separator == ","
    panel._columns_table.selectRow(1)
    assert panel._date_order_combo.isVisible()
    panel._extract_btn.click()
    assert not panel._confirmed and not panel._extract_btn.isEnabled()
    panel._date_order_combo.setCurrentIndex(1)
    panel._extract_btn.click()
    assert panel._confirmed and panel.current_columns()[1].date_order == "DMY"
    documents = tuple(replace(d, units=tuple(replace(u, text="15 March 2026") for u in d.units))
                      for d in panel._documents)
    panel.show_inspection(documents)
    assert not panel._date_order_combo.isVisible()


@pytest.mark.parametrize("thousands,decimal", [("1", "."), ("-", "."), ("+", "."),
                                              ("e", "."), ("$", "."), (",", "1"),
                                              (",", "+"), (",", " ")])
def test_invalid_numeric_separators_cannot_confirm_or_start_extraction(panel, thousands, decimal):
    panel.set_columns((ColumnSpec("amount", "Amount", "Amount", "decimal"),))
    panel._thousands_edit.setText(thousands)
    panel._decimal_edit.setText(decimal)
    panel._extract_btn.click()
    assert not panel._confirmed and not panel._extract_btn.isEnabled()
    with pytest.raises(ValueError, match="separator"):
        panel.snapshot()


@pytest.mark.parametrize("thousands,decimal,value", [(",", ".", "1,234.50"),
                                                    (".", ",", "1.234,50"),
                                                    (" ", ",", "1 234,50"),
                                                    ("'", ".", "1'234.50"),
                                                    ("", ".", "1234.50")])
def test_supported_numeric_separator_overrides_preserve_values(panel, thousands, decimal, value):
    from decimal import Decimal
    from extraction.grounding import validate_field
    from extraction.models import SourceUnit
    panel.set_columns((ColumnSpec("amount", "Amount", "Amount", "decimal"),))
    panel._thousands_edit.setText(thousands)
    panel._decimal_edit.setText(decimal)
    panel._extract_btn.click()
    assert panel._confirmed
    unit = SourceUnit("one#p:0", value, "paragraph", 0)
    result = validate_field(FieldProposal(value, value, unit.anchor), panel.current_columns()[0],
                            {unit.anchor: unit})
    assert not result.flagged and result.typed_value == Decimal("1234.50")


def test_many_coverage_issues_stay_scrollable_and_keep_controls_accessible(panel, qtbot):
    issues = tuple(Issue("low_text", "source-1", f"source-1#page:{page}",
                         "A PDF page has little readable text.") for page in range(2, 102))
    documents = (replace(panel._documents[0], issues=issues, page_cost=101),)
    panel.resize(1000, 650)
    panel.show_inspection(documents)
    qtbot.wait(20)
    assert panel.minimumSizeHint().height() <= 650 and panel.height() <= 650
    assert panel._coverage_view.height() <= 100
    assert panel._source_coverage.verticalScrollBar().maximum() > 0
    assert "page 2" in panel._source_coverage.toPlainText()
    assert "page 101" in panel._source_coverage.toPlainText()
    panel.show_failure("Some pages were not processed.", issues)
    assert panel._coverage_view.height() <= 100
    assert "page 101" in panel._source_coverage.toPlainText()


def test_failure_retains_editable_setup_and_disconnection_gates_model_calls(panel):
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel.set_busy(True)
    panel.show_failure("Could not suggest columns.", ())
    panel.set_busy(False)
    assert panel._prompt_edit.toPlainText() == "Extract the ID"
    assert panel.current_columns()[0].label == "ID"
    panel._extract_btn.click()
    panel.set_ollama_connected(False)
    assert not panel._extract_btn.isEnabled() and not panel._primary_btn.isEnabled()
    assert panel._add_column_btn.isEnabled()
    panel.set_ollama_connected(True)
    assert panel._extract_btn.isEnabled()


def test_saved_grid_literal_flags_evidence_acknowledgement_and_rerun_reset(panel, tmp_path):
    candidate = saved_candidate(tmp_path, coverage=True)
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel._extract_btn.click()
    panel.show_candidate(candidate, read_candidate(candidate))
    cell = panel._grid.item(0, 1)
    assert cell.text() == "000452" and cell.background().color().isValid()
    assert panel._grid.editTriggers() == QTableWidget.EditTrigger.NoEditTriggers
    panel._grid.setCurrentCell(0, 1)
    assert "ID 000452" in panel._evidence.toPlainText()
    assert "Starts with 0" in panel._evidence.toPlainText()
    assert "part of this file was not read." in panel._coverage_view.toPlainText()
    assert not panel._save_btn.isEnabled()
    panel._ack_checkbox.setChecked(True)
    assert panel._save_btn.isEnabled()
    panel.clear_review()
    assert panel._candidate is None and panel._grid.rowCount() == 0
    assert not panel._ack_checkbox.isChecked()


@pytest.mark.parametrize("change", ["request", "column", "sources"])
def test_setup_changes_clear_confirmation_and_review(panel, tmp_path, change):
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel._extract_btn.click()
    candidate = saved_candidate(tmp_path)
    panel.show_candidate(candidate, read_candidate(candidate))
    if change == "request":
        panel._prompt_edit.setPlainText("Another request")
    elif change == "column":
        panel._columns_table.item(0, 0).setText("New ID")
    else:
        panel.set_paths(panel._paths + panel._paths)
    assert not panel._confirmed and panel._candidate is None


def test_busy_setup_is_frozen_and_save_can_finish_without_cancel(panel):
    before = panel._paths
    panel.set_busy(True)
    panel.set_paths((Path("different.docx"),))
    assert panel._paths == before and not panel._primary_btn.isEnabled()
    assert panel._cancel_btn.isEnabled()
    panel.set_busy(True, saving=True)
    assert not panel._cancel_btn.isEnabled()


def test_source_drop_accepts_only_local_supported_files(panel):
    from PySide6.QtCore import QMimeData, QPoint, QPointF, QUrl
    from PySide6.QtGui import QDragEnterEvent, QDropEvent
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(panel._paths[0]))])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.DropAction.CopyAction, mime,
                            Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dragEnterEvent(enter)
    assert enter.isAccepted()
    drop = QDropEvent(QPointF(5, 5), Qt.DropAction.CopyAction, mime,
                      Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dropEvent(drop)
    assert panel._source_list.count() == 2
    mime.setUrls([QUrl("https://example.com/private.pdf")])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.DropAction.CopyAction, mime,
                            Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dragEnterEvent(enter)
    assert not enter.isAccepted()


@pytest.mark.parametrize("page_name", ["_saved_page", "_columns_page", "_review_page"])
def test_source_drop_is_rejected_outside_describe(panel, page_name):
    from PySide6.QtCore import QMimeData, QPoint, QPointF, QUrl
    from PySide6.QtGui import QDragEnterEvent, QDropEvent

    original = panel.snapshot(False).paths
    panel._pages.setCurrentWidget(getattr(panel, page_name))
    requests = []
    panel.inspect_requested.connect(requests.append)
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(original[0]))])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.DropAction.CopyAction, mime,
                           Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dragEnterEvent(enter)
    drop = QDropEvent(QPointF(5, 5), Qt.DropAction.CopyAction, mime,
                      Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dropEvent(drop)
    assert not enter.isAccepted() and not drop.isAccepted()
    assert panel.snapshot(False).paths == original and requests == []


def test_source_remove_buttons_fit_with_application_list_padding(panel, qapp):
    from PySide6.QtWidgets import QPushButton

    original_style = qapp.styleSheet()
    try:
        qapp.setStyleSheet("QListWidget::item { padding: 6px 8px; }")
        panel.set_paths(panel.snapshot(False).paths)
        qapp.processEvents()
        row = panel._source_list.itemWidget(panel._source_list.item(0))
        button = row.findChild(QPushButton)
        assert button.height() >= 28
        button.click()
        assert not panel.snapshot(False).paths
    finally:
        qapp.setStyleSheet(original_style)



def test_failed_column_proposal_offers_manual_columns(qtbot, isolated_profile, tmp_path):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "invoice.docx"
    screen._prompt_edit.setPlainText("Invoice number")
    screen.set_paths((source,))
    screen.show_inspection((_document(source),))
    screen.show_failure("Could not suggest columns.", ())
    assert screen._pages.currentWidget() is screen._columns_page
    screen._add_column_btn.click()
    screen._columns_table.item(0, 0).setText("Invoice number")
    screen._columns_table.item(0, 1).setText("Identifier of the invoice")
    with qtbot.waitSignal(screen.extract_requested):
        screen._extract_btn.click()


def test_absent_saved_value_is_not_reported_as_found(qtbot, isolated_profile, tmp_path):
    from extraction.models import Candidate, ReviewCell
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    entry = {"Data cell": "C2", "Record": "source-1", "File": "invoice.docx", "Field": "PO",
             "Value": None, "Quote": "", "Anchor": "", "Status": "verified", "Issue": ""}
    candidate = Candidate("job", tmp_path / "candidate.xlsx", (), ())
    screen.show_candidate(candidate, (ReviewCell(2, 3, None, "n", (entry,)),))
    assert "Value: Not found" in screen._evidence.toPlainText()
    assert "The source text did not contain a value for this column." in screen._evidence.toPlainText()
    assert "The value was found" not in screen._evidence.toPlainText()
    assert screen._grid.item(0, 1).text() == ""


def test_prompt_hint_follows_sources_then_file_type(qtbot, isolated_profile, tmp_path):
    from app.config.defaults import EXTRACTION_PROMPT_PLACEHOLDER, PROMPT_PLACEHOLDERS
    screen = _screen(qtbot, isolated_profile)
    screen.set_file_type("Excel (.xlsx)")
    assert screen._prompt_edit.placeholderText() == PROMPT_PLACEHOLDERS["Excel (.xlsx)"]
    screen.set_route("pull")
    screen.set_paths((tmp_path / "invoice.docx",))
    assert screen._prompt_edit.placeholderText() == EXTRACTION_PROMPT_PLACEHOLDER
    screen.set_paths(())
    assert screen._prompt_edit.placeholderText() == EXTRACTION_PROMPT_PLACEHOLDER
    screen.set_route("write")
    assert screen._prompt_edit.placeholderText() == PROMPT_PLACEHOLDERS["Excel (.xlsx)"]
    screen.set_file_type("Word (.docx)")
    assert screen._prompt_edit.placeholderText() == PROMPT_PLACEHOLDERS["Word (.docx)"]


def test_reading_problems_start_with_a_truthful_heading(qtbot, isolated_profile, tmp_path):
    from app.widgets.create_screen import COVERAGE_HEADER
    assert COVERAGE_HEADER == (
        "Some information may be missing from the table. Check these problems before you use it:")
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    source = tmp_path / "boxes.docx"
    document = Document()
    document.add_paragraph("Invoice 100")
    document.save(source)
    documents = read_sources((source,))
    source_id = documents[0].source_id
    skipped = Issue("skipped_text_box", source_id, source_id + "#",
                    "Text in text boxes was not read.")
    with_issue = (replace(documents[0], issues=(skipped,)),)
    screen.show_inspection(with_issue)
    assert screen._source_coverage.toPlainText().splitlines() == [
        COVERAGE_HEADER, "boxes.docx: text in text boxes was not read."]
    # Processing problems can keep flagged values, so the heading must not say values are unused.
    screen.show_failure("Could not complete the request.", (
        Issue("request_failed", source_id, source_id + "#p:0", "A section failed."),))
    assert screen._source_coverage.toPlainText().splitlines()[0] == COVERAGE_HEADER
    screen.show_inspection(documents)
    assert screen._source_coverage.isHidden()



def test_browse_sits_directly_under_the_source_box(qtbot, isolated_profile):
    screen = _screen(qtbot, isolated_profile)
    screen.set_route("pull")
    screen.resize(1280, 760)
    screen.show()
    qtbot.waitUntil(screen._browse_btn.isVisible)
    box, browse = screen._source_list.geometry(), screen._browse_btn.geometry()
    assert screen._browse_btn.parentWidget() is screen._source_list.parentWidget()
    assert 0 <= browse.top() - box.bottom() <= 16 and browse.left() == box.left()
    assert screen._browse_btn.text() == "Browse files"
