"""Patch assembly and the six verification checks."""
import io
import zipfile

import pytest
from docx import Document

from editing import patch
from editing.body_map import scan
from editing.models import EditError, Proposal, ProposedParagraph
from editing.package import open_document
from tests.editing.conftest import FIXTURES
from tests.editing.docx_builder import build_docx, p


def proposal(*items: tuple[int, str]) -> Proposal:
    return Proposal(tuple(ProposedParagraph(s, t) for s, t in items))


def index_of(opened, start: str) -> int:
    return next(b.index for b in opened.paragraphs if b.text.startswith(start))


def apply(opened, first, last, *items):
    candidate = patch.build_candidate(opened, first, last, proposal(*items))
    document = Document(io.BytesIO(candidate.data))
    return candidate, document


def xml_of(data: bytes) -> bytes:
    return zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")


def test_replacing_a_bold_value_keeps_it_bold_and_everything_else_identical():
    opened = open_document(FIXTURES / "word-basic.docx")
    k = index_of(opened, "Invoices are payable")
    candidate, document = apply(opened, k, k, (0, "Invoices are payable within 45 days of receipt."))
    runs = [(r.text, bool(r.bold)) for r in document.paragraphs[k].runs]
    assert runs == [("Invoices are payable within ", False), ("45 days", True), (" of receipt.", False)]
    old, new = xml_of(opened.data), xml_of(candidate.data)
    block = opened.paragraphs[k]
    assert new[:block.start] == old[:block.start]
    assert new[len(new) - (len(old) - block.end):] == old[block.end:]
    assert candidate.after_count == 1 and candidate.highlights == (((28, 30),),)


def test_heading_rename_and_split_keep_style_bookmark_tab_and_line_break():
    opened = open_document(FIXTURES / "word-basic.docx")
    h = index_of(opened, "1. Scope")
    candidate, document = apply(
        opened, h, h + 2,
        (0, "1. Project scope"),
        (1, "Review the documents and deliver a concise findings report."),
        (1, "The review covers the agreed documents only."),
        (2, "Deliver the draft within two weeks.\tKeep feedback specific.\nSecond line kept."))
    texts = [x.text for x in document.paragraphs][h:h + 4]
    assert texts[0] == "1. Project scope" and document.paragraphs[h].style.name == "Heading 2"
    assert texts[3] == "Deliver the draft within two weeks.\tKeep feedback specific.\nSecond line kept."
    assert len(document.paragraphs) == len(opened.paragraphs) + 1
    new_xml = xml_of(candidate.data)
    start = new_xml.index(b'w:name="ScopeHeading"')
    assert new_xml.index(b"1. Project scope") > start


def test_shrinking_a_range_and_blank_output_paragraph():
    opened = open_document(FIXTURES / "word-basic.docx")
    k = index_of(opened, "The project includes")
    _, document = apply(opened, k, k + 1, (0, "One merged paragraph."), (0, ""))
    assert [x.text for x in document.paragraphs][k:k + 2] == ["One merged paragraph.", ""]


def test_list_items_continue_the_list():
    opened = open_document(FIXTURES / "word-basic.docx")
    k = index_of(opened, "A short findings report")
    _, document = apply(opened, k, k, (0, "A concise findings report."), (0, "A list of next steps."))
    for paragraph in document.paragraphs[k:k + 2]:
        assert paragraph._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")


def test_goback_paragraph_and_page_break_split():
    opened = open_document(FIXTURES / "word-ranges.docx")
    g = index_of(opened, "The GoBack")
    candidate, _ = apply(opened, g, g, (0, "The hidden bookmark now sits in a rewritten sentence."))
    assert b'w:name="_GoBack"' in xml_of(candidate.data)
    b = index_of(opened, "This paragraph starts on a new page")
    _, document = apply(opened, b, b, (0, "This paragraph starts on a new page."),
                        (0, "This added paragraph must not."))
    assert document.paragraphs[b].paragraph_format.page_break_before is True
    assert document.paragraphs[b + 1].paragraph_format.page_break_before is None


def test_bookmarks_stay_on_the_continuation_paragraph():
    opened = open_document(FIXTURES / "word-ranges.docx")
    h = index_of(opened, "First heading")
    candidate, document = apply(opened, h, h + 1, (0, "Opening heading"), (1, "Closing heading"),
                                (1, "A new body line based on the heading style."))
    xml = xml_of(candidate.data)
    blocks, _ = scan(xml)
    closing = [x for x in blocks if x.kind == "p"][h + 1]
    added = [x for x in blocks if x.kind == "p"][h + 2]
    assert [m.place for m in closing.markers if m.name.startswith("_Toc")] == ["lead"]
    assert any(m.tag == "bookmarkEnd" for m in closing.markers)
    assert added.markers == ()


def test_prefixes_and_self_closing_paragraphs_build_well_formed_parts(tmp_path):
    base = build_docx(p("Alpha") + "<w:p/>")
    xml = xml_of(base)
    variants = {
        "alternate": (xml.replace(b"<w:", b"<x:").replace(b"</w:", b"</x:").replace(b" w:", b" x:")
                      .replace(b"xmlns:w=", b"xmlns:x=")),
        "default": (xml.replace(b"<w:", b"<").replace(b"</w:", b"</").replace(
            b"xmlns:w=", b'xmlns="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w=')),
        "original": xml,
    }
    for name, content in variants.items():
        path = tmp_path / f"{name}.docx"
        path.write_bytes(build_docx(document=content))
        opened = open_document(path)
        for first in (0, 1):
            _, document = apply(opened, first, first, (0, "Replaced text."), (0, "Split text."))
            assert [x.text for x in document.paragraphs][first:first + 2] == ["Replaced text.", "Split text."]


def test_gap_markers_between_selected_paragraphs_are_kept_in_order(make_docx):
    path = make_docx(p("One") + '<w:bookmarkStart w:id="5" w:name="Gap"/>' + p("Two")
                     + '<w:bookmarkEnd w:id="5"/>' + p("Three"))
    opened = open_document(path)
    candidate, _ = apply(opened, 0, 1, (0, "First"), (1, "Second"))
    xml = xml_of(candidate.data)
    assert xml.index(b"First") < xml.index(b'w:name="Gap"') < xml.index(b"Second") < xml.index(b'w:id="5"/>')


def test_empty_original_paragraph_takes_the_paragraph_style(make_docx):
    opened = open_document(make_docx('<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr></w:p>'))
    _, document = apply(opened, 0, 0, (0, "Filled heading"))
    assert document.paragraphs[0].text == "Filled heading"
    assert document.paragraphs[0].runs[0]._r.rPr is None


def test_text_is_escaped_and_new_paragraphs_get_no_paragraph_id():
    opened = open_document(FIXTURES / "word-basic.docx")
    k = index_of(opened, "Use the agreed")
    candidate, document = apply(opened, k, k, (0, "A < B & C > D"), (0, "Second"))
    assert document.paragraphs[k].text == "A < B & C > D"
    assert xml_of(candidate.data).count(opened.paragraphs[k].start_tag) == 1


@pytest.mark.parametrize("corrupt", [
    lambda xml, s, e: xml[:s - 1] + b"X" + xml[s:],
    lambda xml, s, e: xml[:s] + b"<w:p>" + xml[s:],
    lambda xml, s, e: xml[:s] + xml[s:e].replace(b"Invoices", b"Bills") + xml[e:],
])
def test_verification_rejects_corrupted_candidates(monkeypatch, corrupt):
    opened = open_document(FIXTURES / "word-basic.docx")
    k = index_of(opened, "Invoices are payable")
    real = patch._splice

    def broken(*args):
        xml, start, end, highlights = real(*args)
        return corrupt(xml, start, end), start, end, highlights

    monkeypatch.setattr(patch, "_splice", broken)
    with pytest.raises(EditError, match="could not apply this change safely"):
        patch.build_candidate(opened, k, k, proposal((0, "Invoices are payable within 45 days.")))


def test_verification_rejects_a_changed_other_part_and_a_moved_bookmark(monkeypatch):
    opened = open_document(FIXTURES / "word-ranges.docx")
    h = index_of(opened, "First heading")
    real_rebuild = patch.rebuild
    monkeypatch.setattr(patch, "rebuild", lambda data, part, content: real_rebuild(
        real_rebuild(data, part, content), "word/styles.xml", b"<w:styles/>"))
    with pytest.raises(EditError):
        patch.build_candidate(opened, h, h, proposal((0, "Renamed")))
    monkeypatch.setattr(patch, "rebuild", real_rebuild)
    real = patch._splice

    def moved(*args):
        xml, start, end, highlights = real(*args)
        tag = b'w:name="_Toc'
        at = xml.index(tag, start)
        begin = xml.rindex(b"<", 0, at)
        stop = xml.index(b"/>", at) + 2
        marker = xml[begin:stop]
        rest = xml[:begin] + xml[stop:]
        close = rest.index(b"</w:p>", begin)
        return rest[:close] + marker + rest[close:], start, end, highlights

    monkeypatch.setattr(patch, "_splice", moved)
    with pytest.raises(EditError):
        patch.build_candidate(opened, h, h, proposal((0, "Renamed")))


def test_verification_reads_the_packaged_main_part(monkeypatch):
    opened = open_document(FIXTURES / "word-basic.docx")
    k = index_of(opened, "Invoices are payable")
    real = patch.rebuild
    monkeypatch.setattr(patch, "rebuild", lambda data, part, content: real(
        data, part, content.replace(b"45 days", b"99 days")))
    with pytest.raises(EditError, match="could not apply this change safely"):
        patch.build_candidate(opened, k, k, proposal((0, "Invoices are payable within 45 days of receipt.")))


TIME_RUN = '<w:r><w:rPr><w:sz w:val="21"/></w:rPr><w:t xml:space="preserve">12:15 to 13:00 </w:t></w:r>'
TITLE_RUN = '<w:r><w:rPr><w:i/></w:rPr><w:t>The Future of the Software Engineering Profession</w:t></w:r>'


def test_a_rewritten_paragraph_takes_one_format_from_most_of_the_original(make_docx):
    # Alex's devTalks case: a mostly new sentence was split between two original formats.
    opened = open_document(make_docx("<w:p>" + TIME_RUN + TITLE_RUN + "</w:p>"))
    candidate, document = apply(opened, 0, 0, (0, "Superman is the best."))
    runs = document.paragraphs[0].runs
    assert [r.text for r in runs] == ["Superman is the best."]
    assert runs[0].italic and runs[0].font.size is None
    assert candidate.highlights == (((0, 21),),)


def test_a_small_edit_keeps_word_by_word_formatting(make_docx):
    opened = open_document(make_docx("<w:p>" + TIME_RUN + TITLE_RUN + "</w:p>"))
    _, document = apply(opened, 0, 0, (0, "12:15 to 13:30 The Future of the Software Engineering Profession"))
    runs = [(r.text, r.font.size is not None, bool(r.italic)) for r in document.paragraphs[0].runs]
    assert runs == [("12:15 to 13:30 ", True, False),
                    ("The Future of the Software Engineering Profession", False, True)]
