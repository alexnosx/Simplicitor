"""Body map: byte ranges, story-level ranges, bookmark places and editability."""
import io
import zipfile

import pytest
from docx import Document

from editing.body_map import PREPARE_FAILED, scan, selection_problem
from editing.models import EditError
from tests.editing.conftest import FIXTURES
from tests.editing.docx_builder import build_docx, p


def _xml(data: bytes) -> bytes:
    return zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")


def _reasons(data: bytes) -> list[tuple[str, str | None]]:
    blocks, _ = scan(_xml(data))
    return [(b.text, b.reason) for b in blocks if b.kind == "p"]


def test_word_basic_fixture_matches_python_docx_and_names_read_only_parts():
    data = (FIXTURES / "word-basic.docx").read_bytes()
    blocks, _ = scan(_xml(data))
    paragraphs = [b for b in blocks if b.kind == "p"]
    texts = [x.text for x in Document(io.BytesIO(data)).paragraphs]
    assert len(paragraphs) == len(texts)
    assert all(b.text == texts[b.index] for b in paragraphs if b.reason is None)
    reasons = {b.text.split(":")[0]: b.reason for b in paragraphs}
    assert reasons["Related policy"] == "a link"
    assert next(b.reason for b in paragraphs if texts[b.index].startswith("Reference page")) == "a field"
    assert reasons["Invoices are payable within 30 days of receipt."] is None
    assert [b.kind for b in blocks].count("tbl") == 1


def test_word_ranges_fixture_protects_fields_comments_and_inline_bookmarks():
    data = (FIXTURES / "word-ranges.docx").read_bytes()
    texts = [x.text for x in Document(io.BytesIO(data)).paragraphs]
    blocks, _ = scan(_xml(data))
    reason = {texts[b.index]: b.reason for b in blocks if b.kind == "p"}
    assert reason["First heading"] is None and reason["Second heading"] is None
    assert reason["Comment range starts in this paragraph."] == "a comment"
    assert reason["This middle paragraph is inside the comment range."] == "a comment"
    assert reason["Comment range ends in this paragraph."] == "a comment"
    assert reason["Invoices are payable within 30 days of receipt."] == "a bookmark inside the text"
    assert reason["The GoBack bookmark sits inside this sentence."] is None
    assert reason["This paragraph starts on a new page."] is None
    toc = [b.reason for b in blocks if b.kind == "p" and texts[b.index].endswith("\t1")]
    assert toc and set(toc) == {"a field"}


def test_byte_ranges_cover_each_body_child_exactly():
    data = build_docx(p("One") + "<w:bookmarkStart w:id=\"1\" w:name=\"x\"/>" + p("Two")
                      + "<w:bookmarkEnd w:id=\"1\"/>")
    xml = _xml(data)
    blocks, bookmarks = scan(xml)
    assert [b.kind for b in blocks] == ["p", "marker", "p", "marker", "sectPr"]
    assert xml[blocks[0].start:blocks[0].end].startswith(b"<w:p>")
    assert xml[blocks[0].start:blocks[0].end].endswith(b"</w:p>")
    assert xml[blocks[1].start:blocks[1].end] == b'<w:bookmarkStart w:id="1" w:name="x"/>'
    assert [m.tag for m in bookmarks] == ["bookmarkStart", "bookmarkEnd"]


def test_unusual_but_valid_xml_is_mapped():
    body = ('<w:p w:rsidR="00&gt;01" w:rsidRDefault="a>b"><w:r><w:t>Café &amp; té</w:t></w:r></w:p>'
            "<!-- a comment --><?pi data?>\r\n"
            "<w:p><w:r><w:t><![CDATA[x < y]]></w:t></w:r></w:p><w:p/>")
    data = build_docx(body)
    xml = b"\xef\xbb\xbf" + _xml(data)
    blocks, _ = scan(xml)
    paragraphs = [b for b in blocks if b.kind == "p"]
    assert [b.text for b in paragraphs] == ["Café & té", "x < y", ""]
    assert xml[paragraphs[2].start:paragraphs[2].end] == b"<w:p/>"
    assert all(b.reason is None for b in paragraphs)


def test_alternate_and_default_prefixes_are_recorded():
    data = build_docx(p("Alpha"))
    xml = _xml(data)
    alternate = (xml.replace(b"<w:", b"<x:").replace(b"</w:", b"</x:").replace(b" w:", b" x:")
                 .replace(b"xmlns:w=", b"xmlns:x="))
    blocks, _ = scan(alternate)
    assert blocks[0].prefix == b"x:" and blocks[0].text == "Alpha"
    default = (xml.replace(b"<w:", b"<").replace(b"</w:", b"</")
               .replace(b"xmlns:w=", b'xmlns="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w='))
    blocks, _ = scan(default)
    assert blocks[0].prefix == b"" and blocks[0].text == "Alpha"


@pytest.mark.parametrize("inner, reason", [
    ('<w:hyperlink r:id="x"><w:r><w:t>link</w:t></w:r></w:hyperlink>', "a link"),
    ('<w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>', "a field"),
    ('<w:ins w:id="1" w:author="a"><w:r><w:t>new</w:t></w:r></w:ins>', "tracked changes"),
    ('<w:sdt><w:sdtContent><w:r><w:t>form</w:t></w:r></w:sdtContent></w:sdt>', "a form field"),
    ('<w:r><w:drawing/></w:r>', "a picture"),
    ('<w:r><w:footnoteReference w:id="1"/></w:r>', "a footnote"),
    ('<w:r><w:br w:type="page"/></w:r>', "a page break"),
    ('<w:r><w:sym w:font="Symbol" w:char="F020"/></w:r>', "content this version cannot edit"),
    ('<w:smartTag w:element="x"><w:r><w:t>tag</w:t></w:r></w:smartTag>',
     "content this version cannot edit"),
    ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"/>',
     "content this version cannot edit"),
])
def test_paragraph_contents_that_make_it_read_only(inner, reason):
    data = build_docx(f"<w:p><w:r><w:t>Text </w:t></w:r>{inner}</w:p>")
    assert _reasons(data)[0][1] == reason


def test_section_break_and_revision_properties_make_it_read_only():
    data = build_docx(p("A", ppr="<w:sectPr/>") + p("B", rpr='<w:rPrChange w:id="1" w:author="a"/>')
                      + '<w:p><w:pPr><w:rPr><w:ins w:id="2" w:author="a"/></w:rPr></w:pPr>'
                        "<w:r><w:t>C</w:t></w:r></w:p>")
    assert [r for _, r in _reasons(data)] == ["a section break", "tracked changes", "tracked changes"]


def test_complex_field_across_paragraphs_protects_the_plain_middle_paragraph():
    body = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText>TOC</w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>Entry one</w:t></w:r></w:p>'
            + p("Entry two")
            + '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>' + p("After the field"))
    assert [r for _, r in _reasons(build_docx(body))] == ["a field", "a field", "a field", None]


def test_comment_range_across_paragraphs_protects_the_plain_middle_paragraph():
    body = ('<w:p><w:commentRangeStart w:id="4"/><w:r><w:t>Start</w:t></w:r></w:p>' + p("Middle")
            + '<w:p><w:r><w:t>End</w:t></w:r><w:commentRangeEnd w:id="4"/>'
              '<w:r><w:commentReference w:id="4"/></w:r></w:p>' + p("After"))
    assert [r for _, r in _reasons(build_docx(body))] == ["a comment", "a comment", "a comment", None]


@pytest.mark.parametrize("marker", ['<w:commentRangeStart w:id="7"/>', '<w:commentRangeEnd w:id="7"/>'])
def test_point_comments_open_no_range(marker):
    body = (f'<w:p><w:r><w:t>Anchor</w:t></w:r>{marker}<w:r><w:commentReference w:id="7"/></w:r></w:p>'
            + p("Unrelated paragraph"))
    assert [r for _, r in _reasons(build_docx(body))] == ["a comment", None]


@pytest.mark.parametrize("body", [
    '<w:p><w:commentRangeEnd w:id="1"/></w:p><w:p><w:commentRangeStart w:id="1"/></w:p>',
    '<w:p><w:commentRangeStart w:id="1"/><w:commentRangeStart w:id="1"/></w:p>',
    '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>',
    '<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r></w:p>',
])
def test_reversed_repeated_or_unbalanced_ranges_refuse_the_document(body):
    with pytest.raises(EditError, match=PREPARE_FAILED):
        scan(_xml(build_docx(body)))


def test_bookmark_places_and_the_goback_exception():
    body = ('<w:p><w:bookmarkStart w:id="1" w:name="Head"/><w:r><w:t>Heading</w:t></w:r>'
            '<w:bookmarkEnd w:id="1"/></w:p>'
            '<w:p><w:r><w:t>The </w:t></w:r><w:bookmarkStart w:id="2" w:name="_GoBack"/>'
            '<w:bookmarkEnd w:id="2"/><w:r><w:t>cursor</w:t></w:r></w:p>'
            '<w:p><w:r><w:t>Pay in </w:t></w:r><w:bookmarkStart w:id="3" w:name="Term"/>'
            '<w:r><w:t>30 days</w:t></w:r><w:bookmarkEnd w:id="3"/></w:p>')
    blocks, _ = scan(_xml(build_docx(body)))
    paragraphs = [b for b in blocks if b.kind == "p"]
    assert [m.place for m in paragraphs[0].markers] == ["lead", "trail"]
    assert [m.place for m in paragraphs[1].markers] == ["lead", "lead"]
    assert paragraphs[1].reason is None
    assert paragraphs[2].reason == "a bookmark inside the text"


def test_namespace_declaration_inside_a_paragraph_makes_it_read_only():
    data = build_docx('<w:p xmlns:v="urn:x"><w:r><w:t>A</w:t></w:r></w:p>')
    assert _reasons(data)[0][1] == "content this version cannot edit"


def test_selection_rules():
    body = (p("One") + p("Two") + '<w:bookmarkStart w:id="1" w:name="b"/>' + p("Three")
            + '<w:bookmarkEnd w:id="1"/><w:tbl><w:tr><w:tc><w:p/></w:tc></w:tr></w:tbl>' + p("Four")
            + '<w:p><w:hyperlink><w:r><w:t>Five</w:t></w:r></w:hyperlink></w:p>' + p("Six"))
    blocks, _ = scan(_xml(build_docx(body)))
    assert selection_problem(blocks, 0, 2) is None
    assert selection_problem(blocks, 2, 3).startswith("A selection cannot cross a table")
    assert selection_problem(blocks, 3, 5).startswith("This range includes a read-only paragraph")
    assert selection_problem(blocks, 4, 4).startswith("This range includes a read-only paragraph")
    assert selection_problem(blocks, 5, 9) == "Select one or more paragraphs."
