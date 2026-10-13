"""Opening refusals, settings flags and the one-part rebuild."""
import io
import zipfile

import pytest

from editing.models import EditError
from editing.package import open_document, rebuild
from tests.editing.conftest import FIXTURES
from tests.editing.docx_builder import build_docx, p

STANDARD = "Simplicitor can edit standard Word documents (.docx) only."


def _rewrite(data: bytes, name: str, old: bytes, new: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            content = source.read(info)
            target.writestr(info, content.replace(old, new) if info.filename == name else content)
    return out.getvalue()


def test_opens_word_fixtures_without_changing_them():
    for name in ("word-basic.docx", "word-ranges.docx"):
        before = (FIXTURES / name).read_bytes()
        opened = open_document(FIXTURES / name)
        assert opened.part_name == "word/document.xml"
        assert opened.data == before and (FIXTURES / name).read_bytes() == before
        assert not opened.track_changes_on


@pytest.mark.parametrize("name, data, message", [
    ("notes.doc", b"anything", "Choose a Word document"),
    ("locked.docx", b"\xd0\xcf\x11\xe0" + b"\0" * 60, "Choose a Word document"),
    ("broken.docx", b"not a zip", "This Word file could not be read."),
])
def test_refuses_files_that_are_not_readable_docx(tmp_path, name, data, message):
    path = tmp_path / name
    path.write_bytes(data)
    with pytest.raises(EditError, match=message):
        open_document(path)


def test_refuses_macro_documents_strict_packages_and_signed_files(tmp_path):
    base = build_docx(p("Text"))
    cases = {
        "macro.docx": _rewrite(base, "[Content_Types].xml", b"document.main+xml",
                               b"document.macroEnabled.main+xml"),
        "strict.docx": _rewrite(base, "_rels/.rels",
                                b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
                                b"http://purl.oclc.org/ooxml/officeDocument/relationships/officeDocument"),
        "strict-root.docx": _rewrite(base, "word/document.xml",
                                     b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"',
                                     b'xmlns:w="http://purl.oclc.org/ooxml/wordprocessingml/main"'),
    }
    for name, data in cases.items():
        (tmp_path / name).write_bytes(data)
        with pytest.raises(EditError, match="standard Word documents"):
            open_document(tmp_path / name)
    signed = build_docx(p("Text"), extra={"_xmlsignatures/sig1.xml": b"<Signature/>"})
    (tmp_path / "signed.docx").write_bytes(signed)
    with pytest.raises(EditError, match="digitally signed"):
        open_document(tmp_path / "signed.docx")


@pytest.mark.parametrize("setting, refused", [
    ('<w:documentProtection w:edit="readOnly" w:enforcement="1"/>', True),
    ('<w:documentProtection w:edit="readOnly"/>', True),
    ('<w:documentProtection w:edit="readOnly" w:enforcement="0"/>', False),
    ('<w:documentProtection w:edit="readOnly" w:enforcement="false"/>', False),
])
def test_protection_uses_on_off_values(make_docx, setting, refused):
    path = make_docx(p("Text"), settings=setting)
    if refused:
        with pytest.raises(EditError, match="protected in Word"):
            open_document(path)
    else:
        assert open_document(path).paragraphs[0].text == "Text"


@pytest.mark.parametrize("setting, on", [
    ("", False), ("<w:trackRevisions/>", True), ('<w:trackRevisions w:val="true"/>', True),
    ('<w:trackRevisions w:val="0"/>', False), ('<w:trackRevisions w:val="off"/>', False),
])
def test_track_changes_flag_uses_on_off_values(make_docx, setting, on):
    assert open_document(make_docx(p("Text"), settings=setting)).track_changes_on is on


def test_refuses_over_300_page_equivalents(make_docx):
    body = "".join(p("x" * 3000) for _ in range(301))
    with pytest.raises(EditError, match="over 300 pages"):
        open_document(make_docx(body))


def test_refuses_unbalanced_ranges_and_non_utf8(make_docx, tmp_path):
    with pytest.raises(EditError, match="could not be prepared"):
        open_document(make_docx('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r></w:p>'))
    data = _rewrite(build_docx(p("Text")), "word/document.xml", b"encoding='UTF-8'", b"encoding='UTF-16'")
    data = _rewrite(data, "word/document.xml", b'encoding="UTF-8"', b'encoding="UTF-16"')
    (tmp_path / "utf16.docx").write_bytes(data)
    with pytest.raises(EditError, match="could not be read"):
        open_document(tmp_path / "utf16.docx")


def test_rebuild_replaces_only_the_named_part():
    data = (FIXTURES / "word-basic.docx").read_bytes()
    out = rebuild(data, "word/document.xml", b"<new/>")
    old, new = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(out))
    assert [i.filename for i in old.infolist()] == [i.filename for i in new.infolist()]
    for info in old.infolist():
        expected = b"<new/>" if info.filename == "word/document.xml" else old.read(info)
        assert new.read(info.filename) == expected
        assert new.getinfo(info.filename).compress_type == info.compress_type
        assert new.getinfo(info.filename).date_time == info.date_time


def _first_page_header(document, text):
    section = document.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = text


@pytest.mark.parametrize("body, header, merged, refused", [
    (900_000, 0, 0, False),
    (900_001, 0, 0, True),
    (899_500, 1_000, 0, True),
    (0, 0, 451_000, False),
])
def test_page_limit_counts_exactly_like_extraction(tmp_path, body, header, merged, refused):
    from docx import Document
    from extraction.source_readers import _read_docx
    document = Document()
    if body:
        document.add_paragraph("x" * body)
    if header:
        _first_page_header(document, "h" * header)
    if merged:
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "m" * merged
    path = tmp_path / "limit.docx"
    document.save(path)
    from editing.package import _page_cost
    assert _page_cost(Document(path)) == _read_docx(path, "x").page_cost
    if refused:
        with pytest.raises(EditError, match="over 300 pages"):
            open_document(path)
    else:
        assert open_document(path).part_name == "word/document.xml"
