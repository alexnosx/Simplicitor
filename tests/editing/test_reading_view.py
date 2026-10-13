"""Reading view: effective bold/italic as Word shows it, lists, tables, headers, review blocks."""
import pytest

from editing import reading_view
from editing.models import Proposal, ProposedParagraph
from editing.package import open_document
from editing.patch import build_candidate
from tests.editing.conftest import FIXTURES
from tests.editing.docx_builder import p

STYLES = (
    '<w:style w:type="character" w:customStyle="1" w:styleId="CharBoldA"><w:name w:val="CharBoldA"/>'
    '<w:rPr><w:b/></w:rPr></w:style>'
    '<w:style w:type="character" w:customStyle="1" w:styleId="CharBoldB"><w:name w:val="CharBoldB"/>'
    '<w:basedOn w:val="CharBoldA"/><w:rPr><w:b/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="ParaBold"><w:name w:val="ParaBold"/>'
    '<w:basedOn w:val="Normal"/><w:rPr><w:b/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:customStyle="1" w:styleId="ParaItalic"><w:name w:val="ParaItalic"/>'
    '<w:basedOn w:val="Normal"/><w:rPr><w:i/></w:rPr></w:style>'
)
# Word 16 results for the same cases (Claude's and Codex's checks, 2026-10-13).
CASES = [
    (p("case1", rpr='<w:rStyle w:val="CharBoldB"/>'), True),
    (p("case2", ppr='<w:pStyle w:val="ParaBold"/>', rpr='<w:rStyle w:val="CharBoldA"/>'), False),
    (p("case3", ppr='<w:pStyle w:val="ParaBold"/>'), True),
    (p("case4", ppr='<w:pStyle w:val="ParaBold"/>', rpr='<w:b w:val="0"/>'), False),
    (p("case5", rpr='<w:rStyle w:val="CharBoldA"/><w:b w:val="0"/>'), False),
    (p("case6", rpr='<w:rStyle w:val="Strong"/>'), True),
]


def _blocks(make_docx, body, styles=STYLES):
    return [b for b in reading_view.build(open_document(make_docx(body, styles=styles)))
            if b.kind == "paragraph"]


def test_effective_bold_follows_word(make_docx):
    blocks = _blocks(make_docx, "".join(body for body, _ in CASES))
    assert [b.spans[0].bold for b in blocks] == [bold for _, bold in CASES]


def test_italic_from_paragraph_style_and_direct_override(make_docx):
    blocks = _blocks(make_docx, p("a", ppr='<w:pStyle w:val="ParaItalic"/>')
                     + p("b", ppr='<w:pStyle w:val="ParaItalic"/>', rpr='<w:i w:val="false"/>'))
    assert [b.spans[0].italic for b in blocks] == [True, False]


@pytest.mark.parametrize("numpr, listed", [
    ("", True),
    ('<w:numPr><w:ilvl w:val="0"/></w:numPr>', True),
    ('<w:numPr><w:numId w:val="0"/></w:numPr>', False),
    ("<w:numPr/>", True),
])
def test_list_marker_resolves_num_id_per_component(make_docx, numpr, listed):
    blocks = _blocks(make_docx, p("item", ppr='<w:pStyle w:val="ListBullet"/>' + numpr), styles="")
    assert blocks[0].list_item is listed


def test_word_fixture_view_has_headings_lists_table_and_read_only_reasons():
    view = reading_view.build(open_document(FIXTURES / "word-basic.docx"))
    first = view[0]
    assert first.kind == "paragraph" and first.heading == 1 and first.text == "Service outline"
    intro = view[1]
    assert [(s.text, s.bold, s.italic) for s in intro.spans if s.text in ("Example Services", "Sample Studio")] == [
        ("Example Services", True, False), ("Sample Studio", False, True)]
    assert any(b.list_item for b in view)
    table = next(b for b in view if b.kind == "table")
    assert table.cells[0] == ("Milestone", "Example date")
    link = next(b for b in view if b.text.startswith("Related policy"))
    assert link.reason == "a link" and link.text.endswith("example policy link")


def test_review_blocks_come_from_candidate_bytes():
    opened = open_document(FIXTURES / "word-basic.docx")
    k = next(b.index for b in opened.paragraphs if b.text.startswith("Invoices are payable"))
    candidate = build_candidate(opened, k, k, Proposal((ProposedParagraph(0, "Invoices are payable within 45 days."),)))
    before = reading_view.paragraphs(opened.data, k, 1)[0]
    after = reading_view.paragraphs(candidate.data, k, candidate.after_count)[0]
    assert [s.text for s in before.spans if s.bold] == ["30 days"]
    assert [s.text for s in after.spans if s.bold] == ["45 days"]


def test_nested_and_header_tables_are_shown_read_only(tmp_path):
    from docx import Document
    document = Document()
    document.add_paragraph("Body")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "Outer"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested visible text"
    section = document.sections[0]
    section.header.add_table(rows=1, cols=2, width=section.page_width).rows[0].cells[1].text = "Header table text"
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "First page header"
    path = tmp_path / "tables.docx"
    document.save(path)
    view = reading_view.build(open_document(path))
    table = next(b for b in view if b.kind == "table")
    assert "Nested visible text" in table.cells[0][0] and table.reason
    headers = [b.text for b in view if b.kind == "header"]
    assert any("Header table text" in h for h in headers) and "First page header" in headers
