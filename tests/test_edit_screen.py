"""The approved Word editor flow, with synthetic documents and no model calls."""
from dataclasses import replace
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PySide6.QtCore import QPoint, Qt
from PySide6.QtWidgets import QFileDialog, QLabel, QPushButton

from editing import reading_view
from editing.models import Proposal, ProposedParagraph
from editing.package import open_document
from editing.patch import build_candidate


def make_opened(folder, name="service-outline.docx"):
    """A synthetic document with margins, ordinary paragraphs, a table and a field."""
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Example Services | Synthetic outline"
    document.sections[0].footer.paragraphs[0].text = "Fictional document"
    document.add_heading("Service outline", 1)
    paragraph = document.add_paragraph("Invoices are payable within ")
    paragraph.add_run("30 days").bold = True
    paragraph.add_run(" of receipt.")
    document.add_paragraph("Delivery is due in December.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    paragraph = document.add_paragraph("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    document.add_paragraph("Closing paragraph.")
    path = folder / name
    document.save(path)
    return open_document(path)


def candidate_for(opened):
    candidate = build_candidate(opened, 1, 1, Proposal((
        ProposedParagraph(0, "Invoices are payable within 45 days of receipt."),)))
    before = reading_view.paragraphs(opened.data, 1, 1)
    after = reading_view.paragraphs(candidate.data, 1, candidate.after_count)
    candidate = reading_view.with_review_deletions(candidate, before, after)
    return candidate, before, after


@pytest.fixture
def screen(qtbot, tmp_path):
    from app.widgets.edit_screen import EditScreen
    page = EditScreen()
    qtbot.addWidget(page)
    page.resize(1240, 700)
    page.set_model_ready(True)
    opened = make_opened(tmp_path)
    page.show_opened(opened, reading_view.build(opened))
    return page


def click_paragraph(qtbot, screen, index, shift=False):
    view = screen._reader
    row = next(i for i, b in enumerate(view.model().blocks) if b.index == index)
    model_index = view.model().index(row, 0)
    view.scrollTo(model_index)
    qtbot.waitUntil(lambda: view.visualRect(model_index).height() > 0)
    qtbot.mouseClick(view.viewport(), Qt.MouseButton.LeftButton,
                    Qt.KeyboardModifier.ShiftModifier if shift else Qt.KeyboardModifier.NoModifier,
                    pos=view.visualRect(model_index).center())


def describe(screen, qtbot):
    click_paragraph(qtbot, screen, 1)
    screen._edit_part_btn.click()
    screen._description.setPlainText("Change 30 days to 45 days.")


def review(screen, qtbot):
    describe(screen, qtbot)
    screen.show_proposal(*candidate_for(screen._opened))


def test_open_shows_reading_view_and_read_only_reasons(screen, qtbot):
    blocks = screen._reader.model().blocks
    assert blocks[0].kind == "header" and blocks[-1].kind == "footer"
    assert "You are working on: service-outline.docx" == screen._file_label.text()
    assert any(b.kind == "table" and b.cells == (("Item", "Amount"),) for b in blocks)
    row = next(i for i, b in enumerate(blocks) if b.index == 3)
    screen._reader.activate_row(row)
    assert "cannot be edited" in screen._selection_notice.text()
    assert "Read-only" in screen._reader.itemDelegate().row_label(blocks[row])


def test_edit_wording_matches_the_approved_mockup(screen):
    labels = [label.text() for label in screen.findChildren(QLabel)]
    assert "Edit your Word document" in labels
    assert "Select a part, describe the change, then review it before saving a new file." in labels
    assert "1. Open your Word document" in labels
    assert screen._drop_zone.text() == "Drop a .docx here, or browse for a Word document."
    assert screen._original_note.text() == (
        f"Original in {screen._opened.path.parent}. Your original will not be overwritten.")


def test_refused_file_shows_message_and_keeps_drop_zone(screen):
    screen.reset()
    screen.show_failure("This file cannot be read.", "failed")
    assert screen._pages.currentWidget() is screen._open_page
    assert screen._drop_zone.isVisible() and screen._browse_btn.isEnabled()
    assert "This file cannot be read." in screen._error.text()


def test_shift_click_selection_respects_tables_and_read_only_rows(screen, qtbot):
    click_paragraph(qtbot, screen, 1)
    click_paragraph(qtbot, screen, 2, shift=True)
    assert screen._selection == (1, 2)
    click_paragraph(qtbot, screen, 4, shift=True)
    assert screen._selection == (1, 2)
    assert "cannot cross a table" in screen._selection_notice.text()
    click_paragraph(qtbot, screen, 3)
    assert screen._selection == (1, 2)
    assert "read-only" in screen._selection_notice.text().lower() or "cannot be edited" in screen._selection_notice.text()
    screen._clear_btn.click()
    assert screen._selection is None and not screen.has_work()


def test_long_selection_disables_edit_this_part(screen, qtbot):
    paragraphs = list(screen._opened.blocks)
    paragraphs = tuple(replace(b, text="x" * 6001) if b.index == 1 else b for b in paragraphs)
    screen.show_opened(replace(screen._opened, blocks=paragraphs), screen._view)
    click_paragraph(qtbot, screen, 1)
    assert not screen._edit_part_btn.isEnabled()
    assert "Select a shorter part" in screen._selection_notice.text()


def test_description_limit_and_no_model_warning(screen, qtbot):
    describe(screen, qtbot)
    screen._description.setPlainText("x" * 2001)
    assert len(screen._description.toPlainText()) == 2000
    screen.set_model_ready(False)
    assert screen._no_model.isVisible() and screen._settings_btn.isVisible()
    assert not screen._propose_btn.isEnabled()
    screen.set_model_ready(True)
    assert not screen._no_model.isVisible() and screen._propose_btn.isEnabled()
    screen._back_btn.click()
    assert len(screen._description.toPlainText()) == 2000


def test_review_shows_before_after_highlight_and_track_changes_notice(screen, qtbot):
    screen._opened = replace(screen._opened, track_changes_on=True)
    review(screen, qtbot)
    assert "30 days" in screen._before.model().blocks[0].text
    assert "45 days" in screen._after.model().blocks[0].text
    assert screen._after.model().highlights == screen._candidate.highlights
    assert "saved without revision marks" in screen._track_notice.text()
    formats = screen._after.itemDelegate().text_formats(screen._after.model().blocks[0],
                                                      screen._candidate.highlights[0])
    assert any(f.format.background().color().name() == "#dcfce7" for f in formats)
    assert "Longer text can change page breaks." in screen._review_note.text()


def test_try_again_keeps_description_and_increments_attempt(screen, qtbot):
    review(screen, qtbot)
    text = screen._description.toPlainText()
    screen._try_btn.click()
    assert screen._description.toPlainText() == text and screen._attempt == 1
    with qtbot.waitSignal(screen.propose_requested) as signal:
        screen._propose_btn.click()
    assert signal.args == [1, 1, text, 1]


def test_cancel_review_keeps_selection(screen, qtbot):
    review(screen, qtbot)
    screen._cancel_review_btn.click()
    assert screen._selection == (1, 1) and screen._candidate is None
    assert screen._pages.currentWidget() is screen._select_page


def _show_deletion_review(screen):
    """Keep one selected paragraph after an insertion and drop the other, without a model."""
    opened = screen._opened
    screen.select_paragraph(1)
    screen.select_paragraph(2, True)
    screen._edit_part_btn.click()
    screen._description.setPlainText("Rewrite the selected part.")
    kept = opened.paragraphs[1].text
    candidate = build_candidate(opened, 1, 2, Proposal((
        ProposedParagraph(0, "A new introduction."), ProposedParagraph(0, kept))))
    before = reading_view.paragraphs(opened.data, 1, 2)
    after = reading_view.paragraphs(candidate.data, 1, candidate.after_count)
    candidate = reading_view.with_review_deletions(candidate, before, after)
    screen.show_proposal(candidate, before, after)
    return candidate, before


def test_review_calls_out_a_removed_paragraph_and_keeps_the_proposal_unchanged(screen):
    candidate, before = _show_deletion_review(screen)
    assert screen._deletion_notice.isVisible()
    assert "1 original paragraph" in screen._deletion_notice.text()
    assert "removed or merged" in screen._deletion_notice.text()
    assert screen._before.model().deletions == ((), ((0, 28),))
    assert screen._before.model().blocks == before
    assert screen._after.model().highlights == candidate.highlights
    assert screen._after.model().deletions == ()
    assert screen._candidate is candidate
    screen._use_btn.click()
    assert screen._pages.currentWidget() is screen._save_page


def test_before_marks_removed_words_red_without_losing_bold_or_italic(screen):
    from PySide6.QtGui import QFont
    block = reading_view.DisplayBlock("paragraph", list_item=True, spans=(
        reading_view.Span("😀 Pay "), reading_view.Span("30", True, True),
        reading_view.Span(" days.")))
    formats = screen._before.itemDelegate().text_formats(block, (), ((6, 8),))
    deleted = [item for item in formats if item.format.fontStrikeOut()]
    assert len(deleted) == 1
    item = deleted[0]
    assert (item.start, item.length) == (9, 2)  # bullet and emoji use UTF-16 positions
    assert item.format.fontWeight() == QFont.Weight.Bold
    assert item.format.fontItalic()
    color = item.format.foreground().color()
    assert color.red() > color.green()
    color = item.format.background().color()
    assert color.red() > color.green()


def test_review_of_a_small_replacement_has_no_removed_paragraph_notice(screen, qtbot):
    _show_deletion_review(screen)
    screen._cancel_review_btn.click()
    review(screen, qtbot)
    assert screen._before.model().deletions == (((28, 30),),)
    assert screen._deletion_notice.isVisible()
    assert "removed or replaced text" in screen._deletion_notice.text().lower()
    assert "original paragraph" not in screen._deletion_notice.text()


def test_review_insertions_clear_old_deletion_markers_and_notice(screen):
    _show_deletion_review(screen)
    screen._cancel_review_btn.click()
    screen.select_paragraph(1)
    screen._edit_part_btn.click()
    opened = screen._opened
    text = opened.paragraphs[1].text
    candidate = build_candidate(opened, 1, 1, Proposal((
        ProposedParagraph(0, "Please note: " + text),)))
    before = reading_view.paragraphs(opened.data, 1, 1)
    after = reading_view.paragraphs(candidate.data, 1, 1)
    candidate = reading_view.with_review_deletions(candidate, before, after)
    screen.show_proposal(candidate, before, after)
    assert screen._before.model().deletions == ((),)
    assert not screen._deletion_notice.isVisible()


def test_review_of_merged_paragraphs_keeps_their_retained_words_unmarked(screen):
    opened = screen._opened
    text = " ".join(p.text for p in opened.paragraphs[1:3])
    candidate = build_candidate(opened, 1, 2, Proposal((ProposedParagraph(0, text),)))
    before = reading_view.paragraphs(opened.data, 1, 2)
    after = reading_view.paragraphs(candidate.data, 1, 1)
    candidate = reading_view.with_review_deletions(candidate, before, after)
    screen.show_proposal(candidate, before, after)
    assert screen._before.model().deletions == ((), ())
    assert "1 original paragraph" in screen._deletion_notice.text()
    screen._use_btn.click()
    assert screen._pages.currentWidget() is screen._save_page


def test_reset_clears_before_deletions_and_the_review_notice(screen):
    _show_deletion_review(screen)
    screen.reset()
    assert screen._before.model().deletions == ()
    assert not screen._deletion_notice.isVisible()


def test_save_as_cancel_keeps_the_candidate(screen, qtbot, monkeypatch):
    review(screen, qtbot)
    candidate = screen._candidate
    screen._use_btn.click()
    calls = []

    def dialog(*args, **kwargs):
        calls.append((args, kwargs))
        return "", ""

    monkeypatch.setattr(QFileDialog, "getSaveFileName", dialog)
    screen._save_btn.click()
    assert screen._candidate is candidate and screen._pages.currentWidget() is screen._save_page
    assert Path(calls[0][0][2]) == screen._opened.path.with_name("service-outline (edited).docx")
    assert calls[0][0][3] == "Word document (*.docx)"
    assert not calls[0][1]["options"] & QFileDialog.Option.DontConfirmOverwrite


@pytest.mark.parametrize("outcome", ["retry", "destination", "original", "failed"])
def test_save_outcomes_retry_destination_original_failed(screen, qtbot, outcome):
    review(screen, qtbot)
    screen._use_btn.click()
    candidate = screen._candidate
    screen.show_failure("Save problem", outcome)
    if outcome == "original":
        assert screen._candidate is None and screen._opened is None
        assert screen._pages.currentWidget() is screen._open_page
    else:
        assert screen._candidate is candidate
        assert screen._pages.currentWidget() is screen._save_page
    assert screen._error.text() == "Save problem"


def test_saved_page_has_only_open_another_document(screen, qtbot):
    review(screen, qtbot)
    screen.show_saved("C:/synthetic/service-outline (edited).docx")
    assert not screen.has_work()
    assert screen._saved_path.text().endswith("service-outline (edited).docx")
    assert screen._saved_note.text() == "The original is unchanged."
    assert [b.text() for b in screen.findChildren(QPushButton) if b.isVisible()] == ["Open another document"]


def test_choose_another_file_asks_after_the_pick_only_with_work(screen, qtbot, monkeypatch):
    order = []
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (order.append("pick") or "", ""))
    monkeypatch.setattr(screen, "confirm_discard", lambda: order.append("ask") or False)
    describe(screen, qtbot)
    screen._choose_btn.click()
    assert order == ["pick"] and screen._selection == (1, 1)
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (order.append("pick") or "new.docx", ""))
    screen._choose_btn.click()
    assert order == ["pick", "pick", "ask"] and screen._selection == (1, 1)
    screen.reset()
    with qtbot.waitSignal(screen.open_requested):
        screen._browse_btn.click()
    assert order[-1] == "pick"


def test_keep_editing_restores_selection_description_and_scroll_position(screen, qtbot, monkeypatch):
    describe(screen, qtbot)
    screen._back_btn.click()
    screen._reader.verticalScrollBar().setValue(30)
    position = screen._reader.verticalScrollBar().value()
    monkeypatch.setattr(screen, "confirm_discard", lambda: False)
    screen.choose_path(Path("another.docx"))
    assert screen._selection == (1, 1)
    assert screen._description.toPlainText() == "Change 30 days to 45 days."
    assert screen._reader.verticalScrollBar().value() == position


def test_document_text_is_rendered_as_plain_text(screen, qtbot):
    blocks = (reading_view.DisplayBlock("paragraph", 0, spans=(reading_view.Span("<b>Literal & text</b>"),)),)
    screen._reader.set_blocks(blocks)
    layout = screen._reader.itemDelegate().make_layout(blocks[0], (), 400)
    assert layout.text() == "<b>Literal & text</b>"
    assert not any(f.format.fontWeight() > 400 for f in layout.formats())
    assert not screen._reader.grab().isNull()


def test_reading_view_keeps_word_line_breaks(screen):
    block = reading_view.DisplayBlock("paragraph", spans=(reading_view.Span("First line\nSecond line"),))
    layout = screen._reader.itemDelegate().make_layout(block, (), 600)
    assert layout.lineCount() == 2


def test_busy_blocks_drops_and_preserves_current_work(screen, qtbot):
    describe(screen, qtbot)
    screen.set_busy(True)
    screen.choose_path(Path("another.docx"))
    assert screen._opened.path.name == "service-outline.docx"
    assert not screen._choose_btn.isEnabled()
    screen.show_cancelled()
    screen.set_busy(False)
    assert screen._selection == (1, 1)


def test_reading_progress_stays_in_the_open_step(screen):
    screen.reset()
    screen._action = "open"
    screen.set_busy(True)
    assert screen._pages.currentWidget() is screen._open_page
    assert screen._open_busy_label.text() == "Reading the document..."
    assert screen._open_cancel_btn.isVisible()
    assert not screen._browse_btn.isEnabled()



@pytest.mark.parametrize("ratio", [1.0, 1.25, 1.5])
def test_drop_zone_outline_is_visible_on_all_four_edges(qtbot, ratio):
    from PySide6.QtGui import QColor, QImage
    from app.widgets.edit_screen import EditScreen
    page = EditScreen()
    qtbot.addWidget(page)
    page.resize(1240, 700)
    page.show()
    zone = page._drop_zone
    qtbot.waitUntil(zone.isVisible)
    # Render as a display at this scale would; Alex's 125% display lost the top edge.
    image = QImage(round(zone.width() * ratio), round(zone.height() * ratio), QImage.Format.Format_ARGB32)
    image.setDevicePixelRatio(ratio)
    image.fill(QColor("white"))
    zone.render(image, QPoint())
    w, h, band = image.width(), image.height(), round(3 * ratio)

    def line_weight(points):
        """Dark pixels per position along an edge, counted inside the outer band."""
        points = list(points)
        dark = sum(image.pixelColor(x, y).lightness() < 215 for x, y in points)
        return dark / (len(points) / band)

    xs, ys = range(20, w - 20), range(20, h - 20)
    edges = {
        "top": line_weight((x, y) for x in xs for y in range(band)),
        "bottom": line_weight((x, y) for x in xs for y in range(h - band, h)),
        "left": line_weight((x, y) for y in ys for x in range(band)),
        "right": line_weight((x, y) for y in ys for x in range(w - band, w)),
    }
    assert all(weight >= 1.0 for weight in edges.values()), edges


def test_drag_selects_a_range_and_respects_tables(screen, qtbot):
    from PySide6.QtCore import QEvent, QPointF
    from PySide6.QtGui import QMouseEvent
    from PySide6.QtWidgets import QApplication
    view = screen._reader

    def center(index):
        row = next(i for i, b in enumerate(view.model().blocks) if b.index == index)
        model_index = view.model().index(row, 0)
        view.scrollTo(model_index)
        qtbot.waitUntil(lambda: view.visualRect(model_index).height() > 0)
        return view.visualRect(model_index).center()

    def move(index, buttons=Qt.MouseButton.LeftButton):
        pos = QPointF(center(index))
        QApplication.sendEvent(view.viewport(), QMouseEvent(
            QEvent.Type.MouseMove, pos, view.viewport().mapToGlobal(pos), Qt.MouseButton.NoButton,
            buttons, Qt.KeyboardModifier.NoModifier))

    labels = [label.text() for label in screen.findChildren(QLabel)]
    assert "Click a paragraph, or drag across several. Shift-click extends the selection." in labels
    qtbot.mousePress(view.viewport(), Qt.MouseButton.LeftButton, pos=center(1))
    move(2)
    qtbot.mouseRelease(view.viewport(), Qt.MouseButton.LeftButton, pos=center(2))
    assert screen._selection == (1, 2)
    assert screen._selection_status.text() == "2 whole paragraph(s) selected."
    qtbot.mousePress(view.viewport(), Qt.MouseButton.LeftButton, pos=center(1))
    move(4)  # across the table: refused, the press's selection stays
    assert screen._selection == (1, 1) and "table" in screen._selection_notice.text()
    qtbot.mouseRelease(view.viewport(), Qt.MouseButton.LeftButton, pos=center(4))
    move(2, Qt.MouseButton.NoButton)  # moving without the button held selects nothing
    assert screen._selection == (1, 1)
