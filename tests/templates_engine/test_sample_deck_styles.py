# tests/templates_engine/test_sample_deck_styles.py
import importlib.util
from pathlib import Path

import pytest
from pptx import Presentation

from templates_engine.manifest import load_manifest
from tests.templates_engine.decks import company_template, free_form

_SCRIPT = Path(__file__).resolve().parents[2] / "scripts" / "sample_deck_styles.py"
_SPEC = importlib.util.spec_from_file_location("sample_deck_styles", _SCRIPT)
sample_deck_styles = importlib.util.module_from_spec(_SPEC)
_SPEC.loader.exec_module(sample_deck_styles)


def test_every_deck_is_sampled_on_the_route_the_app_takes(tmp_path):
    decks = tmp_path / "decks"
    decks.mkdir()
    free_form(decks / "canva.pptx")
    company_template(decks / "corp.pptx")
    (decks / "broken.pptx").write_bytes(b"not a deck")
    out = tmp_path / "samples"
    lines = sample_deck_styles.main(decks, out)
    corp_types = len(load_manifest(out / "corp" / "manifest.yaml").slide_types)
    assert lines == [
        "broken.pptx: FAILED (ValueError)",
        "canva.pptx: look captured. Taken from your deck: colors, fonts, background, logo. "
        "Its slide designs are not copied.",
        f"corp.pptx: own layouts kept, {corp_types} slide types.",
    ]
    assert len(Presentation(str(out / "canva-sample.pptx")).slides) == 4
    assert len(Presentation(str(out / "corp-sample.pptx")).slides) == corp_types
    assert (out / "report.txt").read_text(encoding="utf-8").count("\n") == 3


def test_sample_text_stays_within_every_field_limit(tmp_path):
    manifest = load_manifest(Path(__file__).parent / "fixtures" / "render_manifest.yaml")
    content = sample_deck_styles.sample_content(manifest, tmp_path / "picture.png")
    for slide in content["slides"]:
        definition = manifest.slide_types[slide["type"]]
        for field in definition.fields:
            value = slide["fields"][field.name]
            if field.max_chars:
                assert 0 < len(value) <= field.max_chars
            if field.max_items:
                assert len(value) == field.max_items


@pytest.mark.parametrize("inside", [False, True], ids=["same folder", "inside"])
def test_output_folder_overlapping_the_decks_is_refused(tmp_path, inside):
    decks = tmp_path / "decks"
    decks.mkdir()
    free_form(decks / "client.pptx")
    free_form(decks / "client-sample.pptx")
    before = {path.name: path.read_bytes() for path in decks.iterdir()}
    with pytest.raises(SystemExit):
        sample_deck_styles.main(decks, decks / "samples" if inside else decks)
    assert {path.name: path.read_bytes() for path in decks.iterdir()} == before
