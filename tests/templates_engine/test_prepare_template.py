# tests/templates_engine/test_prepare_template.py
import zipfile

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from templates_engine.config import MANIFEST_NAME, TEMPLATE_PPTX_NAME, prepare_template
from templates_engine.manifest import load_manifest
from tests.templates_engine.decks import (
    company_template, corrupt_png, damage_pictures, free_form, layout_background, master_branded,
    white_manual)


@pytest.mark.parametrize("factory,captured", [
    (free_form, True), (white_manual, True), (company_template, False), (master_branded, False)])
def test_layouts_are_kept_only_when_the_slides_follow_them(tmp_path, factory, captured):
    folder = tmp_path / "template"
    folder.mkdir()
    result = prepare_template(factory(tmp_path / "deck.pptx"), folder, "deck")
    assert result["captured"] is captured
    manifest = load_manifest(folder / MANIFEST_NAME)
    if captured:
        assert set(manifest.slide_types) == {"title", "agenda", "content", "closing"}
        assert result["report"].startswith("Taken from your deck:")
    else:
        assert "Taken from your deck" not in result["report"]
    assert len(Presentation(str(folder / TEMPLATE_PPTX_NAME)).slides) == 0


@pytest.mark.parametrize("data", [b"not an image", corrupt_png()], ids=["damaged", "undecodable"])
def test_damaged_pictures_are_skipped_and_the_rest_is_captured(tmp_path, data):
    folder = tmp_path / "template"
    folder.mkdir()
    deck = damage_pictures(free_form(tmp_path / "canva.pptx"), data)
    result = prepare_template(deck, folder, "canva")
    assert result["captured"] is True
    assert result["report"] == (
        "Taken from your deck: colors, fonts. Its slide designs are not copied.")
    with zipfile.ZipFile(folder / TEMPLATE_PPTX_NAME) as package:
        assert all(package.read(name) != data for name in package.namelist())


def test_inherited_layout_background_reaches_the_template(tmp_path):
    folder = tmp_path / "template"
    folder.mkdir()
    result = prepare_template(layout_background(tmp_path / "deck.pptx"), folder, "deck")
    assert result["captured"] is True
    assert result["report"] == (
        "Taken from your deck: colors, fonts, background. Its slide designs are not copied.")
    master = Presentation(str(folder / TEMPLATE_PPTX_NAME)).slide_master._element
    color = f"{qn('a:solidFill')}/{qn('a:srgbClr')}"
    assert master.cSld.bg.find(f"{qn('p:bgPr')}/{color}").get("val") == "0B3D2E"
    title = master.find(f"{qn('p:txStyles')}/{qn('p:titleStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
    assert title.find(color).get("val") == "FFFFFF"  # the deck's white text reads on dark green
