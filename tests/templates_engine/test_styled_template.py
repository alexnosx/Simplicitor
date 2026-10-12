# tests/templates_engine/test_styled_template.py
from dataclasses import replace

import yaml
from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from templates_engine.deck_style import Fill, capture_style
from templates_engine.manifest import load_manifest
from templates_engine.render_pptx import render
from templates_engine.styled_template import build_styled_template, readable, styled_manifest
from tests.templates_engine.decks import free_form, hand_built, png

CONTENT = {"slides": [
    {"type": "title", "fields": {"title": "Launch plan", "subtitle": "Q3"}},
    {"type": "agenda", "fields": {"heading": "Agenda", "items": ["Goals", "Timeline"]}},
    {"type": "content", "fields": {"heading": "Goals", "bullets": ["Ship", "Learn"]}},
    {"type": "closing", "fields": {"heading": "Thanks", "statement": "Questions welcome."}},
]}


def _theme(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    return etree.fromstring(part.blob)


def _title_color(prs) -> str:
    style = prs.slide_master._element.find(f"{qn('p:txStyles')}/{qn('p:titleStyle')}")
    path = f"{qn('a:lvl1pPr')}/{qn('a:defRPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}"
    return style.find(path).get("val")


def _build(tmp_path, style):
    out = tmp_path / "template.pptx"
    skipped = build_styled_template(style, out)
    return Presentation(str(out)), skipped, out


def test_free_form_style_becomes_a_wide_fillable_template(tmp_path):
    style = capture_style(free_form(tmp_path / "canva.pptx"))
    prs, skipped, _ = _build(tmp_path, style)
    assert skipped == []
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    title = prs.slide_layouts[0].placeholders[0]
    assert title.left + title.width <= prs.slide_width  # positions scaled to 16:9
    fonts = _theme(prs).find(f"{qn('a:themeElements')}/{qn('a:fontScheme')}")
    assert fonts.find(f"{qn('a:majorFont')}/{qn('a:latin')}").get("typeface") == "Montserrat"
    assert fonts.find(f"{qn('a:minorFont')}/{qn('a:latin')}").get("typeface") == "Open Sans"
    master_bg = prs.slide_master._element.cSld.bg
    assert master_bg.find(f"{qn('p:bgPr')}/{qn('a:blipFill')}") is not None
    cover_bg = prs.slide_layouts[0]._element.cSld.bg
    assert cover_bg is not None and cover_bg.find(f"{qn('p:bgPr')}/{qn('a:blipFill')}") is not None
    assert [s.name for s in prs.slide_master.shapes][-1] == "Logo"
    assert len(prs.slides) == 0


def test_captured_logo_keeps_master_shape_ids_below_signed_integer_limit(tmp_path):
    style = capture_style(free_form(tmp_path / "logo.pptx"))
    assert style.logo is not None
    prs, skipped, _ = _build(tmp_path, style)
    assert skipped == []
    assert any(shape.name == "Logo" for shape in prs.slide_master.shapes)
    assert all(shape.shape_id < 2**31 for shape in prs.slide_master.shapes)


def test_title_slide_keeps_light_text_on_its_dark_cover(tmp_path):
    style = capture_style(free_form(tmp_path / "canva.pptx"))
    prs, _, _ = _build(tmp_path, style)
    assert _title_color(prs) == "102040"  # content slides are light
    cover_title = prs.slide_layouts[0].placeholders[0]._element
    override = cover_title.find(
        f"{qn('p:txBody')}/{qn('a:lstStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}"
        f"/{qn('a:solidFill')}/{qn('a:srgbClr')}")
    assert override is not None and override.get("val") == "FFFFFF"


def test_unreadable_text_color_switches_to_black_or_white(tmp_path):
    style = capture_style(hand_built(tmp_path / "manual.pptx"))
    dark_text = replace(style, title=replace(style.title, color="0A0A0A"))
    prs, _, _ = _build(tmp_path, dark_text)
    assert _title_color(prs) == "FFFFFF"  # dark green background
    assert readable("FFFFFF", "0B3D2E") == "FFFFFF"
    assert readable("EEEEEE", "FFFFFF") == "000000"


def test_unreadable_background_picture_is_skipped_not_fatal(tmp_path):
    style = capture_style(free_form(tmp_path / "canva.pptx"))
    light_text = replace(style, title=replace(style.title, color="F5F5F5"))
    broken = replace(light_text, background=Fill(image=b"not an image", image_ext="png"))
    prs, skipped, _ = _build(tmp_path, broken)
    assert skipped == ["background"]
    assert len(prs.slide_layouts) == 11
    assert _title_color(prs) == "000000"  # judged against the white that is really there


def test_styled_template_renders_every_slide_type(tmp_path):
    style = capture_style(free_form(tmp_path / "canva.pptx"))
    folder = tmp_path / "library" / "canva"
    folder.mkdir(parents=True)
    build_styled_template(style, folder / "template.pptx")
    (folder / "manifest.yaml").write_text(
        yaml.dump(styled_manifest("canva", "canva.pptx"), allow_unicode=True), encoding="utf-8")
    manifest = load_manifest(folder / "manifest.yaml")
    result = render(manifest, CONTENT, tmp_path / "out.pptx", folder)
    deck = Presentation(str(result["path"]))
    assert result["issues"] == []
    assert [slide.shapes.title.text for slide in deck.slides] == [
        "Launch plan", "Agenda", "Goals", "Thanks"]


def test_portrait_deck_keeps_every_layout_inside_the_slide(tmp_path):
    style = capture_style(hand_built(tmp_path / "manual.pptx"))
    portrait = replace(style, slide_width=6858000, slide_height=9144000)
    prs, _, _ = _build(tmp_path, portrait)
    for layout in prs.slide_layouts:
        for shape in layout.placeholders:
            assert shape.left + shape.width <= prs.slide_width
            assert shape.top + shape.height <= prs.slide_height
