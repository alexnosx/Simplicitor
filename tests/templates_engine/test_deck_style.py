# tests/templates_engine/test_deck_style.py
from functools import partial

import pytest
from pptx import Presentation

from templates_engine.deck_style import capture_style, describe_capture, look_is_on_slides
from tests.templates_engine.decks import (
    PALE_ACCENT, company_template, corrupt_png, damage_pictures, dark_theme, free_form,
    hand_built, layout_background, layout_defaults, master_branded, paragraph_styled, png,
    two_masters, webp, white_manual)


def test_free_form_deck_gives_fonts_colors_backgrounds_and_logo(tmp_path):
    style = capture_style(free_form(tmp_path / "canva.pptx"))
    assert (style.slide_width, style.slide_height) == (12192000, 6858000)
    assert style.title.font == "Montserrat" and style.title.bold is True
    assert style.title.size_pt == 40.0
    assert style.body.font == "Open Sans" and style.body.size_pt == 18.0
    assert style.body.color == "333333"  # three light slides outweigh the dark cover
    assert style.title_background.image and style.background.image
    assert style.title_background != style.background
    assert style.logo is not None and 0.89 < style.logo.left < 0.91


def test_hand_built_deck_reads_slide_backgrounds_and_placeholder_titles(tmp_path):
    style = capture_style(hand_built(tmp_path / "manual.pptx"))
    assert style.background.color == "0B3D2E" and style.title_background.color == "0B3D2E"
    assert style.title.font == "Georgia" and style.title.color == "FFFFFF"
    assert style.body.font == "Verdana"
    assert style.logo is None


def test_paragraph_and_master_formatting_are_inherited(tmp_path):
    paragraph = capture_style(paragraph_styled(tmp_path / "paragraph.pptx"))
    assert (paragraph.title.font, paragraph.title.size_pt) == ("Georgia", 36.0)
    assert (paragraph.title.color, paragraph.title.bold) == ("FFFFFF", True)
    assert paragraph.body.font == "Verdana"
    branded = capture_style(master_branded(tmp_path / "branded.pptx"))
    assert (branded.title.font, branded.title.size_pt) == ("Georgia", 40.0)
    assert branded.title.color == "1F3864"


def test_default_paragraph_styles_are_inherited(tmp_path):
    style = capture_style(layout_defaults(tmp_path / "defaults.pptx"))
    assert (style.title.font, style.title.size_pt, style.title.color) == ("Georgia", 36.0, "1F3864")
    assert (style.body.font, style.body.size_pt, style.body.color) == ("Verdana", 18.0, "333333")


def test_each_slide_is_read_against_its_own_master(tmp_path):
    style = capture_style(two_masters(tmp_path / "two.pptx"))
    assert style.title.font == "Georgia" and style.body.font == "Verdana"  # two of three slides


def test_inherited_layout_background_is_captured(tmp_path):
    style = capture_style(layout_background(tmp_path / "layout.pptx"))
    assert style.background.color == "0B3D2E" and style.title_background.color == "0B3D2E"
    assert style.title.color == "FFFFFF"
    assert describe_capture(style) == (
        "Taken from your deck: colors, fonts, background. Its slide designs are not copied.")


@pytest.mark.parametrize("on_layout", [False, True], ids=["master", "layout"])
def test_dark_color_map_keeps_a_dark_background_and_light_text(tmp_path, on_layout):
    style = capture_style(dark_theme(tmp_path / "dark.pptx", on_layout))
    assert style.background.color == "000000" and style.title_background.color == "000000"
    assert style.title.color == "FFFFFF" and style.body.color == "FFFFFF"


def test_lighter_and_darker_theme_colors_are_applied(tmp_path):
    style = capture_style(layout_background(tmp_path / "pale.pptx", PALE_ACCENT))
    assert style.background.color == "F2DCDB"  # accent 2 itself is C0504D


def test_unstyled_text_follows_the_master_and_theme(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Plain title"
    slide.placeholders[1].text = "Plain body"
    prs.save(str(tmp_path / "plain.pptx"))
    style = capture_style(tmp_path / "plain.pptx")
    assert style.title.font == "Calibri" and style.body.font == "Calibri"
    assert style.title.color == "000000" and style.title.size_pt == 44.0  # the master's title
    assert style.background is None and style.title_background is None


def test_deck_without_slides_gives_theme_only(tmp_path):
    Presentation().save(str(tmp_path / "empty.pptx"))
    style = capture_style(tmp_path / "empty.pptx")
    assert style.color_scheme and style.logo is None and style.background is None


@pytest.mark.parametrize("factory,expected", [
    (free_form, True), (hand_built, True), (white_manual, True), (paragraph_styled, True),
    (company_template, False), (master_branded, False)])
def test_look_is_on_slides_when_slides_override_their_layouts(tmp_path, factory, expected):
    assert look_is_on_slides(factory(tmp_path / "deck.pptx")) is expected


@pytest.mark.parametrize("factory,expected", [
    (partial(layout_defaults, slide_fonts=("Georgia", "Verdana")), False),
    (partial(layout_defaults, slide_fonts=("Arial", "Arial")), True),
    (partial(two_masters, repeat_fonts=True), False),
    (partial(layout_background, own_fonts=False), False),
    (layout_background, True)],
    ids=["repeats-layout-fonts", "own-fonts", "repeats-its-master-fonts", "inherited-background",
         "inherited-background-own-fonts"])
def test_inherited_formatting_is_not_an_override(tmp_path, factory, expected):
    assert look_is_on_slides(factory(tmp_path / "deck.pptx")) is expected


def test_summary_names_only_what_was_found(tmp_path):
    full = capture_style(free_form(tmp_path / "canva.pptx"))
    assert describe_capture(full) == (
        "Taken from your deck: colors, fonts, background, logo. Its slide designs are not copied.")
    bare = capture_style(free_form(tmp_path / "bare.pptx", logo=False))
    assert "logo" not in describe_capture(bare)


def test_cover_only_picture_stays_on_the_title_slide(tmp_path):
    prs = Presentation()
    for number in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide {number}"
        if number == 0:
            slide.shapes.add_picture(png((10, 30, 70)), 0, 0, prs.slide_width, prs.slide_height)
    prs.save(str(tmp_path / "cover.pptx"))
    style = capture_style(tmp_path / "cover.pptx")
    assert style.title_background.image is not None
    assert style.background is None  # content slides follow the master


def test_broken_picture_links_are_ignored(tmp_path):
    from pptx.oxml.ns import qn
    prs = Presentation()
    for _ in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        picture = slide.shapes.add_picture(png((1, 2, 3)), 0, 0, prs.slide_width, prs.slide_height)
        slide.part.rels.pop(picture._element.blipFill.find(qn("a:blip")).get(qn("r:embed")))
    prs.save(str(tmp_path / "broken.pptx"))
    style = capture_style(tmp_path / "broken.pptx")
    assert style.background is None and style.logo is None
    assert look_is_on_slides(tmp_path / "broken.pptx") is False


@pytest.mark.parametrize("data", [b"not an image", webp(), corrupt_png()],
                         ids=["damaged", "unsupported", "undecodable"])
def test_unreadable_pictures_are_skipped_and_the_rest_is_kept(tmp_path, data):
    deck = damage_pictures(free_form(tmp_path / "canva.pptx"), data)
    style = capture_style(deck)
    assert style.background is None and style.title_background is None and style.logo is None
    assert style.title.font == "Montserrat"
    assert look_is_on_slides(deck) is True  # the slides still set their own fonts
