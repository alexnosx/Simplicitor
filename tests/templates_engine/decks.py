# tests/templates_engine/decks.py
"""Synthetic decks shaped like the kinds users upload."""
import copy
import io
import struct
import zipfile
import zlib
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.parts.slide import SlideLayoutPart, SlideMasterPart
from pptx.util import Emu, Pt

WIDE = (Emu(12192000), Emu(6858000))  # 16:9


def png(color: tuple[int, int, int], size: tuple[int, int] = (64, 36)) -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", size, color).save(buffer, "PNG")
    buffer.seek(0)
    return buffer


def _text(slide, text: str, top: float, font: str, size: int, color: str, bold: bool = False):
    width, height = WIDE
    box = slide.shapes.add_textbox(round(width * 0.08), round(height * top),
                                   round(width * 0.84), round(height * 0.15))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def free_form(path: Path, slides: int = 4, logo: bool = True) -> Path:
    """Like a Canva export: blank layouts, a full-slide picture, free text boxes, a logo."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = WIDE
    width, height = WIDE
    for number in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
        cover = (10, 30, 70) if number == 0 else (240, 240, 235)
        slide.shapes.add_picture(png(cover), 0, 0, width, height)
        dark = number == 0
        _text(slide, f"Heading {number}", 0.08, "Montserrat", 40,
              "FFFFFF" if dark else "102040", True)
        _text(slide, "Body text that explains the point in a full sentence.", 0.45,
              "Open Sans", 18, "EEEEEE" if dark else "333333")
        if logo:
            slide.shapes.add_picture(png((200, 20, 20), (16, 16)),
                                     round(width * 0.9), round(height * 0.03),
                                     round(width * 0.06), round(height * 0.08))
    prs.save(str(path))
    return path


def hand_built(path: Path) -> Path:
    """Built in PowerPoint slide by slide: standard layouts, own solid backgrounds."""
    prs = Presentation()
    for number in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0x0B, 0x3D, 0x2E)
        title = slide.shapes.title.text_frame.paragraphs[0]
        run = title.add_run()
        run.text = f"Topic {number}"
        run.font.name, run.font.size = "Georgia", Pt(36)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        body = slide.placeholders[1].text_frame.paragraphs[0].add_run()
        body.text = "A point written by hand on this slide."
        body.font.name, body.font.color.rgb = "Verdana", RGBColor(0xE0, 0xE0, 0xE0)
    prs.save(str(path))
    return path


def company_template(path: Path) -> Path:
    """A real template: the look lives in the master, slides follow it."""
    prs = Presentation()
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb = RGBColor(0x20, 0x20, 0x60)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    prs.save(str(path))
    return path


def white_manual(path: Path) -> Path:
    """Built by hand on the plain white theme: own fonts and a logo on every slide."""
    prs = Presentation()
    width, height = prs.slide_width, prs.slide_height
    for number in range(3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title.text_frame.paragraphs[0].add_run()
        title.text, title.font.name = f"Topic {number}", "Georgia"
        body = slide.placeholders[1].text_frame.paragraphs[0].add_run()
        body.text, body.font.name = "A point made by hand.", "Verdana"
        slide.shapes.add_picture(png((200, 20, 20), (16, 16)),
                                 round(width * 0.9), round(height * 0.03),
                                 round(width * 0.06), round(height * 0.08))
    prs.save(str(path))
    return path


def paragraph_styled(path: Path) -> Path:
    """Titles and body text styled on the paragraph instead of on each run."""
    prs = Presentation()
    for number in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title.text_frame.paragraphs[0]
        title.text = f"Topic {number}"
        title.font.name, title.font.size, title.font.bold = "Georgia", Pt(36), True
        title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        body = slide.placeholders[1].text_frame.paragraphs[0]
        body.text = "A point made on this slide."
        body.font.name = "Verdana"
    prs.save(str(path))
    return path


def master_branded(path: Path) -> Path:
    """A company template: the title font, size and color are set once, on the master."""
    prs = Presentation()
    styles = prs.slide_master._element.find(f"{qn('p:txStyles')}/{qn('p:titleStyle')}")
    title = styles.find(f"{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
    title.set("sz", "4000")
    title.remove(title.find(qn("a:solidFill")))
    navy = f'<a:solidFill {nsdecls("a")}><a:srgbClr val="1F3864"/></a:solidFill>'
    title.insert(0, parse_xml(navy))
    title.find(qn("a:latin")).set("typeface", "Georgia")
    for number in range(2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Quarter {number}"
        slide.placeholders[1].text = "Results for the quarter."
    prs.save(str(path))
    return path


DARK_GREEN = f'<a:solidFill {nsdecls("a")}><a:srgbClr val="0B3D2E"/></a:solidFill>'
# What PowerPoint writes for the color "Accent 2, Lighter 80%".
PALE_ACCENT = (f'<a:solidFill {nsdecls("a")}><a:schemeClr val="accent2"><a:lumMod val="20000"/>'
               '<a:lumOff val="80000"/></a:schemeClr></a:solidFill>')


def layout_defaults(path: Path, slide_fonts: tuple[str, str] | None = None) -> Path:
    """A template whose layout sets fonts in its placeholders' default paragraph style.

    With *slide_fonts*, the slides also set (title font, body font) on every run.
    """
    prs = Presentation()
    layout = prs.slide_layouts[1]
    for shape, font, size, color in ((layout.placeholders[0], "Georgia", 3600, "1F3864"),
                                     (layout.placeholders[1], "Verdana", 1800, "333333")):
        styles = shape._element.find(f"{qn('p:txBody')}/{qn('a:lstStyle')}")
        styles.insert(0, parse_xml(
            f'<a:defPPr {nsdecls("a")}><a:defRPr sz="{size}"><a:solidFill>'
            f'<a:srgbClr val="{color}"/></a:solidFill><a:latin typeface="{font}"/></a:defRPr>'
            '</a:defPPr>'))
    for number in range(2):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Topic {number}"
        slide.placeholders[1].text = "A point made on this slide."
        if slide_fonts:
            for shape, font in zip((slide.shapes.title, slide.placeholders[1]), slide_fonts):
                shape.text_frame.paragraphs[0].runs[0].font.name = font
    prs.save(str(path))
    return path


def layout_background(path: Path, fill: str = DARK_GREEN, own_fonts: bool = True) -> Path:
    """Slides inheriting their layout's background; with *own_fonts*, white text set on each."""
    prs = Presentation()
    layout = prs.slide_layouts[1]
    properties = layout._element.cSld.get_or_add_bgPr()
    properties.replace(properties[0], parse_xml(fill))
    for number in range(3):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Topic {number}"
        slide.placeholders[1].text = "A point made on this slide."
        if own_fonts:
            for shape, font in ((slide.shapes.title, "Georgia"), (slide.placeholders[1], "Verdana")):
                run = shape.text_frame.paragraphs[0].runs[0]
                run.font.name, run.font.color.rgb = font, RGBColor(0xFF, 0xFF, 0xFF)
    prs.save(str(path))
    return path


def dark_theme(path: Path, on_layout: bool = False) -> Path:
    """Built by hand on a dark theme variant: the color map swaps background and text.

    The swap is on the master, or with *on_layout* only on the Blank layout the slides use.
    """
    prs = Presentation()
    layout = prs.slide_layouts[6]  # Blank
    mapping = prs.slide_master._element.find(qn("p:clrMap"))
    if on_layout:
        override = layout._element.find(qn("p:clrMapOvr"))
        mapping = parse_xml(f'<a:overrideClrMapping {nsdecls("a")}/>')
        for name, slot in prs.slide_master._element.find(qn("p:clrMap")).attrib.items():
            mapping.set(name, slot)
        override.replace(override[0], mapping)
    for name, slot in {"bg1": "dk1", "tx1": "lt1", "bg2": "dk2", "tx2": "lt2"}.items():
        mapping.set(name, slot)
    width, height = prs.slide_width, prs.slide_height
    for number in range(3):
        slide = prs.slides.add_slide(layout)
        for text, top, font, size in ((f"Topic {number}", 0.08, "Georgia", 40),
                                      ("Body text in the theme's text color.", 0.45, "Verdana", 18)):
            box = slide.shapes.add_textbox(round(width * 0.08), round(height * top),
                                           round(width * 0.84), round(height * 0.15))
            run = box.text_frame.paragraphs[0].add_run()
            run.text, run.font.name, run.font.size = text, font, Pt(size)
    prs.save(str(path))
    return path


def two_masters(path: Path, repeat_fonts: bool = False) -> Path:
    """Slides pasted from another deck: a second master whose theme uses Georgia and Verdana.

    One slide uses the first master and two use the second. With *repeat_fonts*, the
    second master's slides also set its fonts on every run.
    """
    prs = Presentation()
    first, package = prs.slide_master.part, prs.part.package
    theme = first.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    for kind, font in (("a:majorFont", "Georgia"), ("a:minorFont", "Verdana")):
        root.find(f".//{qn(kind)}/{qn('a:latin')}").set("typeface", font)
    second_theme = Part(PackURI("/ppt/theme/theme2.xml"), theme.content_type, package,
                        etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                                       standalone=True))
    master = SlideMasterPart(PackURI("/ppt/slideMasters/slideMaster2.xml"),
                             first.content_type, package, copy.deepcopy(first._element))
    source = prs.slide_layouts[1].part
    layout = SlideLayoutPart(PackURI("/ppt/slideLayouts/slideLayout12.xml"),
                             source.content_type, package, copy.deepcopy(source._element))
    layout.relate_to(master, RT.SLIDE_MASTER)
    master.relate_to(second_theme, RT.THEME)
    layouts = master._element.find(qn("p:sldLayoutIdLst"))
    for entry in list(layouts)[1:]:
        layouts.remove(entry)
    layouts[0].set("id", "2147483700")
    layouts[0].set(qn("r:id"), master.relate_to(layout, RT.SLIDE_LAYOUT))
    masters = prs.part._element.find(qn("p:sldMasterIdLst"))
    entry = copy.deepcopy(masters[0])
    entry.set("id", "2147483699")
    entry.set(qn("r:id"), prs.part.relate_to(master, RT.SLIDE_MASTER))
    masters.append(entry)
    pasted = prs.slide_masters[1].slide_layouts[0]
    for number, chosen in enumerate((prs.slide_layouts[1], pasted, pasted)):
        slide = prs.slides.add_slide(chosen)
        slide.shapes.title.text = f"Topic {number}"
        slide.placeholders[1].text = "A point made on this slide."
        if repeat_fonts and chosen is pasted:
            for shape, font in ((slide.shapes.title, "Georgia"), (slide.placeholders[1], "Verdana")):
                shape.text_frame.paragraphs[0].runs[0].font.name = font
    prs.save(str(path))
    return path


def corrupt_png() -> bytes:
    """A PNG whose header and checksums are valid but whose pixel data cannot be decoded."""
    buffer = io.BytesIO()
    Image.new("RGB", (16, 16), (9, 9, 9)).save(buffer, "PNG")
    data = buffer.getvalue()
    start = data.index(b"IDAT") - 4
    length = struct.unpack(">I", data[start:start + 4])[0]
    payload = b"\x78\x9c" + b"\xff" * 10
    chunk = (struct.pack(">I", len(payload)) + b"IDAT" + payload
             + struct.pack(">I", zlib.crc32(b"IDAT" + payload)))
    return data[:start] + chunk + data[start + 12 + length:]


def webp() -> bytes:
    """A real picture in a format python-pptx does not support."""
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (9, 9, 9)).save(buffer, "WEBP")
    return buffer.getvalue()


def damage_pictures(path: Path, data: bytes) -> Path:
    """Replace the bytes of every embedded picture; the deck itself stays readable."""
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            content = source.read(item.filename)
            target.writestr(item, data if item.filename.startswith("ppt/media/") else content)
    return path
