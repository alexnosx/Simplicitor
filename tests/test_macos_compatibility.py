"""Portable source harness and preservation contracts on the actual test filesystem."""
import os
from pathlib import Path
import sys

import pytest


ROOT = Path(__file__).parents[1]


@pytest.mark.parametrize("platform", ["darwin", "win32"])
@pytest.mark.parametrize("script,boundary", [
    ("scripts/capture_screens.py", "from PySide6.QtCore"),
    ("tests/conftest.py", "\nimport gc"),
])
def test_source_harness_font_override_is_windows_only(tmp_path, monkeypatch, platform, script, boundary):
    monkeypatch.setattr(os, "environ", os.environ.copy())
    monkeypatch.setattr(sys, "platform", platform)
    monkeypatch.setattr(sys, "path", sys.path.copy())
    monkeypatch.delenv("QT_QPA_FONTDIR", raising=False)
    prefix = (ROOT / script).read_text(encoding="utf-8").split(boundary, 1)[0]
    namespace = {"__file__": str(tmp_path / script)}
    exec(compile(prefix, script, "exec"), namespace)
    if platform == "darwin":
        assert "QT_QPA_FONTDIR" not in os.environ
    else:
        assert os.environ["QT_QPA_FONTDIR"].endswith("Fonts")


def test_capture_isolates_darwin_home_before_importing_settings(tmp_path, monkeypatch):
    monkeypatch.setattr(os, "environ", os.environ.copy())
    monkeypatch.setattr(sys, "platform", "darwin")
    monkeypatch.setattr(sys, "path", sys.path.copy())
    monkeypatch.setenv("HOME", str(tmp_path / "outer"))
    source = (ROOT / "scripts/capture_screens.py").read_text(encoding="utf-8")
    prefix = source.split("from PySide6.QtCore", 1)[0]
    namespace = {"__file__": str(tmp_path / "scripts/capture_screens.py")}
    exec(compile(prefix, "capture-prefix", "exec"), namespace)
    assert os.environ["HOME"] == str(namespace["PROFILE"])
    assert source.index("from app.config.settings") > len(prefix)


def test_isolated_profile_sets_home_under_darwin_semantics(tmp_path, monkeypatch, request):
    monkeypatch.setattr(sys, "platform", "darwin")
    monkeypatch.setenv("HOME", str(tmp_path / "outer"))
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: Path(os.environ["HOME"])))
    settings = request.getfixturevalue("isolated_profile")
    assert Path.home() == tmp_path / "home"
    assert Path(settings.generated_dir).is_relative_to(Path.home())


def _candidate(tmp_path, route):
    from editing.models import Proposal, ProposedParagraph
    from editing.package import open_document
    from editing.patch import build_candidate
    from extraction.jobs import create_job
    from extraction.models import Candidate
    from openpyxl import Workbook
    from tests.editing.docx_builder import build_docx, p

    source = tmp_path / "Pilot É 100%.docx"
    source.write_bytes(build_docx(p("Pay in 30 days.")))
    if route == "edit":
        candidate = build_candidate(open_document(source), 0, 0,
                                    Proposal((ProposedParagraph(0, "Pay in 45 days."),)))
    else:
        job = create_job(tmp_path / "app")
        path = job / "candidate.xlsx"
        book = Workbook()
        book.active.title = "Data"
        book.active.append(["ID"])
        book.active.append(["000452"])
        book.create_sheet("Evidence").append(["Synthetic source evidence"])
        book.save(path)
        book.close()
        candidate = Candidate(job.name, path, (source,), ())
    return source, candidate


def _save(candidate, destination, route):
    if route == "edit":
        from editing.save import save_candidate
        return save_candidate(candidate, destination)
    from extraction.jobs import save_candidate
    return save_candidate(candidate, destination, False)


@pytest.mark.parametrize("route", ["edit", "extract"])
@pytest.mark.parametrize("kind", ["symlink", "hardlink", "case"])
def test_filesystem_aliases_cannot_overwrite_originals(tmp_path, record_property, route, kind):
    from editing.models import EditError
    source, candidate = _candidate(tmp_path, route)
    before = source.read_bytes()
    alias = tmp_path / ("alias.docx" if route == "edit" else "alias.xlsx")
    try:
        if kind == "symlink":
            alias.symlink_to(source)
        elif kind == "hardlink":
            os.link(source, alias)
        else:
            alias = source.with_name(source.name.upper())
            if not alias.exists() or not source.samefile(alias):
                record_property("filesystem_capability", f"{kind}: not supported")
                pytest.skip("Test volume does not provide a case alias")
    except OSError as exc:
        record_property("filesystem_capability", f"{kind}: unavailable ({type(exc).__name__})")
        pytest.skip(f"Test volume cannot create {kind}: {type(exc).__name__}")
    record_property("filesystem_capability", f"{kind}: supported")
    with pytest.raises((EditError, ValueError), match="original|source"):
        _save(candidate, alias, route)
    assert source.read_bytes() == before


@pytest.mark.parametrize("route", ["edit", "extract"])
def test_unicode_save_and_denied_write_preserve_sources(tmp_path, monkeypatch, route):
    from editing.models import EditError
    from openpyxl import load_workbook
    source, candidate = _candidate(tmp_path, route)
    before = source.read_bytes()
    destination = tmp_path / ("Edited É 100%.docx" if route == "edit" else "Table É 100%.xlsx")
    assert _save(candidate, destination, route) == destination
    assert source.read_bytes() == before
    if route == "edit":
        assert destination.read_bytes() == candidate.data
    else:
        book = load_workbook(destination)
        assert book["Data"]["A2"].value == "000452"
        assert "Evidence" in book.sheetnames
        book.close()
    previous_output = destination.read_bytes()

    def deny(*args, **kwargs):
        raise PermissionError("Synthetic denied write")

    monkeypatch.setattr(os, "replace", deny)
    with pytest.raises((EditError, OSError), match="save|denied|Synthetic"):
        _save(candidate, destination, route)
    assert source.read_bytes() == before
    assert destination.read_bytes() == previous_output


def test_denied_document_open_has_a_visible_error_and_preserves_bytes(tmp_path, monkeypatch):
    from editing.models import EditError
    from editing.package import open_document
    from tests.editing.docx_builder import build_docx, p
    source = tmp_path / "Synthetic.docx"
    source.write_bytes(build_docx(p("Synthetic paragraph.")))
    before = source.read_bytes()
    read = Path.read_bytes

    def deny(path):
        if path == source:
            raise PermissionError("Synthetic denied open")
        return read(path)

    monkeypatch.setattr(Path, "read_bytes", deny)
    with pytest.raises(EditError):
        open_document(source)
    assert read(source) == before
