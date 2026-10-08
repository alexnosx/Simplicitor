"""Source workspace behavior, saved evidence, and native Save As delegation."""
from dataclasses import replace
from pathlib import Path
import json
from threading import Event

from docx import Document
from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QTableWidget
import pytest

from app.config.settings import Settings
from app.widgets.extraction_panel import ExtractionPanel
from extraction.models import ColumnSpec, ExtractionResult, FieldProposal, FieldResult, Issue
from extraction.source_readers import read_sources
from extraction.xlsx_writer import read_candidate, write_candidate


@pytest.fixture
def panel(qtbot, tmp_path):
    widget = ExtractionPanel()
    qtbot.addWidget(widget)
    widget.set_ollama_connected(True)
    widget.set_model("qwen3:8b")
    path = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("ID 00123; date 01/02/2026")
    document.save(path)
    widget.set_paths((path,))
    widget._request_edit.setPlainText("Extract the ID")
    widget.show_inspection(read_sources((path,)))
    return widget


def saved_candidate(tmp_path, *, flagged=True, coverage=False):
    issue = Issue("request_failed", "source-1", "source-1#p:1", "A section failed.")
    field = FieldResult(FieldProposal("000452", "ID 000452", "source-1#p:0"),
                        "000452", flagged, ("leading_zero",) if flagged else ())
    result = ExtractionResult(("source-1",), {("source-1", "id"): field},
                              {"source-1": tmp_path / "source.docx"},
                              (issue,) if coverage else (), {})
    return write_candidate(result, (ColumnSpec("id", "ID", "Identifier"),),
                           tmp_path / "candidate.xlsx")


@pytest.mark.parametrize("value,kind,expected", [
    ("12500.00", "decimal", "12500.00"),
    ("12.5000", "decimal", "12.5000"),
    ("-2.50", "decimal", "-2.50"),
    ("123.100000000000000", "decimal", "123.100000000000000"),
    ("9999999999999.9000", "decimal", "9999999999999.9000"),
    ("0.00000000000000001", "decimal", "1E-17"),
    ("12500", "integer", "12500"),
    ("000452", "text", "000452"),
])
def test_review_uses_saved_number_format_and_file_without_record(panel, tmp_path,
                                                                value, kind, expected):
    from decimal import Decimal
    typed = Decimal(value) if kind == "decimal" else int(value) if kind == "integer" else value
    field = FieldResult(FieldProposal(value, value, "source-1#p:0"), typed, False, ())
    result = ExtractionResult(("source-1",), {("source-1", "value"): field},
                              {"source-1": tmp_path / "source.docx"}, (), {})
    candidate = write_candidate(result, (ColumnSpec("value", "Value", "Value", kind),),
                                tmp_path / "candidate.xlsx")
    panel.show_candidate(candidate, read_candidate(candidate))
    assert panel._grid.columnCount() == 2
    assert panel._grid.horizontalHeaderItem(0).text() == "File"
    assert panel._grid.item(0, 0).text() == "source.docx"
    assert panel._grid.item(0, 1).text() == expected
    panel._grid.setCurrentCell(0, 1)
    assert "Record: source-1" in panel._evidence.toPlainText()
    book = load_workbook(candidate.path)
    assert book["Evidence"]["B1"].value == "Record"
    book.close()


def test_source_picker_preserves_duplicate_attachments_without_copying(panel, monkeypatch):
    path = panel._paths[0]
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *a, **k:
                        ([str(path), str(path)], ""))
    panel._add_btn.click()
    assert panel._paths == (path, path, path)
    assert panel._source_list.count() == 3
    assert all(p == path for p in panel._paths)
    assert not panel._confirmed


def test_columns_add_edit_remove_and_freeze_before_extract(panel):
    panel.set_columns((ColumnSpec("id", "ID", "Identifier", "unknown"),))
    assert panel.current_columns()[0].kind == "text"
    assert not panel._extract_btn.isEnabled()
    panel._add_column_btn.click()
    assert len(panel.current_columns()) == 2
    panel._columns_table.selectRow(1)
    panel._remove_column_btn.click()
    panel._columns_table.item(0, 0).setText("Customer ID")
    panel._confirm_btn.click()
    assert panel._extract_btn.isEnabled()
    with pytest.raises(Exception):  # dataclass is frozen
        panel.snapshot().request = "changed"
    with pytest.raises(ValueError):
        panel.set_columns((ColumnSpec("id", "", ""),))
        panel.current_columns()


def test_numeric_defaults_overrides_and_conditional_date_order(panel):
    panel.set_columns((ColumnSpec("amount", "Amount", "Amount", "decimal"),
                       ColumnSpec("date", "Date", "Date", "date")))
    panel._columns_table.selectRow(0)
    assert panel._thousands_edit.text() == "," and panel._decimal_edit.text() == "."
    panel._thousands_edit.setText(".")
    panel._decimal_edit.setText(",")
    assert panel.current_columns()[0].decimal_separator == ","
    panel._columns_table.selectRow(1)
    assert panel._date_order_combo.isVisible()
    panel._confirm_btn.click()
    assert not panel._confirmed and not panel._extract_btn.isEnabled()
    panel._date_order_combo.setCurrentIndex(1)
    panel._confirm_btn.click()
    assert panel._confirmed and panel.current_columns()[1].date_order == "DMY"
    documents = tuple(replace(d, units=tuple(replace(u, text="15 March 2026") for u in d.units))
                      for d in panel._documents)
    panel.show_inspection(documents)
    assert not panel._date_order_combo.isVisible()


@pytest.mark.parametrize("thousands,decimal", [("1", "."), ("-", "."), ("+", "."),
                                              ("e", "."), ("$", "."), (",", "1"),
                                              (",", "+"), (",", " ")])
def test_invalid_numeric_separators_cannot_confirm_or_start_extraction(panel, thousands, decimal):
    panel.set_columns((ColumnSpec("amount", "Amount", "Amount", "decimal"),))
    panel._thousands_edit.setText(thousands)
    panel._decimal_edit.setText(decimal)
    panel._confirm_btn.click()
    assert not panel._confirmed and not panel._extract_btn.isEnabled()
    with pytest.raises(ValueError, match="separator"):
        panel.snapshot()


@pytest.mark.parametrize("thousands,decimal,value", [(",", ".", "1,234.50"),
                                                    (".", ",", "1.234,50"),
                                                    (" ", ",", "1 234,50"),
                                                    ("'", ".", "1'234.50"),
                                                    ("", ".", "1234.50")])
def test_supported_numeric_separator_overrides_preserve_values(panel, thousands, decimal, value):
    from decimal import Decimal
    from extraction.grounding import validate_field
    from extraction.models import SourceUnit
    panel.set_columns((ColumnSpec("amount", "Amount", "Amount", "decimal"),))
    panel._thousands_edit.setText(thousands)
    panel._decimal_edit.setText(decimal)
    panel._confirm_btn.click()
    assert panel._confirmed
    unit = SourceUnit("one#p:0", value, "paragraph", 0)
    result = validate_field(FieldProposal(value, value, unit.anchor), panel.current_columns()[0],
                            {unit.anchor: unit})
    assert not result.flagged and result.typed_value == Decimal("1234.50")


def test_many_coverage_issues_stay_scrollable_and_keep_controls_accessible(panel, qtbot):
    issues = tuple(Issue("low_text", "source-1", f"source-1#page:{page}",
                         "A PDF page has little readable text.") for page in range(2, 102))
    documents = (replace(panel._documents[0], issues=issues, page_cost=101),)
    panel.resize(1000, 650)
    panel.show_inspection(documents)
    qtbot.wait(20)
    assert panel.minimumSizeHint().height() <= 650 and panel.height() <= 650
    assert panel._coverage_view.height() <= 100
    assert panel._coverage_view.verticalScrollBar().maximum() > 0
    assert "source-1#page:2" in panel._coverage_view.toPlainText()
    assert "source-1#page:101" in panel._coverage_view.toPlainText()
    panel.show_failure("Some pages were not processed.", issues)
    assert panel._coverage_view.height() <= 100
    assert "source-1#page:101" in panel._coverage_view.toPlainText()


def test_failure_retains_editable_setup_and_disconnection_gates_model_calls(panel):
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel.set_busy(True)
    panel.show_failure("Could not suggest columns.", ())
    panel.set_busy(False)
    assert panel._request_edit.toPlainText() == "Extract the ID"
    assert panel.current_columns()[0].label == "ID"
    panel._confirm_btn.click()
    panel.set_ollama_connected(False)
    assert not panel._extract_btn.isEnabled() and not panel._suggest_btn.isEnabled()
    assert panel._add_column_btn.isEnabled()
    panel.set_ollama_connected(True)
    assert panel._extract_btn.isEnabled()


def test_saved_grid_literal_flags_evidence_acknowledgement_and_rerun_reset(panel, tmp_path):
    candidate = saved_candidate(tmp_path, coverage=True)
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel._confirm_btn.click()
    panel.show_candidate(candidate, read_candidate(candidate))
    cell = panel._grid.item(0, 1)
    assert cell.text() == "000452" and cell.background().color().isValid()
    assert panel._grid.editTriggers() == QTableWidget.EditTrigger.NoEditTriggers
    panel._grid.setCurrentCell(0, 1)
    assert "ID 000452" in panel._evidence.toPlainText()
    assert "leading_zero" in panel._evidence.toPlainText()
    assert "A section failed." in panel._coverage_view.toPlainText()
    assert not panel._save_btn.isEnabled()
    panel._ack_checkbox.setChecked(True)
    assert panel._save_btn.isEnabled()
    panel.clear_review()
    assert panel._candidate is None and panel._grid.rowCount() == 0
    assert not panel._ack_checkbox.isChecked()


@pytest.mark.parametrize("change", ["request", "column", "sources"])
def test_setup_changes_clear_confirmation_and_review(panel, tmp_path, change):
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel._confirm_btn.click()
    candidate = saved_candidate(tmp_path)
    panel.show_candidate(candidate, read_candidate(candidate))
    if change == "request":
        panel._request_edit.setPlainText("Another request")
    elif change == "column":
        panel._columns_table.item(0, 0).setText("New ID")
    else:
        panel.set_paths(panel._paths + panel._paths)
    assert not panel._confirmed and panel._candidate is None


def test_busy_setup_is_frozen_and_save_can_finish_without_cancel(panel):
    before = panel._paths
    panel.set_busy(True)
    panel.set_paths((Path("different.docx"),))
    assert panel._paths == before and not panel._suggest_btn.isEnabled()
    assert panel._cancel_btn.isEnabled()
    panel.set_busy(True, saving=True)
    assert not panel._cancel_btn.isEnabled()


@pytest.fixture
def window(qtbot, tmp_path, monkeypatch):
    from app.main_window import MainWindow
    from app.services.ollama_client import OllamaClient
    def offline_start(self):
        self._ollama_client = OllamaClient("http://localhost:11434")
    monkeypatch.setattr(MainWindow, "_start_ollama_worker", offline_start)
    settings = Settings(tmp_path)
    settings.set("generated_dir", str(tmp_path / "outputs"))
    widget = MainWindow(settings)
    qtbot.addWidget(widget)
    yield widget
    widget.close()


def test_create_workspace_defaults_to_sources_and_preserves_prompt_host(window):
    from app.widgets.create_panel import CreatePanel
    assert window._workspace._tabs.tabText(0) == "From source files"
    assert window._workspace._tabs.tabText(1) == "From prompt"
    assert window._workspace._tabs.currentIndex() == 0
    assert window._extraction_panel.isVisible()
    assert isinstance(window._create_panel, CreatePanel)
    assert not window._edit_panel.isVisible()
    assert not window._extraction_panel._cancel_btn.isEnabled()
    assert window.width() >= 1200 and window.height() >= 800
    window._workspace._tabs.setCurrentIndex(1)
    assert window._create_panel.isVisible()


@pytest.mark.parametrize("count,warn", [(0, True), (7_999_999_999, True),
                                       (8_000_000_000, False), (27_000_000_000, False)])
def test_selected_model_warning(window, count, warn):
    window._on_ollama_connected(["selected", "loaded"], "loaded")
    window._top_bar._model_combo.setCurrentText("selected")
    window._on_model_params_ready("selected", count)
    assert window._capability_banner.isVisible() is warn
    assert window._capability_banner._info_label.text() == (
        "Simplicitor works best with models of 8B parameters or more.")
    window._create_panel._prompt_edit.setPlainText("Write a report")
    assert window._create_panel._generate_btn.isEnabled()
    window._on_model_params_ready("loaded", 1)
    window._on_ollama_connected(["selected", "loaded"], "loaded")
    assert window._current_model == "selected"
    assert window._top_bar.current_model() == "selected"
    assert window._capability_banner.isVisible() is warn


def test_banner_dismissal_is_per_selection_and_metadata_lookup_is_queued(window, qtbot):
    with qtbot.waitSignal(window._model_params_requested) as request:
        window._on_ollama_connected(["small", "other"], "small")
    assert request.args == ["small"]
    window._on_model_params_ready("small", 1)
    window._capability_banner._dismiss_btn.click()
    window._on_model_params_ready("small", 1)
    assert not window._capability_banner.isVisible()
    window._top_bar._model_combo.setCurrentText("other")
    window._on_model_params_ready("other", 0)
    assert window._capability_banner.isVisible()


def test_empty_model_discovery_clears_the_selection_and_warning(window):
    window._on_ollama_connected(["removed-model"], "")
    window._on_model_params_ready("removed-model", 1)
    assert window._capability_banner.isVisible()
    window._on_ollama_connected([], "")
    assert window._top_bar.current_model() == ""
    assert window._current_model == "" and window._extraction_panel._model == ""
    assert not window._capability_banner.isVisible()


@pytest.mark.parametrize("cancel", [False, True])
def test_native_save_as_xlsx_filter_confirmation_and_cancellation(window, tmp_path, monkeypatch, cancel):
    candidate = saved_candidate(tmp_path, flagged=False)
    panel = window._extraction_panel
    panel.show_candidate(candidate, read_candidate(candidate))
    destination = tmp_path / "chosen.xlsx"
    requests = []
    def dialog(parent, caption, directory, filter, **kwargs):
        requests.append((filter, kwargs.get("options", QFileDialog.Option(0))))
        return ("" if cancel else str(destination), "Excel workbook (*.xlsx)")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", dialog)
    operations = []
    monkeypatch.setattr(window, "_start_extraction", lambda action, setup, **kw:
                        operations.append((action, kw)))
    panel._save_btn.click()
    assert requests[0][0] == "Excel workbook (*.xlsx)"
    assert not requests[0][1] & QFileDialog.Option.DontConfirmOverwrite
    assert not requests[0][1] & QFileDialog.Option.DontUseNativeDialog
    assert bool(operations) is not cancel
    assert panel._candidate == candidate


def test_save_acknowledgement_blocks_native_dialog(window, tmp_path, monkeypatch):
    candidate = saved_candidate(tmp_path, flagged=True)
    panel = window._extraction_panel
    panel.show_candidate(candidate, read_candidate(candidate))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: pytest.fail("unacknowledged"))
    window._on_extraction_save_requested(candidate, False)
    assert panel._candidate == candidate


def configure_window(window, tmp_path):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("ID 00123")
    doc.save(path)
    window._on_ollama_connected(["qwen3:8b"], "")
    panel = window._extraction_panel
    panel.set_paths((path,))
    panel._request_edit.setPlainText("Extract the ID")
    panel.show_inspection(read_sources((path,)))
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    panel._confirm_btn.click()
    return panel, path


class ExtractionClient:
    def __init__(self, pause=False):
        self.calls = 0
        self.entered, self.release = Event(), Event()
        if not pause:
            self.release.set()

    def generate(self, model, prompt, system, **kwargs):
        self.calls += 1
        self.entered.set()
        assert self.release.wait(5)
        schema = kwargs["output_format"]
        if "columns" in schema["properties"]:
            return json.dumps({"columns": [{"label": "ID", "description": "Identifier",
                                            "kind": "text"}]})
        fields = schema["properties"]["records"]["items"]["properties"]["fields"]["properties"]
        return json.dumps({"records": [{"record_id": "source-1", "fields": {
            key: {"value": "00123", "quote": "ID 00123", "anchor": "source-1#p:0"}
            for key in fields}}]})


def test_complete_column_review_extract_save_and_confirmed_overwrite(window, tmp_path, qtbot, monkeypatch):
    panel, source = configure_window(window, tmp_path)
    original = source.read_bytes()
    client = ExtractionClient()
    window._ollama_client = client
    panel._suggest_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert panel.current_columns()[0].label == "ID" and not panel._confirmed
    panel._confirm_btn.click()
    panel._extract_btn.click()
    window._start_extraction("extract", panel.snapshot())  # duplicate ignored
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    candidate = panel._candidate
    assert candidate and panel._grid.item(0, 1).text() == "00123"
    assert "ID 00123" in panel._evidence.toPlainText()
    assert client.calls == 2
    destination = tmp_path / "output.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
    for existing in (False, True):
        if existing:
            destination.write_bytes(b"previous output")
        panel._save_btn.click()
        qtbot.waitUntil(lambda: window._extraction_thread is None)
        assert destination.read_bytes() == candidate.path.read_bytes()
        assert "saved successfully" in panel._status.text()
    assert source.read_bytes() == original
    assert panel._candidate == candidate


def test_save_failure_retains_review_and_previous_destination(window, tmp_path, qtbot, monkeypatch):
    from extraction import jobs
    panel, _ = configure_window(window, tmp_path)
    window._ollama_client = ExtractionClient()
    panel._extract_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    candidate = panel._candidate
    destination = tmp_path / "output.xlsx"
    destination.write_bytes(b"previous bytes")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
    def fail(*args):
        raise PermissionError("private path should not appear")
    monkeypatch.setattr(jobs.os, "replace", fail)
    panel._save_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert destination.read_bytes() == b"previous bytes"
    assert panel._candidate == candidate and panel._save_btn.isEnabled()
    assert "Could not save" in panel._status.text() and "private" not in panel._status.text()


@pytest.mark.parametrize("cancel", [False, True])
def test_rerun_invalidates_old_candidate_and_late_cancel_cannot_restore_it(window, tmp_path, qtbot, cancel):
    panel, _ = configure_window(window, tmp_path)
    window._ollama_client = ExtractionClient()
    panel._extract_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    old = panel._candidate
    client = ExtractionClient(pause=True)
    window._ollama_client = client
    panel._extract_btn.click()
    assert panel._candidate is None and not old.path.exists()
    qtbot.waitUntil(client.entered.is_set)
    if cancel:
        panel._cancel_btn.click()
    client.release.set()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    if cancel:
        assert panel._candidate is None and window._extraction_job is None
    else:
        assert panel._candidate.path != old.path
    assert panel._request_edit.toPlainText() == "Extract the ID"


def test_close_waits_for_cancelled_worker_before_removing_job(window, tmp_path, qtbot):
    panel, _ = configure_window(window, tmp_path)
    client = ExtractionClient(pause=True)
    window._ollama_client = client
    panel._extract_btn.click()
    qtbot.waitUntil(client.entered.is_set)
    job = window._extraction_job
    assert not window.close()
    assert job.exists() and window._extraction_thread.isRunning()
    client.release.set()
    qtbot.waitUntil(lambda: not window.isVisible(), timeout=5000)
    assert not job.exists()


def test_close_waits_for_queued_finished_handler_even_after_thread_stops(window, tmp_path, qtbot):
    from PySide6.QtGui import QCloseEvent
    panel, _ = configure_window(window, tmp_path)
    client = ExtractionClient()
    window._ollama_client = client
    panel._extract_btn.click()
    job = window._extraction_job
    assert window._extraction_thread.wait(5000)  # GUI result/finished handlers still queued
    event = QCloseEvent()
    window.closeEvent(event)
    assert not event.isAccepted()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert not job.exists()


def test_cancel_after_candidate_signal_clears_review_before_thread_finishes(window, tmp_path, qtbot, monkeypatch):
    from app.workers.extraction_worker import ExtractionWorker
    entered, release = Event(), Event()
    def paused_worker(*args, **kwargs):
        worker = ExtractionWorker(*args, **kwargs)
        def pause_finish():
            entered.set()
            assert release.wait(5)
        worker.finished.connect(pause_finish, Qt.ConnectionType.DirectConnection)
        return worker
    monkeypatch.setattr("app.main_window.ExtractionWorker", paused_worker)
    panel, _ = configure_window(window, tmp_path)
    window._ollama_client = ExtractionClient()
    panel._extract_btn.click()
    try:
        qtbot.waitUntil(lambda: panel._candidate is not None and entered.is_set())
        panel._cancel_btn.click()
    finally:
        release.set()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert panel._candidate is None and panel._grid.rowCount() == 0
    assert not panel._save_btn.isEnabled()


def test_prompt_generation_cannot_overlap_an_extraction(window, tmp_path, qtbot, monkeypatch):
    panel, _ = configure_window(window, tmp_path)
    client = ExtractionClient(pause=True)
    window._ollama_client = client
    panel._extract_btn.click()
    qtbot.waitUntil(client.entered.is_set)
    monkeypatch.setattr("app.main_window.GenerateWorker", lambda *a, **k:
                        pytest.fail("Prompt generation overlapped extraction"))
    try:
        window._on_generate_requested("Word (.docx)", str(tmp_path / "outputs"), "Write a report")
    finally:
        client.release.set()
        qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert not (tmp_path / "outputs").exists()


def test_source_path_refusal_and_wrong_suffix_keep_review(window, tmp_path, monkeypatch):
    candidate = saved_candidate(tmp_path, flagged=False)
    panel = window._extraction_panel
    panel.show_candidate(candidate, read_candidate(candidate))
    for destination, message in ((candidate.source_paths[0], "source file"),
                                  (tmp_path / "output.csv", ".xlsx")):
        monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
        panel._save_btn.click()
        assert message in panel._status.text()
        assert panel._candidate == candidate and window._extraction_thread is None


def test_source_drop_accepts_only_local_supported_files(panel):
    from PySide6.QtCore import QMimeData, QPoint, QPointF, QUrl
    from PySide6.QtGui import QDragEnterEvent, QDropEvent
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(str(panel._paths[0]))])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.DropAction.CopyAction, mime,
                            Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dragEnterEvent(enter)
    assert enter.isAccepted()
    drop = QDropEvent(QPointF(5, 5), Qt.DropAction.CopyAction, mime,
                      Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dropEvent(drop)
    assert panel._source_list.count() == 2
    mime.setUrls([QUrl("https://example.com/private.pdf")])
    enter = QDragEnterEvent(QPoint(5, 5), Qt.DropAction.CopyAction, mime,
                            Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier)
    panel.dragEnterEvent(enter)
    assert not enter.isAccepted()
