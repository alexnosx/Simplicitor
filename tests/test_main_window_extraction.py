"""Source workspace behavior, saved evidence, and native Save As delegation."""
from dataclasses import replace
from pathlib import Path
import json
from threading import Event

from docx import Document
from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QFileDialog, QTableWidget
import pytest

from app.config.settings import Settings
from extraction.models import ColumnSpec, ExtractionResult, FieldProposal, FieldResult, Issue
from extraction.source_readers import read_sources
from extraction.xlsx_writer import read_candidate, write_candidate


@pytest.fixture
def window(qtbot, tmp_path, monkeypatch):
    from app.main_window import MainWindow
    from app.services.ollama_client import OllamaClient
    def offline_start(self):
        self._ollama_client = OllamaClient("http://localhost:11434")
    monkeypatch.setattr(MainWindow, "_start_ollama_worker", offline_start)
    settings = Settings(tmp_path)
    settings.set("generated_dir", str(tmp_path / "outputs"))
    widget = MainWindow(settings)
    qtbot.addWidget(widget)
    yield widget
    widget.close()


def test_create_defaults_to_describe_with_optional_sources(window):
    from app.widgets.create_screen import CreateScreen
    assert isinstance(window._modes.currentWidget(), CreateScreen)
    assert window._create_screen._pages.currentWidget() is window._create_screen._describe_page
    assert not window._edit_panel.isVisible()
    assert not window._create_screen._cancel_btn.isEnabled()
    assert window.width() >= 1200 and window.height() >= 800


@pytest.mark.parametrize("count,warn", [(0, True), (7_999_999_999, True),
                                       (8_000_000_000, False), (27_000_000_000, False)])
def test_selected_model_warning(window, count, warn):
    window._on_ollama_connected(["selected", "loaded"], "loaded")
    window._top_bar._model_combo.setCurrentText("selected")
    window._on_model_params_ready("selected", count)
    assert window._capability_banner.isVisible() is warn
    assert window._capability_banner._info_label.text() == (
        "Simplicitor works best with models of 8B parameters or more.")
    window._create_screen._prompt_edit.setPlainText("Write a report")
    assert window._create_screen._primary_btn.isEnabled()
    window._on_model_params_ready("loaded", 1)
    window._on_ollama_connected(["selected", "loaded"], "loaded")
    assert window._current_model == "selected"
    assert window._top_bar.current_model() == "selected"
    assert window._capability_banner.isVisible() is warn


def test_banner_dismissal_is_per_selection_and_metadata_lookup_is_queued(window, qtbot):
    with qtbot.waitSignal(window._model_params_requested) as request:
        window._on_ollama_connected(["small", "other"], "small")
    assert request.args == ["small"]
    window._on_model_params_ready("small", 1)
    window._capability_banner._dismiss_btn.click()
    window._on_model_params_ready("small", 1)
    assert not window._capability_banner.isVisible()
    window._top_bar._model_combo.setCurrentText("other")
    window._on_model_params_ready("other", 0)
    assert window._capability_banner.isVisible()


def test_empty_model_discovery_clears_the_selection_and_warning(window):
    window._on_ollama_connected(["removed-model"], "")
    window._on_model_params_ready("removed-model", 1)
    assert window._capability_banner.isVisible()
    window._on_ollama_connected([], "")
    assert window._top_bar.current_model() == ""
    assert window._current_model == "" and window._create_screen._model == ""
    assert not window._capability_banner.isVisible()


@pytest.mark.parametrize("cancel", [False, True])
def test_native_save_as_xlsx_filter_confirmation_and_cancellation(window, tmp_path, monkeypatch, cancel):
    candidate = saved_candidate(tmp_path, flagged=False)
    panel = window._create_screen
    panel.set_route("pull")
    panel.show_candidate(candidate, read_candidate(candidate))
    destination = tmp_path / "chosen.xlsx"
    requests = []
    def dialog(parent, caption, directory, filter, **kwargs):
        requests.append((filter, kwargs.get("options", QFileDialog.Option(0))))
        return ("" if cancel else str(destination), "Excel workbook (*.xlsx)")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", dialog)
    operations = []
    monkeypatch.setattr(window, "_start_extraction", lambda action, setup, **kw:
                        operations.append((action, kw)))
    panel._save_btn.click()
    assert requests[0][0] == "Excel workbook (*.xlsx)"
    assert not requests[0][1] & QFileDialog.Option.DontConfirmOverwrite
    assert not requests[0][1] & QFileDialog.Option.DontUseNativeDialog
    assert bool(operations) is not cancel
    assert panel._candidate == candidate


def test_save_acknowledgement_blocks_native_dialog(window, tmp_path, monkeypatch):
    candidate = saved_candidate(tmp_path, flagged=True)
    panel = window._create_screen
    panel.set_route("pull")
    panel.show_candidate(candidate, read_candidate(candidate))
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: pytest.fail("unacknowledged"))
    window._on_extraction_save_requested(candidate, False)
    assert panel._candidate == candidate


def configure_window(window, tmp_path, qtbot):
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_paragraph("ID 00123")
    doc.save(path)
    window._on_ollama_connected(["qwen3:8b"], "")
    panel = window._create_screen
    panel.set_route("pull")
    panel.set_paths((path,))
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    panel._prompt_edit.setPlainText("Extract the ID")
    panel.show_inspection(read_sources((path,)))
    panel.set_columns((ColumnSpec("id", "ID", "Identifier"),))
    return panel, path


class ExtractionClient:
    def __init__(self, pause=False):
        self.calls = 0
        self.entered, self.release = Event(), Event()
        if not pause:
            self.release.set()

    def generate(self, model, prompt, system, **kwargs):
        self.calls += 1
        self.entered.set()
        assert self.release.wait(5)
        schema = kwargs["output_format"]
        if "columns" in schema["properties"]:
            return json.dumps({"columns": [{"label": "ID", "description": "Identifier",
                                            "kind": "text"}]})
        fields = schema["properties"]["records"]["items"]["properties"]["fields"]["properties"]
        return json.dumps({"records": [{"record_id": "source-1", "fields": {
            key: {"value": "00123", "quote": "ID 00123", "anchor": "source-1#p:0"}
            for key in fields}}]})


def test_complete_column_review_extract_save_and_confirmed_overwrite(window, tmp_path, qtbot, monkeypatch):
    panel, source = configure_window(window, tmp_path, qtbot)
    original = source.read_bytes()
    client = ExtractionClient()
    window._ollama_client = client
    panel._primary_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert panel.current_columns()[0].label == "ID" and not panel._confirmed
    panel._extract_btn.click()
    window._start_extraction("extract", panel.snapshot())  # duplicate ignored
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    candidate = panel._candidate
    assert candidate and panel._grid.item(0, 1).text() == "00123"
    assert "ID 00123" in panel._evidence.toPlainText()
    assert client.calls == 2
    destination = tmp_path / "output.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
    for existing in (False, True):
        if existing:
            destination.write_bytes(b"previous output")
        panel._pages.setCurrentWidget(panel._review_page)
        panel._save_btn.click()
        qtbot.waitUntil(lambda: window._extraction_thread is None)
        assert destination.read_bytes() == candidate.path.read_bytes()
        assert panel._pages.currentWidget() is panel._saved_page
    assert panel._saved_label.text() == "Saved: output.xlsx"
    assert source.read_bytes() == original
    assert panel._candidate == candidate


def test_save_failure_retains_review_and_previous_destination(window, tmp_path, qtbot, monkeypatch):
    from extraction import jobs
    panel, _ = configure_window(window, tmp_path, qtbot)
    window._ollama_client = ExtractionClient()
    panel._extract_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    candidate = panel._candidate
    destination = tmp_path / "output.xlsx"
    destination.write_bytes(b"previous bytes")
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
    def fail(*args):
        raise PermissionError("private path should not appear")
    monkeypatch.setattr(jobs.os, "replace", fail)
    panel._save_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert destination.read_bytes() == b"previous bytes"
    assert panel._candidate == candidate and panel._save_btn.isEnabled()
    assert "Could not save" in panel._status.text() and "private" not in panel._status.text()


@pytest.mark.parametrize("cancel", [False, True])
def test_rerun_invalidates_old_candidate_and_late_cancel_cannot_restore_it(window, tmp_path, qtbot, cancel):
    panel, _ = configure_window(window, tmp_path, qtbot)
    window._ollama_client = ExtractionClient()
    panel._extract_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    old = panel._candidate
    client = ExtractionClient(pause=True)
    window._ollama_client = client
    panel._extract_btn.click()
    assert panel._candidate is None and not old.path.exists()
    qtbot.waitUntil(client.entered.is_set)
    if cancel:
        panel._cancel_btn.click()
    client.release.set()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    if cancel:
        assert panel._candidate is None and window._extraction_job is None
    else:
        assert panel._candidate.path != old.path
    assert panel._prompt_edit.toPlainText() == "Extract the ID"


def test_close_waits_for_cancelled_worker_before_removing_job(window, tmp_path, qtbot):
    panel, _ = configure_window(window, tmp_path, qtbot)
    client = ExtractionClient(pause=True)
    window._ollama_client = client
    panel._extract_btn.click()
    qtbot.waitUntil(client.entered.is_set)
    job = window._extraction_job
    assert not window.close()
    assert job.exists() and window._extraction_thread.isRunning()
    client.release.set()
    qtbot.waitUntil(lambda: not window.isVisible(), timeout=5000)
    assert not job.exists()


def test_close_waits_for_queued_finished_handler_even_after_thread_stops(window, tmp_path, qtbot):
    from PySide6.QtGui import QCloseEvent
    panel, _ = configure_window(window, tmp_path, qtbot)
    client = ExtractionClient()
    window._ollama_client = client
    panel._extract_btn.click()
    job = window._extraction_job
    assert window._extraction_thread.wait(5000)  # GUI result/finished handlers still queued
    event = QCloseEvent()
    window.closeEvent(event)
    assert not event.isAccepted()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert not job.exists()


@pytest.mark.parametrize("cleanup_fails", [False, True])
def test_cancel_after_candidate_signal_clears_review_before_thread_finishes(
        window, tmp_path, qtbot, monkeypatch, cleanup_fails):
    from app.workers.extraction_worker import ExtractionWorker
    entered, release = Event(), Event()
    def paused_worker(*args, **kwargs):
        worker = ExtractionWorker(*args, **kwargs)
        def pause_finish():
            entered.set()
            assert release.wait(5)
        if worker.action == "extract":
            worker.finished.connect(pause_finish, Qt.ConnectionType.DirectConnection)
        return worker
    monkeypatch.setattr("app.main_window.ExtractionWorker", paused_worker)
    panel, _ = configure_window(window, tmp_path, qtbot)
    window._ollama_client = ExtractionClient()
    panel._extract_btn.click()
    with monkeypatch.context() as cleanup_patch:
        if cleanup_fails:
            def fail_cleanup(job):
                raise OSError("temporary workbook is open")
            cleanup_patch.setattr("app.main_window.discard_job", fail_cleanup)
        try:
            qtbot.waitUntil(lambda: panel._candidate is not None and entered.is_set())
            job = window._extraction_job
            panel._cancel_btn.click()
        finally:
            release.set()
        qtbot.waitUntil(lambda: window._extraction_thread is None)
        if cleanup_fails:
            assert window._extraction_job == job and job.exists()
            assert not panel._status.isHidden()
            assert panel._status.text().startswith("Could not clear the temporary job")
    assert panel._candidate is None and panel._grid.rowCount() == 0
    assert not panel._save_btn.isEnabled()


def test_prompt_generation_cannot_overlap_an_extraction(window, tmp_path, qtbot, monkeypatch):
    panel, _ = configure_window(window, tmp_path, qtbot)
    client = ExtractionClient(pause=True)
    window._ollama_client = client
    panel._extract_btn.click()
    qtbot.waitUntil(client.entered.is_set)
    monkeypatch.setattr("app.main_window.GenerateWorker", lambda *a, **k:
                        pytest.fail("Prompt generation overlapped extraction"))
    try:
        window._on_generate_requested("Word (.docx)", str(tmp_path / "outputs"), "Write a report")
    finally:
        client.release.set()
        qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert not (tmp_path / "outputs").exists()


def test_source_path_refusal_and_wrong_suffix_keep_review(window, tmp_path, monkeypatch):
    candidate = saved_candidate(tmp_path, flagged=False)
    panel = window._create_screen
    panel.set_route("pull")
    panel.show_candidate(candidate, read_candidate(candidate))
    for destination, message in ((candidate.source_paths[0], "source file"),
                                  (tmp_path / "output.csv", ".xlsx")):
        monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(destination), ""))
        panel._save_btn.click()
        assert message in panel._status.text()
        assert panel._candidate == candidate and window._extraction_thread is None




def saved_candidate(tmp_path, *, flagged=True, coverage=False):
    issue = Issue("request_failed", "source-1", "source-1#p:1", "A section failed.")
    field = FieldResult(FieldProposal("000452", "ID 000452", "source-1#p:0"),
                        "000452", flagged, ("leading_zero",) if flagged else ())
    result = ExtractionResult(("source-1",), {("source-1", "id"): field},
                              {"source-1": tmp_path / "source.docx"},
                              (issue,) if coverage else (), {})
    return write_candidate(result, (ColumnSpec("id", "ID", "Identifier"),),
                           tmp_path / "candidate.xlsx")


@pytest.mark.parametrize("kind", ["extraction", "ollama"])
def test_thread_cleanup_keeps_gui_responsive_until_native_thread_exits(
        window, qtbot, kind):
    from types import SimpleNamespace
    from PySide6.QtCore import QThread, QTimer

    entered, release = Event(), Event()
    thread = QThread(window)
    attribute = f"_{kind}_thread"
    setattr(window, attribute, thread)
    setattr(window, f"_{kind}_worker", SimpleNamespace(action="inspect"))
    handler = getattr(window, f"_on_{kind}_thread_finished")
    thread.finished.connect(handler)

    def finishing():
        entered.set()
        release.wait(0.5)  # Bounded fallback makes the broken implementation fail, not hang.

    thread.finished.connect(finishing, Qt.ConnectionType.DirectConnection)
    thread.started.connect(thread.quit, Qt.ConnectionType.DirectConnection)
    thread.start()
    assert entered.wait(2)  # Its GUI cleanup is queued before the heartbeat below.
    heartbeat = []

    def acknowledge_cleanup():
        heartbeat.append(getattr(window, attribute) is thread)
        release.set()

    QTimer.singleShot(0, acknowledge_cleanup)
    try:
        qtbot.waitUntil(lambda: getattr(window, attribute) is None and bool(heartbeat),
                        timeout=5000)
    finally:
        release.set()
        thread.wait(2000) if getattr(window, attribute) is not None else None
    assert heartbeat == [True]


def test_finished_steps_clear_their_progress_line(window, tmp_path, qtbot):
    panel, _ = configure_window(window, tmp_path, qtbot)
    assert panel._status.isHidden()  # "Reading sources..." is gone once reading ends
    window._ollama_client = ExtractionClient()
    panel._primary_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert panel._pages.currentWidget() is panel._columns_page
    assert panel._status.isHidden()  # so is "Suggesting columns..."


def test_failure_message_outlives_its_finished_step(window, tmp_path, qtbot):
    panel, _ = configure_window(window, tmp_path, qtbot)

    class Failing(ExtractionClient):
        def generate(self, *args, **kwargs):
            raise RuntimeError("model crashed")

    window._ollama_client = Failing()
    panel._primary_btn.click()
    qtbot.waitUntil(lambda: window._extraction_thread is None)
    assert not panel._status.isHidden()
    assert panel._status.text().startswith("Could not complete the request")
