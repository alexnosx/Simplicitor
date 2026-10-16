"""Limited display of a Word document: style, effective bold/italic, lists, tables, headers."""
from dataclasses import dataclass, replace
from difflib import SequenceMatcher
import io
import re

from docx import Document
from docx.oxml.ns import qn
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph

from editing.models import EditCandidate, OpenedDocument

_ON = {"true", "1", "on"}


@dataclass(frozen=True)
class Span:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass(frozen=True)
class DisplayBlock:
    """kind: paragraph, table, other, header or footer. index is the body paragraph index."""
    kind: str
    index: int = -1
    style: str = ""
    heading: int = 0
    list_item: bool = False
    spans: tuple[Span, ...] = ()
    cells: tuple[tuple[str, ...], ...] = ()
    reason: str | None = None

    @property
    def text(self) -> str:
        return "".join(s.text for s in self.spans)


@dataclass(frozen=True)
class DeletionReview:
    """Display-only removed text ranges per Before paragraph and removed/merged count."""
    ranges: tuple[tuple[tuple[int, int], ...], ...]
    removed_paragraphs: int


def review_deletions(before: list[str], after: list[str],
                     sources: tuple[int, ...] | None = None) -> DeletionReview:
    """Compare the whole selection, preserving words carried across paragraph boundaries.

    Virtual spaces make a split or merge equivalent to an ordinary space for the word diff.
    Paragraph matching counts removals separately; it does not judge the model's meaning.
    Validated proposal sources disambiguate identical paragraphs with different formatting.
    Ranges refer to the original strings, never to modified document bytes.
    """
    def tokens(paragraphs):
        values, positions = [], []
        for row, text in enumerate(paragraphs):
            if row:
                values.append(" ")
                positions.append(None)
            for match in re.finditer(r"\w+|\s+|[^\w\s]", text):
                values.append(match.group())
                positions.append((row, match.start(), match.end()))
        return values, positions

    ranges: list[list[tuple[int, int]]] = [[] for _ in before]
    removed = 0
    old_keys = list(enumerate(before)) if sources is not None else before
    new_keys = list(zip(sources, after)) if sources is not None else after
    paragraphs = SequenceMatcher(None, old_keys, new_keys, autojunk=False)
    for operation, a1, a2, b1, b2 in paragraphs.get_opcodes():
        if operation in ("equal", "insert"):
            continue
        removed += max(0, (a2 - a1) - (b2 - b1))
        if operation == "delete":
            for row in range(a1, a2):
                if before[row]:
                    ranges[row].append((0, len(before[row])))
            continue
        old, positions = tokens(before[a1:a2])
        new, _ = tokens(after[b1:b2])
        words = SequenceMatcher(None, old, new, autojunk=False)
        for word_operation, first, last, _, _ in words.get_opcodes():
            if word_operation not in ("delete", "replace"):
                continue
            for position in positions[first:last]:
                if position is None:
                    continue
                row, start, end = position
                row += a1
                if ranges[row] and ranges[row][-1][1] == start:
                    ranges[row][-1] = (ranges[row][-1][0], end)
                else:
                    ranges[row].append((start, end))
    return DeletionReview(tuple(tuple(row) for row in ranges), removed)


def with_review_deletions(candidate: EditCandidate, before: tuple[DisplayBlock, ...],
                         after: tuple[DisplayBlock, ...],
                         sources: tuple[int, ...] | None = None) -> EditCandidate:
    """Attach display markers off the GUI thread, preserving the verified document bytes."""
    deleted = review_deletions([b.text for b in before], [b.text for b in after], sources)
    return replace(candidate, before_deletions=deleted.ranges,
                   removed_paragraphs=deleted.removed_paragraphs)


def _on_off(rpr, tag: str) -> bool | None:
    """Direct on/off value of w:<tag> in an rPr element, or None when it is not set."""
    if rpr is None:
        return None
    element = rpr.find(qn("w:" + tag))
    if element is None:
        return None
    value = element.get(qn("w:val"))
    return value is None or value.lower() in _ON


def _chain(style, tag: str) -> bool | None:
    """Nearest value in a style's basedOn chain (ordinary inheritance), or None."""
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        value = _on_off(style.element.rPr, tag)
        if value is not None:
            return value
        style = style.base_style
    return None


def _effective(run, paragraph, defaults, tag: str) -> bool:
    """Direct value wins; else paragraph-style XOR character-style value; else defaults."""
    direct = _on_off(run._r.rPr, tag)
    if direct is not None:
        return direct
    paragraph_value = _chain(paragraph.style, tag)
    character_value = _chain(run.style, tag)
    if paragraph_value is None and character_value is None:
        return bool(_on_off(defaults, tag))
    return bool(paragraph_value) ^ bool(character_value)


def _list_item(paragraph) -> bool:
    """First present numId in direct properties, then the style chain; 0 cancels."""
    values = paragraph._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
    style, seen = paragraph.style, set()
    while not values and style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        values = style.element.xpath("./w:pPr/w:numPr/w:numId/@w:val")
        style = style.base_style
    return bool(values) and values[0] != "0"


def _heading(style_name: str) -> int:
    if style_name == "Title":
        return 1
    if style_name.startswith("Heading ") and style_name[8:].isdigit():
        return int(style_name[8:])
    return 0


def _paragraph(paragraph, index: int, defaults, reason: str | None = None) -> DisplayBlock:
    spans = []
    for item in paragraph.iter_inner_content():
        for run in (item.runs if isinstance(item, Hyperlink) else [item]):
            if run.text:
                spans.append(Span(run.text, _effective(run, paragraph, defaults, "b"),
                                  _effective(run, paragraph, defaults, "i")))
    if paragraph._p.xpath(".//w:drawing | .//w:pict"):
        spans.append(Span("[Picture]"))
    style = paragraph.style.name if paragraph.style is not None else "Normal"
    return DisplayBlock("paragraph", index, style, _heading(style), _list_item(paragraph),
                        tuple(spans), reason=reason)


def _defaults(document):
    found = document.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr")
    return found[0] if found else None


def _cells(table) -> tuple[tuple[str, ...], ...]:
    """Each physical cell's text, with nested tables flattened into it, read-only."""
    return tuple(tuple("\n".join(Paragraph(p, table).text for p in tc.iter(qn("w:p")))
                       for tc in row._tr.tc_lst) for row in table.rows)


def _container_text(container) -> str:
    """Paragraphs and tables of a header or footer, tables one row per line."""
    lines = []
    for block in container.iter_inner_content():
        if isinstance(block, Paragraph):
            lines.append(block.text)
        else:
            lines.extend("\t".join(row) for row in _cells(block))
    return "\n".join(lines).strip()


def build(opened: OpenedDocument) -> tuple[DisplayBlock, ...]:
    """Headers, then body blocks in document order, then footers."""
    document = Document(io.BytesIO(opened.data))
    defaults = _defaults(document)
    paragraphs, tables = document.paragraphs, iter(document.tables)
    body = []
    for block in opened.blocks:
        if block.kind == "p":
            body.append(_paragraph(paragraphs[block.index], block.index, defaults, block.reason))
        elif block.kind == "tbl":
            body.append(DisplayBlock("table", cells=_cells(next(tables)),
                                     reason="Tables are read-only in the first Word editor."))
        elif block.kind not in ("marker", "sectPr"):
            body.append(DisplayBlock("other", reason=block.reason))
    margins = {"header": [], "footer": []}
    seen = set()
    for section in document.sections:
        for kind in ("header", "footer"):
            for variant in (kind, "first_page_" + kind, "even_page_" + kind):
                part = getattr(section, variant)
                if part.is_linked_to_previous or part.part.partname in seen:
                    continue
                seen.add(part.part.partname)
                text = _container_text(part)
                if text:
                    margins[kind].append(DisplayBlock(kind, spans=(Span(text),),
                                                      reason=f"The {kind} is shown for context and cannot be edited."))
    return tuple(margins["header"] + body + margins["footer"])


def paragraphs(data: bytes, first: int, count: int) -> tuple[DisplayBlock, ...]:
    """Display blocks for body paragraphs first..first+count-1 of a package (review)."""
    document = Document(io.BytesIO(data))
    defaults = _defaults(document)
    found = document.paragraphs
    return tuple(_paragraph(found[i], i, defaults) for i in range(first, first + count))
