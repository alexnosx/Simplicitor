"""Anchored DOCX body/header/footer and PDF readers; never write to sources."""
from math import ceil
from pathlib import Path
from typing import Iterable, Sequence
from xml.etree import ElementTree

import pdfplumber
from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.config.defaults import (
    EXTRACTION_DOCX_CHARS_PER_PAGE,
    EXTRACTION_MAX_FILE_BYTES,
    EXTRACTION_MAX_JOB_PAGES,
    EXTRACTION_PDF_MIN_TEXT_CHARS,
)
from extraction.models import Issue, SourceDocument, SourceUnit

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class SourceReadError(ValueError):
    """A safe reader failure, optionally carrying page/structure issues."""

    def __init__(self, message: str, issues: tuple[Issue, ...] = ()) -> None:
        super().__init__(message)
        self.issues = issues


def read_source(path: Path, source_id: str) -> SourceDocument:
    """Read the complete supported source with stable, job-local anchors."""
    path = Path(path)
    try:
        if path.stat().st_size > EXTRACTION_MAX_FILE_BYTES:
            raise SourceReadError("Input exceeds the 50 MiB per-file limit.")
        if path.suffix.lower() == ".docx":
            document = _read_docx(path, source_id)
        elif path.suffix.lower() == ".pdf":
            document = _read_pdf(path, source_id)
        else:
            raise SourceReadError("Choose a DOCX or text-layer PDF file.")
        if document.page_cost > EXTRACTION_MAX_JOB_PAGES:
            raise SourceReadError("Source exceeds the 300-page job limit.")
        return document
    except SourceReadError:
        raise
    except Exception:
        raise SourceReadError("Could not read source. Check its format or password protection.") from None


def read_sources(paths: Sequence[Path]) -> tuple[SourceDocument, ...]:
    """Read attachments independently and enforce the aggregate page budget."""
    documents = []
    pages = 0
    for index, path in enumerate(paths, 1):
        document = read_source(path, f"source-{index}")
        pages += document.page_cost
        if pages > EXTRACTION_MAX_JOB_PAGES:
            raise SourceReadError("Attachments exceed the 300-page job limit.")
        documents.append(document)
    return tuple(documents)


def _read_docx(path: Path, source_id: str) -> SourceDocument:
    doc = Document(path)
    units, issues = [], []
    containers = [(source_id + "#", doc, doc.element.body)]
    seen_parts = set()
    counts = {"header": 0, "footer": 0}
    for section in doc.sections:
        for kind in ("header", "footer"):
            for variant in (kind, "first_page_" + kind, "even_page_" + kind):
                container = getattr(section, variant)
                if container.is_linked_to_previous:
                    continue
                part = container.part
                if part.partname in seen_parts:
                    continue
                seen_parts.add(part.partname)
                prefix = f"{source_id}#{kind}:{counts[kind]}:"
                counts[kind] += 1
                containers.append((prefix, container, part.element))
    for prefix, container, element in containers:
        _read_blocks(container.iter_inner_content(), prefix, source_id, units, issues)
        if any(list(element.iter(_W + tag))
               for tag in ("sdt", "ins", "del", "txbxContent", "altChunk", "fldSimple")):
            issues.append(Issue("unsupported_structure", source_id, prefix,
                                "Content controls, revisions, fields, or embedded text need review."))
    for part in doc.part.package.parts:
        if any(token in str(part.partname) for token in ("/footnotes", "/endnotes")):
            root = ElementTree.fromstring(part.blob)
            if any((n.text or "").strip() for n in root.iter(_W + "t")):
                issues.append(Issue("unsupported_structure", source_id, source_id,
                                    "Footnotes and endnotes are outside supported reading."))
                break
    count = sum(len(u.text) for u in units)
    if not any(u.text.strip() for u in units):
        raise SourceReadError("DOCX has no readable supported text.", tuple(issues))
    cost = max(1, ceil(count / EXTRACTION_DOCX_CHARS_PER_PAGE))
    return SourceDocument(source_id, path.resolve(), path.name, tuple(units), cost, tuple(issues))


def _read_blocks(
    blocks: Iterable[Paragraph | Table], prefix: str, source_id: str,
    units: list[SourceUnit], issues: list[Issue],
) -> None:
    paragraph_index = table_index = 0
    for block in blocks:
        if isinstance(block, Paragraph):
            anchor = f"{prefix}p:{paragraph_index}"
            units.append(SourceUnit(anchor, block.text, "paragraph", len(units), anchor))
            paragraph_index += 1
        elif isinstance(block, Table):
            for row_index, row in enumerate(block.rows):
                group = f"{prefix}t:{table_index}:r:{row_index}"
                for cell_index, tc in enumerate(row._tr.tc_lst):
                    cell = _Cell(tc, block)
                    text = "\n".join(p.text for p in cell.paragraphs)
                    anchor = f"{group}:c:{cell_index}"
                    units.append(SourceUnit(anchor, text, "cell", len(units), group))
                    if cell.tables:
                        issues.append(Issue("unsupported_structure", source_id, anchor,
                                            "Nested tables are outside supported reading."))
            table_index += 1


def _read_pdf(path: Path, source_id: str) -> SourceDocument:
    units, issues = [], []
    with pdfplumber.open(path) as pdf:
        if len(pdf.pages) > EXTRACTION_MAX_JOB_PAGES:
            raise SourceReadError("Source exceeds the 300-page job limit.")
        for number, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            anchor = f"{source_id}#page:{number}"
            units.append(SourceUnit(anchor, text, "page", number - 1, anchor))
            if sum(not ch.isspace() for ch in text) < EXTRACTION_PDF_MIN_TEXT_CHARS:
                issues.append(Issue("unreadable_page", source_id, anchor,
                                    "Page has too little extracted text; coverage is incomplete."))
    if not units or len(issues) == len(units):
        raise SourceReadError("PDF has no usable text layer.", tuple(issues))
    return SourceDocument(source_id, path.resolve(), path.name, tuple(units),
                          len(units), tuple(issues))
