"""Typed/literal Data and Evidence output, verified after saving and used for review."""
from collections import defaultdict
from copy import copy
from dataclasses import replace
from datetime import date, datetime
from decimal import Decimal
import math
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config.defaults import (
    APP_FONT_FAMILY, BODY_TEXT_COLOR, EXTRACTION_FLAG_COLOR, FONT_SIZE_BODY_PT,
    PRIMARY_ACCENT_COLOR, WHITE,
)
from extraction.jobs import same_path
from extraction.models import Candidate, ColumnSpec, ExtractionResult, Issue, ReviewCell

_EVIDENCE_HEADERS = ("Data cell", "Record", "File", "Field", "Value", "Quote", "Anchor",
                     "Status", "Issue")
_FLAG = PatternFill("solid", fgColor=EXTRACTION_FLAG_COLOR.lstrip("#"))


def _text(value: str) -> None:
    if (any(not (c in "\t\n\r" or 0x20 <= ord(c) <= 0xD7FF
                 or 0xE000 <= ord(c) <= 0xFFFD or 0x10000 <= ord(c) <= 0x10FFFF)
            for c in value)
            or len(value.encode("utf-16-le")) // 2 > 32767):
        raise ValueError("Excel cannot represent this text. Shorten or correct the field/evidence.")


def _numeric(value: int | Decimal) -> int | float:
    number = Decimal(value)
    if not number.is_finite():
        raise ValueError("Excel cannot represent this non-finite number.")
    normalized = number.normalize()
    floating = float(number)
    if (len(normalized.as_tuple().digits) > 15 or not math.isfinite(floating)
            or (number and abs(floating) < 2.2250738585072014e-308)
            or Decimal(str(floating)) != number):
        raise ValueError("Excel cannot represent this number without losing precision.")
    return value if isinstance(value, int) else floating


def _put(sheet, row: int, column: int, value, *, flagged: bool = False,
         number_format: str = "General") -> None:
    if isinstance(value, str):
        _text(value)
    elif isinstance(value, (int, Decimal)) and not isinstance(value, bool):
        value = _numeric(value)
    elif value is not None and not isinstance(value, date):
        raise ValueError("Excel field has an unsupported value type.")
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"  # Never infer formulas, even in quotes, labels, or file names.
    cell.number_format = number_format
    cell.font = Font(name=APP_FONT_FAMILY, size=FONT_SIZE_BODY_PT,
                     color=BODY_TEXT_COLOR.lstrip("#"))
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if flagged:
        cell.fill = _FLAG


def _style(sheet) -> None:
    sheet.freeze_panes = "C2" if sheet.title == "Data" else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=PRIMARY_ACCENT_COLOR.lstrip("#"))
        cell.font = Font(name=APP_FONT_FAMILY, size=FONT_SIZE_BODY_PT,
                         bold=True, color=WHITE.lstrip("#"))
    for column in sheet.columns:
        width = min(60, max(14, max(len(str(c.value or "")) for c in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def _verify_saved(workbook: Workbook, path: Path) -> None:
    with path.open("rb") as stream:
        saved = load_workbook(stream, data_only=False)
    try:
        for expected_sheet in workbook:
            actual_sheet = saved[expected_sheet.title]
            if actual_sheet.max_row != expected_sheet.max_row:
                raise ValueError("Excel saved workbook lost records or evidence.")
            for row in expected_sheet:
                for cell in row:
                    actual = actual_sheet[cell.coordinate]
                    expected = cell.value
                    if isinstance(expected, date) and not isinstance(expected, datetime):
                        expected = datetime.combine(expected, datetime.min.time())
                    if expected == "":
                        expected = None
                    if (actual.value != expected or actual.number_format != cell.number_format
                            or actual.fill != copy(cell.fill)
                            or (expected is not None and actual.data_type != cell.data_type)
                            or actual.hyperlink is not None):
                        raise ValueError("Excel cannot preserve this workbook value or evidence.")
    finally:
        saved.close()


def write_candidate(result: ExtractionResult, columns: tuple[ColumnSpec, ...],
                    path: Path) -> Candidate:
    """Save/reopen one candidate; preserve imprecise numbers as flagged literals.

    Replacing a candidate invalidates the previous run immediately, even if this
    write fails. Callers clear their review before each run; cancelled extraction
    discards its job without calling this writer.
    """
    path = Path(path)
    sources = tuple(result.source_paths[source] for source in result.ordered_source_ids)
    if any(same_path(path, source) for source in sources):
        raise ValueError("A candidate cannot replace a source file.")
    if path.name != "candidate.xlsx":
        raise ValueError("Use candidate.xlsx in the job directory.")
    path.unlink(missing_ok=True)
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    evidence = workbook.create_sheet("Evidence")
    temporary = None
    issues = list(result.issues)
    # Coverage still appears if a caller forgot to supply a corresponding issue.
    for anchor, status in result.coverage.items():
        if status != "processed" and not any(i.anchor == anchor for i in issues):
            issues.append(Issue("coverage_" + status, anchor.split("#", 1)[0], anchor,
                                "A source unit was not processed."))
    coverage_issues = tuple(issues)
    try:
        if not columns or len(columns) > 16382 or len(sources) > 1048575:
            raise ValueError("Excel output requires columns and must fit worksheet dimensions.")
        if len({c.id for c in columns}) != len(columns):
            raise ValueError("Excel columns must have unique identifiers.")
        for col, label in enumerate(("Record", "File", *(c.label for c in columns)), 1):
            _put(data, 1, col, label)
        for col, label in enumerate(_EVIDENCE_HEADERS, 1):
            _put(evidence, 1, col, label)

        def evidence_row(values, flagged=False):
            row = evidence.max_row + 1
            if row > 1048576:
                raise ValueError("Excel evidence exceeds the worksheet row limit.")
            for col, value in enumerate(values, 1):
                _put(evidence, row, col, value, flagged=flagged)

        for row, source in enumerate(result.ordered_source_ids, 2):
            label = result.source_paths[source].name
            _put(data, row, 1, source)
            _put(data, row, 2, label)
            for col, spec in enumerate(columns, 3):
                field = result.fields.get((source, spec.id))
                if field is None:
                    raise ValueError("Excel output is missing a requested field.")
                value = field.data_value
                if isinstance(value, (int, Decimal)):
                    try:
                        _numeric(value)
                    except (ValueError, ArithmeticError):
                        field = replace(field, typed_value=None, flagged=True,
                                        issues=field.issues + ("excel_precision",))
                        value = field.data_value
                format_string = "General"
                if not field.flagged and value is not None:
                    if spec.kind == "date":
                        format_string = "yyyy-mm-dd"
                    elif spec.kind == "integer":
                        format_string = "0"
                    elif spec.kind == "decimal":
                        places = max(2, -value.as_tuple().exponent) if isinstance(value, Decimal) else 2
                        format_string = ("0." + "0" * places if places <= 15
                                         else "0.##############E+00")
                _put(data, row, col, value, flagged=field.flagged, number_format=format_string)
                reference = f"{get_column_letter(col)}{row}"
                proposal = field.proposal
                evidence_row((reference, source, label, spec.label, proposal.value,
                              proposal.quote, proposal.anchor,
                              "flagged" if field.flagged else "verified", "; ".join(field.issues)),
                             field.flagged)
                if field.flagged:
                    code = "excel_precision" if "excel_precision" in field.issues else "flagged_field"
                    issues.append(Issue(code, source, reference,
                                        "A field needs review before saving."))
                for alternative in field.alternatives:
                    evidence_row((reference, source, label, spec.label, alternative.value,
                                  alternative.quote, alternative.anchor, "alternative", ""),
                                 field.flagged)
            for issue in coverage_issues:
                if issue.source_id in (source, ""):
                    evidence_row(("", source, label, "", None, "", issue.anchor, "coverage",
                                  issue.code + ": " + issue.safe_message), True)
        _style(data)
        _style(evidence)
        descriptor, name = tempfile.mkstemp(prefix=".candidate-", suffix=".tmp", dir=path.parent)
        os.close(descriptor)
        temporary = Path(name)
        workbook.save(temporary)
        _verify_saved(workbook, temporary)
        os.replace(temporary, path)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
    return Candidate(path.parent.name, path, sources, tuple(issues))


def read_candidate(candidate: Candidate) -> tuple[ReviewCell, ...]:
    """Build review entirely from saved Data/Evidence, including coverage issues."""
    workbook = load_workbook(candidate.path, data_only=False)
    try:
        data, evidence = workbook["Data"], workbook["Evidence"]
        by_cell, by_record = defaultdict(list), defaultdict(list)
        for row in evidence.iter_rows(min_row=2, values_only=True):
            entry = dict(zip(_EVIDENCE_HEADERS, row))
            if entry["Status"] == "coverage":
                by_record[entry["Record"]].append(entry)
            else:
                by_cell[entry["Data cell"]].append(entry)
        return tuple(ReviewCell(
            cell.row, cell.column, cell.value, cell.data_type,
            tuple(by_cell[cell.coordinate] + by_record[data.cell(cell.row, 1).value]),
        ) for row in data.iter_rows(min_row=2, min_col=3) for cell in row)
    finally:
        workbook.close()
