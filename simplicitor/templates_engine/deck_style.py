# simplicitor/templates_engine/deck_style.py
"""Read the look of any .pptx deck, best effort and read-only.

Every slide is read against its own layout, master and theme. Text styles follow
PowerPoint's inheritance: the run, the paragraph's defaults, the shape's own list
style, then (for placeholders) the layout and master placeholders and the master's
title, body or other style, then the theme. Within each list style a level's
properties come before the style's default paragraph. Colors go through the
slide's color map and PowerPoint's lighter and darker variants. A slide's
background is its own, else its layout's, else its master's.
"""
import colorsys
import io
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

_STANDARD_MAP = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}
_TITLES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_OTHERS = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER,
           PP_PLACEHOLDER.HEADER}
# Linked, missing, damaged, unsupported or oversized images are skipped, never fatal.
_IMAGE_ERRORS = (AttributeError, KeyError, ValueError, OSError, Image.DecompressionBombError)


@dataclass(frozen=True)
class TextStyle:
    font: str | None = None
    size_pt: float | None = None
    color: str | None = None  # "RRGGBB"
    bold: bool | None = None


@dataclass(frozen=True)
class Fill:
    color: str | None = None  # "RRGGBB"
    image: bytes | None = None
    image_ext: str = ""


# A background style giving the theme's light color, which a new template shows anyway.
_PLAIN = Fill()


@dataclass(frozen=True)
class Logo:
    image: bytes
    image_ext: str
    left: float  # fractions of the slide width and height
    top: float
    width: float
    height: float


@dataclass(frozen=True)
class DeckStyle:
    slide_width: int
    slide_height: int
    color_scheme: bytes | None  # the deck theme's <a:clrScheme> element
    title: TextStyle
    body: TextStyle
    title_background: Fill | None  # None: the slide follows the template's master
    background: Fill | None
    logo: Logo | None


@dataclass(frozen=True)
class _Theme:
    slots: dict  # color slot (dk1, lt1, accent1...) to "RRGGBB"
    scheme: bytes | None
    major: str | None
    minor: str | None


@dataclass(frozen=True)
class _Run:
    style: TextStyle  # the effective style
    weight: int  # characters
    on_slide: bool  # its font or color is set on the slide and differs from its layout's


@dataclass(frozen=True)
class _Deck:
    width: int
    height: int
    colors: dict  # the colors of the master most slides use
    theme: _Theme  # that master's theme
    title_runs: list
    body_runs: list
    fills: list  # each slide's own background, or None
    backgrounds: list  # each slide's effective background; None is the plain default
    logo: Logo | None


def _open(path: str | Path):
    from templates_engine.breakdown import _open_presentation
    return _open_presentation(Path(path))


def _theme(master) -> _Theme:
    """A master's theme: colors by slot name, the raw color scheme, the two fonts."""
    part = master.part.part_related_by(RT.THEME)
    elements = etree.fromstring(part.blob).find(qn("a:themeElements"))
    slots: dict[str, str] = {}
    scheme = elements.find(qn("a:clrScheme")) if elements is not None else None
    if scheme is not None:
        for slot in scheme:
            if len(slot):
                value = slot[0].get("lastClr") or slot[0].get("val")
                if value:
                    slots[etree.QName(slot).localname] = value.upper()
    fonts = elements.find(qn("a:fontScheme")) if elements is not None else None

    def typeface(kind: str) -> str | None:
        latin = fonts.find(f"{qn(kind)}/{qn('a:latin')}") if fonts is not None else None
        return (latin.get("typeface") or None) if latin is not None else None

    raw = etree.tostring(scheme) if scheme is not None else None
    return _Theme(slots, raw, typeface("a:majorFont"), typeface("a:minorFont"))


def _color_map(*owners):
    """The color map in force: the first override among *owners*, else the last one's map."""
    for owner in owners[:-1]:
        override = owner._element.find(f"{qn('p:clrMapOvr')}/{qn('a:overrideClrMapping')}")
        if override is not None:
            return override
    return owners[-1]._element.find(qn("p:clrMap"))


def _colors(theme: _Theme, mapping) -> dict[str, str]:
    """The theme's slots plus bg1, tx1, bg2, tx2 and any other name the color map gives."""
    names = {**_STANDARD_MAP, **(dict(mapping.attrib) if mapping is not None else {})}
    return {**theme.slots,
            **{name: theme.slots[slot] for name, slot in names.items() if slot in theme.slots}}


def _fraction(color, tag: str) -> float | None:
    element = color.find(qn(tag))
    value = element.get("val", "") if element is not None else ""
    return int(value) / 100000 if value.isdigit() else None


def _variant(value: str, color) -> str:
    """Apply PowerPoint's lighter and darker variants (lumMod, lumOff) to "RRGGBB"."""
    scale, shift = _fraction(color, "a:lumMod"), _fraction(color, "a:lumOff")
    if scale is None and shift is None:
        return value
    hue, light, saturation = colorsys.rgb_to_hls(
        *(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4)))
    light = min(max(light * (1 if scale is None else scale) + (shift or 0), 0), 1)
    return "".join(f"{round(c * 255):02X}" for c in colorsys.hls_to_rgb(hue, light, saturation))


def _color(element, colors: dict[str, str]) -> str | None:
    if element is None:
        return None
    srgb = element.find(qn("a:srgbClr"))
    if srgb is not None and srgb.get("val"):
        return _variant(srgb.get("val").upper(), srgb)
    scheme = element.find(qn("a:schemeClr"))
    if scheme is not None:
        value = colors.get(scheme.get("val", ""))
        return _variant(value, scheme) if value else None
    system = element.find(qn("a:sysClr"))
    if system is not None and system.get("lastClr"):
        return _variant(system.get("lastClr").upper(), system)
    return None


def _first(*styles: TextStyle) -> TextStyle:
    """Each property from the first style that sets it."""
    def pick(name: str):
        return next((getattr(s, name) for s in styles if getattr(s, name) is not None), None)
    return TextStyle(pick("font"), pick("size_pt"), pick("color"), pick("bold"))


def _props(element, colors: dict[str, str], major: str | None, minor: str | None) -> TextStyle:
    """What one a:rPr or a:defRPr element sets; unset properties are None."""
    if element is None:
        return TextStyle()
    latin = element.find(qn("a:latin"))
    font = latin.get("typeface") if latin is not None else None
    font = {"+mj-lt": major, "+mn-lt": minor}.get(font, font) or None
    size = element.get("sz")
    size = int(size) / 100 if size and size.isdigit() else None
    bold = {"1": True, "true": True, "0": False, "false": False}.get(element.get("b"))
    return TextStyle(font, size, _color(element.find(qn("a:solidFill")), colors), bold)


def _level(styles, level: int, colors, major, minor) -> TextStyle:
    """One paragraph level of a list or text style, over that style's default paragraph."""
    if styles is None:
        return TextStyle()
    return _first(*(_props(styles.find(f"{qn(tag)}/{qn('a:defRPr')}"), colors, major, minor)
                    for tag in (f"a:lvl{level}pPr", "a:defPPr")))


def _list_style(shape):
    return shape._element.find(f"{qn('p:txBody')}/{qn('a:lstStyle')}")


def _inherited_styles(shape, master, presentation) -> list:
    """Style elements the shape's text inherits, nearest first, excluding the shape itself."""
    if not shape.is_placeholder:
        return [presentation.find(qn("p:defaultTextStyle"))]
    sources, base = [], shape
    for _ in range(2):  # the layout placeholder, then the master placeholder
        try:
            base = base._base_placeholder
        except (AttributeError, KeyError):
            base = None
        if base is None:
            break
        sources.append(_list_style(base))
    kind = shape.placeholder_format.type
    name = ("p:titleStyle" if kind in _TITLES
            else "p:otherStyle" if kind in _OTHERS else "p:bodyStyle")
    sources.append(master._element.find(f"{qn('p:txStyles')}/{qn(name)}"))
    return sources


def _runs(shape, master, presentation, colors, major, minor) -> list[_Run]:
    """The effective style of every non-empty run, and whether the slide overrides it."""
    title = shape.is_placeholder and shape.placeholder_format.type in _TITLES
    theme = TextStyle(major if title else minor, None, colors.get("tx1"), None)
    sources = _inherited_styles(shape, master, presentation)
    own_styles = _list_style(shape)
    result = []
    for paragraph in shape.text_frame.paragraphs:
        level = min(paragraph.level + 1, 9)
        inherited = _first(*(_level(source, level, colors, major, minor) for source in sources),
                           theme)
        defaults = _first(
            _props(paragraph._p.find(f"{qn('a:pPr')}/{qn('a:defRPr')}"), colors, major, minor),
            _level(own_styles, level, colors, major, minor))
        for run in paragraph.runs:
            text = run.text.strip()
            if not text:
                continue
            own = _first(_props(run._r.find(qn("a:rPr")), colors, major, minor), defaults)
            on_slide = own.font not in (None, inherited.font) or own.color not in (
                None, inherited.color)
            result.append(_Run(_first(own, inherited), len(text), on_slide))
    return result


def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _title_shape(shapes: list, height: int):
    """The title placeholder, else the shape near the top with the largest text."""
    for shape, _ in shapes:
        if shape.is_placeholder and shape.placeholder_format.type in _TITLES:
            return shape
    top = [(shape, runs) for shape, runs in shapes
           if shape.top is not None and shape.top < height * 0.4]
    if not top:
        return None
    return max(top, key=lambda item: (max((r.style.size_pt or 0 for r in item[1]), default=0),
                                      -item[0].top))[0]


def _weighted_median(values: list[tuple[float, int]]) -> float | None:
    values = sorted(values)
    total = sum(weight for _, weight in values)
    running = 0
    for value, weight in values:
        running += weight
        if running * 2 >= total:
            return value
    return None


def _dominant(runs: list[_Run], font: str | None, color: str | None) -> TextStyle:
    """The most used font, color and boldness, and the median size, weighted by characters."""
    def most(attribute: str):
        counts: Counter = Counter()
        for run in runs:
            value = getattr(run.style, attribute)
            if value is not None:
                counts[value] += run.weight
        return counts.most_common(1)[0][0] if counts else None

    size = _weighted_median([(r.style.size_pt, r.weight) for r in runs if r.style.size_pt])
    return TextStyle(most("font") or font, size, most("color") or color, most("bold"))


def _readable(blob: bytes) -> bool:
    """True when the whole image decodes; opening an image reads only its header."""
    try:
        with Image.open(io.BytesIO(blob)) as image:
            image.load()
    except _IMAGE_ERRORS:
        return False
    return True


def _picture(shape):
    """(blob, extension, SHA-1) of a picture's embedded image, or None when unusable."""
    try:
        image = shape.image
        found = (image.blob, image.ext, image.sha1)  # ext reads the image header
    except _IMAGE_ERRORS:
        return None
    return found if _readable(found[0]) else None


def _fill(parent, part, colors: dict[str, str]) -> Fill | None:
    """Read a solid, gradient (first stop) or picture fill under *parent*."""
    if parent is None:
        return None
    for tag in ("a:solidFill", "a:gradFill"):
        element = parent.find(qn(tag))
        if element is not None:
            if tag == "a:gradFill":
                element = element.find(f"{qn('a:gsLst')}/{qn('a:gs')}")
            color = _color(element, colors)
            return Fill(color=color) if color else None
    blip = parent.find(f"{qn('a:blipFill')}/{qn('a:blip')}")
    if blip is not None and blip.get(qn("r:embed")):
        try:
            image = part.related_part(blip.get(qn("r:embed")))
            blob, ext = image.blob, image.ext
        except _IMAGE_ERRORS:
            return None  # a broken relationship: treat as no background
        return Fill(image=blob, image_ext=ext) if _readable(blob) else None
    return None


def _covers(shape, width: int, height: int) -> bool:
    if None in (shape.left, shape.top, shape.width, shape.height):
        return False
    return (shape.left <= width * 0.02 and shape.top <= height * 0.02
            and shape.left + shape.width >= width * 0.98
            and shape.top + shape.height >= height * 0.98)


def _own_fill(owner, colors: dict[str, str], width: int, height: int) -> Fill | None:
    """The background a slide, layout or master sets itself, or None when it sets none.

    That is its own fill, a background style (taken as its base color, or _PLAIN when
    that is the theme's light color), or a full-slide picture or box.
    """
    bg = owner._element.cSld.bg
    if bg is not None:
        ref = bg.find(qn("p:bgRef"))
        color = _color(ref, colors)
        if color:
            return _PLAIN if color == colors.get("lt1") else Fill(color=color)
        found = _fill(bg.find(qn("p:bgPr")), owner.part, colors)
        if found is not None:
            return found
    for shape in list(owner.shapes)[:3]:
        if not _covers(shape, width, height):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            picture = _picture(shape)
            if picture is not None:
                return Fill(image=picture[0], image_ext=picture[1])
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            found = _fill(shape._element.spPr, owner.part, colors)
            if found is not None and found.color:
                return found
    return None


def _key(fill: Fill | None):
    return None if fill is None else (fill.color or hash(fill.image))


def _common_background(fills: list[Fill | None]) -> Fill | None:
    """The most common background; the plain default (None) counts too."""
    if not fills:
        return None
    winner = Counter(_key(fill) for fill in fills).most_common(1)[0][0]
    return next(fill for fill in fills if _key(fill) == winner)


def _logo(slides, width: int, height: int) -> Logo | None:
    """A small picture repeated in the same place on at least half the slides."""
    if len(slides) < 2:
        return None
    counts: Counter = Counter()
    samples = {}
    for slide in slides:
        seen = set()
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE or None in (shape.width, shape.height):
                continue
            picture = _picture(shape)
            if picture is None or shape.width * shape.height > 0.05 * width * height:
                continue
            key = (picture[2], round(shape.left / width, 2), round(shape.top / height, 2))
            if key not in seen:
                seen.add(key)
                counts[key] += 1
                samples.setdefault(key, (picture, shape))
    if not counts:
        return None
    key, count = counts.most_common(1)[0]
    if count < 2 or count * 2 < len(slides):
        return None
    picture, shape = samples[key]
    return Logo(picture[0], picture[1], shape.left / width, shape.top / height,
                shape.width / width, shape.height / height)


def _read_deck(path: str | Path) -> _Deck:
    prs = _open(path)
    width, height = prs.slide_width, prs.slide_height
    presentation = prs.part._element
    slides = list(prs.slides)
    themes: dict = {}  # master part name: (master, theme)
    used: Counter = Counter()
    title_runs: list[_Run] = []
    body_runs: list[_Run] = []
    fills, backgrounds = [], []
    for slide in slides:
        layout = slide.slide_layout
        master = layout.slide_master
        name = master.part.partname
        if name not in themes:
            themes[name] = (master, _theme(master))
        theme = themes[name][1]
        used[name] += 1
        colors = _colors(theme, _color_map(slide, layout, master))
        shapes = [(shape, _runs(shape, master, presentation, colors, theme.major, theme.minor))
                  for shape in _iter_shapes(slide.shapes)
                  if shape.has_text_frame and shape.text_frame.text.strip()]
        title = _title_shape(shapes, height)
        for shape, runs in shapes:
            (title_runs if shape is title else body_runs).extend(runs)
        own = _own_fill(slide, colors, width, height)
        fills.append(own)
        effective = (own or _own_fill(layout, colors, width, height)
                     or _own_fill(master, colors, width, height))
        backgrounds.append(None if effective == _PLAIN else effective)
    if used:
        master, theme = themes[used.most_common(1)[0][0]]
    else:
        master = prs.slide_master
        theme = _theme(master)
    colors = _colors(theme, _color_map(master))
    if not slides:
        own = _own_fill(master, colors, width, height)
        backgrounds = [None if own == _PLAIN else own]
    return _Deck(width, height, colors, theme, title_runs, body_runs, fills, backgrounds,
                 _logo(slides, width, height))


def _mostly_on_slide(runs: list[_Run]) -> bool:
    total = sum(run.weight for run in runs)
    return total > 0 and 2 * sum(run.weight for run in runs if run.on_slide) >= total


def look_is_on_slides(path: str | Path) -> bool:
    """True when the slides override their layouts.

    That is: at least half the slides set their own background, a logo repeats on the
    slides, or at least half the title or body text has a font or color that its
    layout does not give. A background the slides only inherit does not count.
    """
    deck = _read_deck(path)
    if not deck.fills:
        return False
    own = sum(1 for fill in deck.fills if fill is not None)
    return (own * 2 >= len(deck.fills) or deck.logo is not None
            or _mostly_on_slide(deck.title_runs) or _mostly_on_slide(deck.body_runs))


def capture_style(path: str | Path) -> DeckStyle:
    """Read size, theme colors, text styles, backgrounds and a repeated logo."""
    deck = _read_deck(path)
    text = deck.colors.get("tx1")
    backgrounds = deck.backgrounds
    return DeckStyle(deck.width, deck.height, deck.theme.scheme,
                     _dominant(deck.title_runs, deck.theme.major, text),
                     _dominant(deck.body_runs, deck.theme.minor, text),
                     backgrounds[0] if backgrounds else None,
                     _common_background(backgrounds[1:] or backgrounds), deck.logo)


def describe_capture(style: DeckStyle, skipped: tuple[str, ...] = ()) -> str:
    """One plain sentence naming what the new template took from the deck."""
    taken = ["colors"]
    if style.title.font or style.body.font:
        taken.append("fonts")
    if (style.background or style.title_background) and "background" not in skipped:
        taken.append("background")
    if style.logo and "logo" not in skipped:
        taken.append("logo")
    return f"Taken from your deck: {', '.join(taken)}. Its slide designs are not copied."
