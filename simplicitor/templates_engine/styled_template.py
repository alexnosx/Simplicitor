# simplicitor/templates_engine/styled_template.py
"""Build a fillable template that carries a captured deck look.

The base is the app's bundled copy of python-pptx's default template (the
package's own copy is unreachable in the packaged app). Its layouts have
title and content placeholders, so the existing renderer can fill it.
"""
import io
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml.shapes.picture import CT_Picture

from templates_engine.deck_style import DeckStyle, Fill, Logo, TextStyle

BASE_TEMPLATE = Path(__file__).resolve().parent.parent / "templates" / "pptx_default.pptx"
_FILLS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")
# Reading an image happens in memory: a failure skips that image. Saving the template
# is outside these guards, so a disk error still raises.
_IMAGE_ERRORS = (OSError, ValueError, Image.DecompressionBombError)


def _field(name: str, idx: int, kind: str, required: bool, **limit) -> dict:
    return {"name": name, "placeholder_idx": idx, "kind": kind, "required": required, **limit}


# Layout indices of the default template: 0 Title Slide, 1 Title and Content,
# 2 Section Header. Field limits match the built-in business_pitch template.
STYLED_SLIDE_TYPES = {
    "title": {"layout_index": 0, "fields": [
        _field("title", 0, "text", True, max_chars=60),
        _field("subtitle", 1, "text", False, max_chars=120)]},
    "agenda": {"layout_index": 1, "fields": [
        _field("heading", 0, "text", True, max_chars=40),
        _field("items", 1, "bullets", True, max_items=5)]},
    "content": {"layout_index": 1, "fields": [
        _field("heading", 0, "text", True, max_chars=60),
        _field("bullets", 1, "bullets", False, max_items=6)]},
    "closing": {"layout_index": 2, "fields": [
        _field("heading", 0, "text", True, max_chars=60),
        _field("statement", 1, "text", True, max_chars=160)]},
}


def styled_manifest(name: str, deck_name: str) -> dict:
    """The manifest for a template built by build_styled_template()."""
    return {"name": name, "type": "pptx", "template_file": "template.pptx",
            "description": f"Made from the look of {deck_name}.",
            "slide_types": STYLED_SLIDE_TYPES}


def _luminance(color: str) -> float:
    def channel(value: int) -> float:
        value /= 255
        return value / 12.92 if value <= 0.03928 else ((value + 0.055) / 1.055) ** 2.4
    red, green, blue = (int(color[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * channel(red) + 0.7152 * channel(green) + 0.0722 * channel(blue)


def _contrast(first: str, second: str) -> float:
    light, dark = sorted((_luminance(first), _luminance(second)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


def _average(fill: Fill | None) -> str:
    """The background's color, or a picture's average color; white when unknown."""
    if fill is None:
        return "FFFFFF"
    if fill.color:
        return fill.color
    try:
        with Image.open(io.BytesIO(fill.image)) as image:
            red, green, blue = image.convert("RGB").resize((1, 1)).getpixel((0, 0))
    except _IMAGE_ERRORS:
        return "FFFFFF"
    return f"{red:02X}{green:02X}{blue:02X}"


def readable(color: str | None, background: str) -> str:
    """Keep the deck's text color when it reads at 3:1 or better; else black or white."""
    if color and _contrast(color, background) >= 3.0:
        return color
    black, white = _contrast("000000", background), _contrast("FFFFFF", background)
    return "000000" if black >= white else "FFFFFF"


def _set_color(rpr, color: str) -> None:
    for tag in _FILLS:
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
    solid = parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{color}"/></a:solidFill>')
    line = rpr.find(qn("a:ln"))
    rpr.insert(list(rpr).index(line) + 1 if line is not None else 0, solid)


def _resize(prs, width: int, height: int) -> None:
    """Change the slide size and scale every master and layout shape that has its own position."""
    scale_x, scale_y = width / prs.slide_width, height / prs.slide_height
    prs.slide_width, prs.slide_height = width, height
    for owner in (prs.slide_master, *prs.slide_layouts):
        for shape in owner.shapes:
            if shape._element.spPr.find(qn("a:xfrm")) is None:
                continue  # inherits its position from the master
            shape.left, shape.top = round(shape.left * scale_x), round(shape.top * scale_y)
            shape.width, shape.height = round(shape.width * scale_x), round(shape.height * scale_y)


def _apply_theme(prs, style: DeckStyle) -> None:
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    elements = root.find(qn("a:themeElements"))
    if style.color_scheme:
        elements.replace(elements.find(qn("a:clrScheme")), etree.fromstring(style.color_scheme))
    fonts = elements.find(qn("a:fontScheme"))
    for kind, font in (("a:majorFont", style.title.font), ("a:minorFont", style.body.font)):
        if font:
            fonts.find(f"{qn(kind)}/{qn('a:latin')}").set("typeface", font)
    part.blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _apply_text(master, style_name: str, text: TextStyle, color: str, size_range: tuple) -> None:
    levels = master._element.find(f"{qn('p:txStyles')}/{qn('p:' + style_name)}")
    for level in levels:
        rpr = level.find(qn("a:defRPr"))
        if rpr is not None:
            _set_color(rpr, color)
    first = levels.find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    if text.size_pt:
        low, high = size_range
        first.set("sz", str(round(min(max(text.size_pt, low), high) * 100)))
    if text.bold is not None:
        first.set("b", "1" if text.bold else "0")


def _title_slide_colors(prs, title: str, body: str) -> None:
    """Override title-slide text colors where its background differs from the others."""
    for shape in prs.slide_layouts[0].placeholders:
        color = title if shape.placeholder_format.idx == 0 else body
        styles = shape._element.find(f"{qn('p:txBody')}/{qn('a:lstStyle')}")
        if styles is None:
            continue
        level = styles.find(qn("a:lvl1pPr"))
        if level is None:
            level = parse_xml(f'<a:lvl1pPr {nsdecls("a")}/>')
            default = styles.find(qn("a:defPPr"))
            styles.insert(list(styles).index(default) + 1 if default is not None else 0, level)
        rpr = level.find(qn("a:defRPr"))
        if rpr is None:
            rpr = parse_xml(f'<a:defRPr {nsdecls("a")}/>')
            extension = level.find(qn("a:extLst"))
            level.insert(list(level).index(extension) if extension is not None else len(level), rpr)
        _set_color(rpr, color)


def _set_background(owner, fill: Fill) -> None:
    if fill.color:
        background = owner.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor.from_string(fill.color)
        return
    _, rid = owner.part.get_or_add_image_part(io.BytesIO(fill.image))
    properties = owner._element.cSld.get_or_add_bgPr()
    for tag in _FILLS:
        for old in properties.findall(qn(tag)):
            properties.remove(old)
    properties.insert(0, parse_xml(
        f'<a:blipFill {nsdecls("a", "r")} dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/>'
        '<a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'))


def _add_logo(prs, logo: Logo) -> None:
    master = prs.slide_master
    _, rid = master.part.get_or_add_image_part(io.BytesIO(logo.image))
    width, height = prs.slide_width, prs.slide_height
    # On a master, _next_shape_id also counts the large slide-layout IDs.
    shape_id = max(int(value) for value in master.shapes._spTree.xpath(".//p:cNvPr/@id")) + 1
    picture = CT_Picture.new_pic(shape_id, "Logo", "", rid,
                                 round(logo.left * width), round(logo.top * height),
                                 round(logo.width * width), round(logo.height * height))
    master.shapes._spTree.append(picture)


def _apply_pictures(prs, style: DeckStyle) -> list[str]:
    """Set backgrounds and the logo; return the parts whose images could not be read."""
    skipped: list[str] = []
    steps = [("background", style.background, prs.slide_master)]
    if style.title_background != style.background:
        steps.append(("title slide background", style.title_background, prs.slide_layouts[0]))
    for kind, fill, owner in steps:
        if fill is None:
            continue
        try:
            _set_background(owner, fill)
        except _IMAGE_ERRORS:
            skipped.append(kind)
    if style.logo is not None:
        try:
            _add_logo(prs, style.logo)
        except _IMAGE_ERRORS:
            skipped.append("logo")
    return skipped


def build_styled_template(style: DeckStyle, out_path: str | Path) -> list[str]:
    """Write a fillable template with *style*; return parts of the look that were skipped."""
    out_path = Path(out_path)
    prs = Presentation(str(BASE_TEMPLATE))
    _resize(prs, style.slide_width, style.slide_height)
    _apply_theme(prs, style)
    skipped = _apply_pictures(prs, style)
    # Choose text colors against the backgrounds that were really applied.
    content = _average(None if "background" in skipped else style.background)
    _apply_text(prs.slide_master, "titleStyle", style.title,
                readable(style.title.color, content), (24, 54))
    _apply_text(prs.slide_master, "bodyStyle", style.body,
                readable(style.body.color, content), (14, 32))
    if (style.title_background is not None and style.title_background != style.background
            and "title slide background" not in skipped):
        cover = _average(style.title_background)
        _title_slide_colors(prs, readable(style.title.color, cover),
                            readable(style.body.color, cover))
    temporary = out_path.with_name(out_path.name + ".tmp")
    try:
        prs.save(str(temporary))
        temporary.replace(out_path)
    finally:
        temporary.unlink(missing_ok=True)
    return skipped
