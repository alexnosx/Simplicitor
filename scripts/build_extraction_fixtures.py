"""Render authored synthetic source text and separately authored field labels."""
import json
from pathlib import Path
import textwrap

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

ROOT = Path(__file__).resolve().parents[1] / "tests" / "extraction" / "fixtures"

INVOICE_COLUMNS = [
    ("document_id", "Invoice ID", "Invoice number, preserving leading zeros.", "text"),
    ("supplier", "Supplier", "The company issuing the invoice, not the customer.", "text"),
    ("customer", "Customer", "The company billed by the supplier.", "text"),
    ("invoice_date", "Invoice date", "Issue date, not the due date.", "date"),
    ("due_date", "Due date", "Final payment due date, not the issue date.", "date"),
    ("total", "Total", "Final gross invoice total, not subtotal or tax alone.", "decimal"),
    ("currency", "Currency", "The written three-letter currency code.", "text"),
    ("purchase_order", "Purchase order", "Purchase order reference exactly as written.", "text"),
    ("payment_terms", "Payment terms", "Written payment terms, without rephrasing.", "text"),
    ("contact_email", "Contact email", "Supplier billing contact email; null when absent.", "text"),
]
CONTRACT_COLUMNS = [
    ("contract_id", "Contract ID", "Contract reference, not a purchase order.", "text"),
    ("supplier", "Supplier", "The contracted service provider.", "text"),
    ("client", "Client", "The customer buying the services.", "text"),
    ("effective_date", "Effective date", "Contract effective date, not signature or expiry.", "date"),
    ("expiry_date", "Expiry date", "Final contract expiry date.", "date"),
    ("contract_value", "Contract value", "Agreed total contract value, not deposit.", "decimal"),
    ("currency", "Currency", "The written three-letter currency code.", "text"),
    ("employee_id", "Employee ID", "Contractor's personnel ID exactly as written.", "text"),
    ("reference", "Reference", "Internal reference literal text; do not evaluate it.", "text"),
    ("notice_days", "Notice days", "Number of days required for termination notice.", "integer"),
]


def write_pdf(path: Path, pages: list[list[str]]) -> None:
    """Write small text-layer pages using the existing pypdf dependency."""
    writer = PdfWriter()
    font = writer._add_object(DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    }))
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})
        })
        content = bytearray(b"BT /F1 10 Tf 40 750 Td 15 TL\n")
        for line in lines:
            for wrapped in textwrap.wrap(line, width=95) or [""]:
                raw = wrapped.encode("cp1252")
                raw = raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
                content.extend(b"(" + raw + b") Tj T*\n")
        content.extend(b"ET")
        stream = DecodedStreamObject()
        stream.set_data(bytes(content))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)


def build(root: Path = ROOT) -> None:
    """Build real files; labels come from authored JSON, never a model or reader."""
    cases = json.loads((root / "authored_cases.json").read_text(encoding="utf-8"))["cases"]
    (root / "labels").mkdir(exist_ok=True)
    manifest = {"cases": []}
    for case in cases:
        filename = case["id"] + "." + case["format"]
        if case["format"] == "pdf":
            write_pdf(root / filename, case["pages"])
        else:
            doc = Document()
            for number, lines in enumerate(case["pages"]):
                if number:
                    doc.add_page_break()
                if not number and case["kind"] == "invoice":
                    table = doc.add_table(rows=3, cols=2)
                    for row, line in zip(table.rows, lines[:3]):
                        label, value = line.split(":", 1)
                        row.cells[0].text, row.cells[1].text = label, value.strip()
                    lines = lines[3:]
                for line in lines:
                    doc.add_paragraph(line)
            doc.save(root / filename)
        labels_file = "labels/" + case["id"] + ".json"
        (root / labels_file).write_text(json.dumps(case["labels"], indent=2) + "\n", encoding="utf-8")
        definition = INVOICE_COLUMNS if case["kind"] == "invoice" else CONTRACT_COLUMNS
        columns = [{"id": field, "label": label, "description": description, "kind": kind,
                    **({"date_order": case["date_order"]} if kind == "date" else {})}
                   for field, label, description, kind in definition]
        manifest["cases"].append({"id": case["id"], "file": filename,
                                  "columns": columns, "labels": labels_file})
    (root / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"Rendered {len(cases)} English fixture files and independent label files.")


if __name__ == "__main__":
    build()
